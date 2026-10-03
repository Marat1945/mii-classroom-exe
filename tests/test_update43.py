"""4.3: ZIP дня й стеження за папками, спільні назви, наперед, спливні підказки, колесо, OpenAI в дні, пам'ять програми."""
import json
import os
import shutil
import tempfile
import time
import tkinter as tk
import unittest
import zipfile
from contextlib import ExitStack
from pathlib import Path
from unittest import mock

from docx import Document
from PIL import Image

from test_update36 import TkCase
from test_update38 import config_with
from test_update42 import SharedCourseEditorTests, make_lecture, make_picture

from classroom_assistant import (attachments, chatgpt_ui, datepicker, data_tools, editor_core, editor_ui, engine, file_match, gui,
                                 lecture_inbox, material_library, samples_ui, toast, wheel)
from classroom_assistant.engine import Lesson, build_calendar, combined_label, lesson_base_name, same_day_label

REPO_DATA = Path(__file__).resolve().parents[1] / "data"
FIXTURE = Path(__file__).resolve().parent / "fixtures" / "Зразок_КТП_вчителя.docx"


def L(day, period, stream, topic="Тема", number=1):
    return Lesson(day, 0, period, "08:30", "09:15", stream, stream, "p", number, topic, "§ 1", "ktp.docx")


class CombinedLabelTests(unittest.TestCase):
    def test_same_day_parallels_share_one_label(self):
        self.assertEqual(combined_label(["8-Б ГО", "8-В ГО", "8-Г ГО"]), "8-Б-В-Г ГО")
        self.assertEqual(combined_label(["8-Г ГО", "8-Б ГО"]), "8-Б-Г ГО")
        self.assertEqual(combined_label(["9-Б ІУ", "9-Б ВІ"]), "9-Б ІУ")                 # різні предмети — не склеюємо
        self.assertEqual(combined_label(["10 ІУ профіль"]), "10 ІУ профіль")
        self.assertEqual(combined_label(["9-Б Право", "9-В Право"]), "9-Б-В Право")

    def test_other_days_keep_their_own_class(self):
        mine = L("2026-10-05", 1, "8-Б ГО")
        group = [mine, L("2026-10-05", 2, "8-В ГО"), L("2026-10-06", 1, "8-Г ГО")]
        self.assertEqual(same_day_label(mine, group), "8-Б-В ГО")
        self.assertEqual(same_day_label(group[2], group), "8-Г ГО")

    def test_name_and_matching_of_a_combined_lesson(self):
        lessons = [L("2026-10-05", 1, "8-Б ГО", "Суспільство"), L("2026-10-05", 2, "8-В ГО", "Суспільство"),
                   L("2026-10-05", 3, "8-Г ГО", "Суспільство"), L("2026-10-06", 1, "9-Б ВІ", "Інше")]
        name = lesson_base_name(lessons[0], label="8-Б-В-Г ГО")
        self.assertEqual(name, "8-Б-В-Г ГО, Урок 05.10 — Суспільство")
        self.assertEqual(file_match.expand_stream("8-Б-В-Г ГО"), ["8-Б ГО", "8-В ГО", "8-Г ГО"])
        found = file_match.find_lesson(lessons, file_match.parse_name(name))
        self.assertEqual(found.stream, "8-Б ГО")                                      # головний — найраніший урок
        self.assertEqual(file_match.expand_stream("8-Б ГО"), ["8-Б ГО"])


class PromptTests(unittest.TestCase):
    def test_day_prompt_uses_combined_names_and_asks_for_one_zip(self):
        from classroom_assistant.chatgpt_bridge import build_prompt
        lessons = [x for x in build_calendar(engine.read_json("Налаштування.json"))
                   if x.day == "2026-10-05" and x.stream == "8-Б ГО"]
        lesson = lessons[0]
        text = chatgpt_ui.build_day_prompt([lesson], {lesson.unique_key: "8-Б-В-Г ГО"})
        self.assertIn("ZIP-архів «05.10.26.zip»", text)
        self.assertIn("«8-Б-В-Г ГО, Урок 05.10 — ", text)
        self.assertIn("«8-Б-В-Г ГО, Урок 05.10 — " + lesson.topic.strip(". ")[:20], text)
        self.assertIn("для класів «8-Б-В-Г ГО»", build_prompt(lesson, label="8-Б-В-Г ГО"))
        self.assertNotIn("для класів", build_prompt(lesson))


class ArchiveAndWatcherTests(unittest.TestCase):
    def test_zip_is_unpacked_safely_and_only_word_and_images(self):
        with tempfile.TemporaryDirectory() as folder:
            archive = Path(folder) / "05.10.26.zip"
            with zipfile.ZipFile(archive, "w") as z:
                z.writestr("8-Б ГО, Урок 05.10 — Тема.docx", b"word")
                z.writestr("підпапка/8-Б ГО, Урок 05.10 — Тема.png", b"png")
                z.writestr("../evil.docx", b"x")
                z.writestr("запускати.exe", b"MZ")
                z.writestr("нотатки.txt", b"t")
            out = lecture_inbox.extract_archive(archive, Path(folder) / "розпаковано")
            names = sorted(p.name for p in out)
            self.assertEqual(names, ["8-Б ГО, Урок 05.10 — Тема.docx", "8-Б ГО, Урок 05.10 — Тема.png", "evil.docx"])
            self.assertTrue(all(Path(folder, "розпаковано") in p.parents for p in out))      # нічого поза папкою

    def test_candidate_rules_for_downloads_and_own_inbox(self):
        ok = lambda name, strict: lecture_inbox.is_candidate(Path(name), strict)
        self.assertTrue(ok("9-Б ВІ, Урок 05.10 — Тема.docx", True))
        self.assertTrue(ok("05.10.26.zip", True))
        self.assertFalse(ok("Звіт.docx", True))                     # чужий файл у «Завантаженнях» не чіпаємо
        self.assertFalse(ok("9-Б ВІ, Урок 05.10 — Тема.docx.crdownload", True))
        self.assertFalse(ok("~$9-Б ВІ, Урок 05.10 — Тема.docx", True))
        self.assertTrue(ok("Звіт.docx", False))                     # власна папка програми — будь-який Word
        self.assertFalse(ok("музика.mp3", False))

    def test_watcher_waits_until_the_file_stops_growing_and_never_repeats(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            file = folder / "9-Б ВІ, Урок 05.10 — Тема.docx"
            file.write_bytes(b"a")
            watcher = lecture_inbox.InboxWatcher()
            self.assertEqual(watcher.poll([(folder, True)]), [])                  # перший раз — лише запам'ятали
            file.write_bytes(b"aaaa")                                              # ще дописується
            self.assertEqual(watcher.poll([(folder, True)]), [])
            self.assertEqual([p.name for p, _ in watcher.poll([(folder, True)])], [file.name])   # стабільний
            watcher.mark_done(file)
            self.assertEqual(watcher.poll([(folder, True)]), [])
            self.assertEqual(lecture_inbox.InboxWatcher(watcher.export()).poll([(folder, True)]), [])   # і після перезапуску

    def test_old_downloads_are_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            file = Path(folder) / "9-Б ВІ, Урок 05.10 — Тема.docx"
            file.write_bytes(b"a")
            watcher = lecture_inbox.InboxWatcher()
            future = time.time_ns() + 10**9
            watcher.poll([(folder, True)], since_ns=future)
            self.assertEqual(watcher.poll([(folder, True)], since_ns=future), [])

    def test_move_to_done_keeps_the_file(self):
        with tempfile.TemporaryDirectory() as folder:
            file = Path(folder) / "a.docx"
            file.write_bytes(b"x")
            target = lecture_inbox.move_to_done(file)
            self.assertTrue(target.is_file() and not file.exists())
            self.assertEqual(target.parent.name, lecture_inbox.DONE_NAME)


class WheelTests(TkCase):
    NEEDS_ROOT = True

    def test_direction_and_canvas_scrolling(self):
        fake = lambda **k: type("E", (), k)()
        self.assertEqual(wheel.direction(fake(delta=120)), -1)
        self.assertEqual(wheel.direction(fake(delta=-120)), 1)
        self.assertEqual(wheel.direction(fake(num=4)), -1)
        self.assertEqual(wheel.direction(fake(num=5)), 1)
        top = tk.Toplevel(self.root)
        canvas = tk.Canvas(top, height=100, yscrollcommand=lambda *a: None)
        canvas.pack()
        inner = tk.Frame(canvas)
        canvas.create_window(0, 0, window=inner, anchor="nw")
        for i in range(60):
            tk.Label(inner, text=str(i)).pack()
        inner.update_idletasks()
        canvas.configure(scrollregion=canvas.bbox("all"))
        top.update()
        self.assertIs(wheel.scroll_target(inner.winfo_children()[3]), canvas)       # Label усередині Canvas
        event = fake(delta=-120, widget=canvas, x_root=canvas.winfo_rootx() + 5, y_root=canvas.winfo_rooty() + 5)
        self.assertEqual(wheel.on_wheel(event), "break")
        self.assertGreater(canvas.yview()[0], 0.0)
        listbox = tk.Listbox(top)
        listbox.pack()
        self.assertIsNone(wheel.scroll_target(listbox))                               # Listbox прокручує сам Tk
        top.destroy()

    def test_steps_helper_big_step_with_shift(self):
        top = tk.Toplevel(self.root)
        calls = []
        wheel.bind_steps(top, lambda step, big: calls.append((step, big)))
        top.update()
        handler = top.bind("<MouseWheel>")
        self.assertTrue(handler)
        top.destroy()


class DateWheelAndBannerTests(TkCase):
    NEEDS_ROOT = False

    def test_wheel_over_the_date_moves_days_and_week(self):
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(gui.messagebox, "showinfo"))
            app = gui.MainApp()
            self.addCleanup(app.destroy)
            app.datevar.set("05.10.2026")
            app.update_day()
            app.shift(1)
            self.assertEqual(app.datevar.get(), "06.10.2026")
            app.date_entry.event_generate("<Button-5>")                                   # коліщатко вниз
            app.update()
            self.assertEqual(app.datevar.get(), "07.10.2026")                              # +1 день
            app.date_entry.event_generate("<Button-4>")                                   # вгору
            app.update()
            self.assertEqual(app.datevar.get(), "06.10.2026")
            app.weekday.event_generate("<Shift-Button-5>")                                # Shift: +1 тиждень
            app.update()
            self.assertEqual(app.datevar.get(), "13.10.2026")

    def test_author_is_drawn_on_the_banner(self):
        with ExitStack() as stack:
            app = gui.MainApp()
            self.addCleanup(app.destroy)
            app.update()
            texts = []
            for child in app.winfo_children():
                if child.winfo_class() == "Canvas":
                    for item in child.find_all():
                        if child.type(item) == "text":
                            texts.append(child.itemcget(item, "text"))
            self.assertIn("Розробник програми — вчитель історії Пасічник Іван Олегович", texts)
            self.assertTrue(any("Помічник учителя Classroom" in x for x in texts))


class CalendarWheelTests(TkCase):
    def test_wheel_turns_the_calendar_month_and_shift_turns_the_year(self):
        from datetime import date
        host = tk.Toplevel(self.root)
        host.geometry("300x200+10+10")
        host.update()

        def close():
            try:
                host.destroy()
            except tk.TclError:
                pass
        self.addCleanup(close)
        dialog = datepicker.CalendarDialog(host, initial=date(2026, 10, 5))
        dialog.update()

        def day_button():
            dialog.update()                                                       # після перемальовування
            return next(k for k in dialog.body.winfo_children()
                        if k.winfo_class() == "Button" and k.winfo_ismapped())
        day_button().event_generate("<Button-5>", x=2, y=2)                  # коліщатко вниз над кнопкою дня
        dialog.update()
        self.assertEqual((dialog.year, dialog.month), (2026, 11))               # → наступний місяць
        day_button().event_generate("<Button-4>", x=2, y=2)
        day_button().event_generate("<Button-4>", x=2, y=2)
        self.assertEqual((dialog.year, dialog.month), (2026, 9))
        day_button().event_generate("<Shift-Button-5>", x=2, y=2)               # Shift → рік
        dialog.update()
        self.assertEqual((dialog.year, dialog.month), (2027, 9))


class ToastTests(TkCase):
    def test_toast_has_no_buttons_and_disappears_by_itself(self):
        toast_window = toast.show_toast(self.root, "✓ Прикріплено", ms=150)
        self.assertIsNotNone(toast_window)
        self.root.update()
        buttons = [w for w in toast_window.winfo_children() if w.winfo_class() in ("Button", "TButton")]
        self.assertEqual(buttons, [])
        deadline = time.time() + 2
        while time.time() < deadline and toast_window.winfo_exists():
            self.root.update()
            time.sleep(0.02)
        self.assertFalse(toast_window.winfo_exists())


class SampleKtpTests(unittest.TestCase):
    def test_sample_puts_each_class_on_its_own_line_like_the_teachers_file(self):
        with tempfile.TemporaryDirectory() as folder:
            built = samples_ui.build_sample_ktp(Path(folder) / "s.docx")
            mine = Document(built).tables[0]
            theirs = Document(FIXTURE).tables[0]
            self.assertEqual([c.text.strip() for c in theirs.rows[2].cells][1].count("\n"), 2)
            self.assertEqual(mine.rows[2].cells[1].text.count("\n"), 2)
            self.assertEqual(len(mine.rows), len(theirs.rows))

    def test_both_files_are_read_to_the_same_dates_for_every_class(self):
        from classroom_assistant import ktp_import
        with tempfile.TemporaryDirectory() as folder:
            built = samples_ui.build_sample_ktp(Path(folder) / "s.docx")
            a = ktp_import.load_ktp(built, 2026)
            b = ktp_import.load_ktp(FIXTURE, 2026)
        self.assertEqual(len(a), len(b))
        self.assertEqual([x["topic"] for x in a], [x["topic"] for x in b])
        for lesson in b:
            self.assertEqual(set(lesson["class_dates"]), {"8-Б", "8-В", "8-Г"}, lesson["topic"]) \
                if "class_dates" in lesson else None


class PreferencesSurviveResetTests(unittest.TestCase):
    def test_reset_keeps_window_and_gpt_preferences_and_install_marker_is_not_backed_up(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "root"
            data = root / "data"
            data.mkdir(parents=True)
            for name in ("Налаштування.json", "Календарні плани.json"):
                shutil.copy(REPO_DATA / name, data / name)
            (data / "Стан.json").write_text(json.dumps({
                "drafts": {"a": 1}, "chatgpt_url": "https://chatgpt.com/c/abc", "watch_downloads": False,
                "column_order": ["a", "b"]}), encoding="utf-8")
            data_tools.write_install({"clean_start": "4.3"}, data)
            data_tools.reset_to_blank(root / "копія.zip", root=root, data_dir=data)
            state = json.loads((data / "Стан.json").read_text("utf-8"))
            self.assertEqual(state, {"chatgpt_url": "https://chatgpt.com/c/abc", "watch_downloads": False,
                                     "column_order": ["a", "b"]})
            self.assertEqual(data_tools.read_install(data), {"clean_start": "4.3"})         # мітка пережила скидання
            self.assertNotIn("data/install_state.json", zipfile.ZipFile(root / "копія.zip").namelist())


class DayDialogTests(TkCase):
    NEEDS_ROOT = False

    def setUp(self):
        super().setUp()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root_dir = Path(self.tmp.name) / "root"
        (self.root_dir / "data").mkdir(parents=True)
        for name in ("Налаштування.json", "Календарні плани.json"):
            shutil.copy(REPO_DATA / name, self.root_dir / "data" / name)
        stack = ExitStack()
        for patcher in (mock.patch.object(engine, "DATA", self.root_dir / "data"),
                        mock.patch.object(engine, "ROOT", self.root_dir),
                        mock.patch.object(gui, "DATA", self.root_dir / "data"),
                        mock.patch.object(gui, "ROOT", self.root_dir),
                        mock.patch.object(material_library, "LIBRARY_ROOT", self.root_dir / "Бібліотека"),
                        mock.patch.object(material_library, "INDEX", self.root_dir / "Бібліотека" / "Індекс.json"),
                        mock.patch.object(attachments, "ATTACHMENT_ROOT", self.root_dir / "Вкладення"),
                        mock.patch.object(lecture_inbox, "downloads_dir", lambda: self.root_dir / "Завантаження")):
            stack.enter_context(patcher)
        self.addCleanup(stack.close)
        (self.root_dir / "Завантаження").mkdir()
        self.toasts, self.warnings = [], []
        stack.enter_context(mock.patch.object(gui, "show_toast", side_effect=lambda p, text, *a, **k: self.toasts.append(text)))
        stack.enter_context(mock.patch.object(chatgpt_ui, "show_toast", side_effect=lambda p, text, *a, **k: self.toasts.append(text)))
        stack.enter_context(mock.patch.object(gui.messagebox, "showwarning", side_effect=lambda *a, **k: self.warnings.append(a[1])))
        stack.enter_context(mock.patch.object(gui.messagebox, "showinfo"))
        stack.enter_context(mock.patch.object(chatgpt_ui.messagebox, "showwarning", side_effect=lambda *a, **k: self.warnings.append(a[1])))
        stack.enter_context(mock.patch.object(gui.messagebox, "askyesno", side_effect=AssertionError("запитання!")))
        self.app = gui.MainApp()
        self.addCleanup(lambda: self._close())
        self.app.sync_requests = []
        self.app.request_sync = lambda delay=900: self.app.sync_requests.append(delay)
        self.lessons = [x for x in build_calendar(self.app.cfg) if x.status == "готово"]
        self.files = Path(self.tmp.name) / "файли"
        self.files.mkdir()

    def _close(self):
        try:
            self.app.destroy()
        except tk.TclError:
            pass

    def lecture(self, lesson, label=None):
        name = lesson_base_name(lesson, label=label)
        first = f"{label or lesson.stream}, Урок {lesson.day[8:10]}.{lesson.day[5:7]} — {lesson.topic}"
        return make_lecture(self.files / (name + ".docx"), first), make_picture(self.files / (name + ".png"))

    def day_lessons(self, day="2026-10-05"):
        seen, result = set(), []
        for x in self.lessons:
            key = (x.day, x.stream)
            if x.day == day and key not in seen:
                seen.add(key)
                result.append(x)
        return result

    def test_zip_of_the_day_is_imported_in_one_go_and_toast_names_word_and_images(self):
        picked = [self.day_lessons()[0], self.day_lessons()[3]]
        archive = Path(self.tmp.name) / "05.10.26.zip"
        with zipfile.ZipFile(archive, "w") as z:
            for lesson in picked:
                word, picture = self.lecture(lesson)
                z.write(word, word.name)
                z.write(picture, picture.name)
        unmatched = self.app.import_lecture_files([archive])
        self.assertEqual(unmatched, [])
        self.assertIn(f"Розпізнано уроків: {len(picked)} · Word: {len(picked)} · зображень: {len(picked)}", self.toasts[-1])
        for lesson in picked:
            self.assertTrue(self.app.state["files"][lesson.unique_key]["complete"])
        self.assertEqual(self.warnings, [])                                              # жодного вікна з «ОК»
        self.assertTrue(self.app.sync_requests)                                          # Classroom буде перевірено

    def test_parallel_on_another_day_is_marked_as_prepared_in_advance(self):
        first = parallels = later = None
        for candidate in self.lessons:
            if not (candidate.stream.startswith("8-") and candidate.stream.endswith("ІУ")):
                continue
            found = material_library.parallel_matches(candidate, self.app.cfg, calendar=self.lessons)
            after = [x for x in found if x.day > candidate.day]
            if after:
                first, parallels, later = candidate, found, after
                break
        self.assertIsNotNone(first, "у даних має бути урок із паралеллю, що йде пізніше")
        own_label = same_day_label(first, parallels)
        word, picture = self.lecture(first, label=own_label)
        self.app.import_lecture_files([word, picture])
        for lesson in later:
            entry = self.app.state["files"][lesson.unique_key]
            self.assertTrue(entry["ahead"])
            self.assertEqual(entry["from"], f"{first.stream}, {first.day[8:10]}.{first.day[5:7]}")
            texts = [p.text for p in Document(entry["path"]).paragraphs if p.text.strip()]
            expected = same_day_label(lesson, parallels)               # у цей день клас(и) цього дня
            self.assertTrue(texts[0].startswith(f"{expected}, Урок {lesson.day[8:10]}.{lesson.day[5:7]}"), texts[0])
        stream_names = {x.stream for x in later}
        self.assertGreater(len(stream_names), 1)
        self.assertTrue(any("-" in same_day_label(x, parallels).split()[0][2:] for x in later))   # «8-Б-В ІУ»
        own = self.app.state["files"][first.unique_key]
        self.assertNotIn("from", own)                                  # сам урок — не «наперед»
        day = later[0].day
        self.app.datevar.set(f"{day[8:10]}.{day[5:7]}.{day[:4]}")
        self.app.update_day()
        statuses = [self.app.grid.item(i, "values")[-1] for i in self.app.grid.get_children()]
        marker = "наперед з " + f"{first.stream}, {first.day[8:10]}.{first.day[5:7]}"
        self.assertTrue(any(marker in s for s in statuses), statuses)

    def test_inbox_folder_and_downloads_are_picked_up_without_clicks(self):
        a, b = self.day_lessons()[0], self.day_lessons()[3]
        inbox = lecture_inbox.inbox_dir(self.root_dir)
        word, picture = self.lecture(a)
        shutil.copy(word, inbox / word.name)
        shutil.copy(picture, inbox / picture.name)
        word2, _ = self.lecture(b)
        shutil.copy(word2, self.root_dir / "Завантаження" / word2.name)
        (self.root_dir / "Завантаження" / "Звіт.docx").write_bytes("чужий".encode("utf-8"))
        self.app._watch_since_ns = 0
        for _ in range(3):                                                                # опитування: бачить → стабільний → обробка
            self.app._inbox_poll()
        self.assertTrue(self.app.state["files"][a.unique_key]["complete"])
        self.assertTrue(self.app.state["files"][b.unique_key]["complete"])
        self.assertTrue((inbox / lecture_inbox.DONE_NAME / word.name).is_file())          # власна папка → «Оброблено»
        self.assertTrue((self.root_dir / "Завантаження" / word2.name).is_file())          # «Завантаження» не чіпаємо
        self.assertTrue((self.root_dir / "Завантаження" / "Звіт.docx").is_file())
        self.app.state["watch_downloads"] = False
        self.assertEqual([f for f, strict in self.app._watch_sources() if strict], [])

    def test_day_dialog_shows_files_that_arrived_by_themselves(self):
        day = self.day_lessons()
        lessons = [day[0], day[3], day[4], day[5]]                      # без паралелей між собою
        dialog = chatgpt_ui.ChatGPTDayDialog(self.app, lessons)
        self.addCleanup(lambda: dialog.winfo_exists() and dialog.destroy())
        self.assertIn("Word — 0 · зображень — 0", dialog.note.cget("text"))
        word, picture = self.lecture(lessons[1])
        self.app.import_lecture_files([word, picture])
        dialog._poll_state()
        row = dialog.table.item("1", "values")
        self.assertEqual(str(row[5]), "✅")
        self.assertEqual(str(row[6]), "1")
        self.assertIn("Word — 1 · зображень — 1", dialog.note.cget("text"))

    def test_openai_button_exists_in_the_day_dialog_and_runs_batch_for_checked_lessons(self):
        lessons = self.day_lessons()[:3]
        dialog = chatgpt_ui.ChatGPTDayDialog(self.app, lessons)
        texts, stack = [], [dialog]
        while stack:
            widget = stack.pop()
            stack.extend(widget.winfo_children())
            if widget.winfo_class() == "TButton":
                texts.append(str(widget.cget("text")))
        self.assertIn("Платно через OpenAI API…", texts)
        dialog.checked = {0, 2}
        with mock.patch.object(chatgpt_ui.messagebox, "askyesno", return_value=True), \
                mock.patch.object(self.app, "batch_ai") as batch:
            dialog.use_api()
        batch.assert_called_once()
        self.assertEqual([x.unique_key for x in batch.call_args.kwargs["only"]],
                         [lessons[0].unique_key, lessons[2].unique_key])

    def test_attach_all_shows_a_short_toast_and_checks_classroom(self):
        lessons = self.day_lessons()[:2]
        dialog = chatgpt_ui.ChatGPTDayDialog(self.app, lessons)
        word, picture = self.lecture(lessons[0])
        dialog.docx[0], dialog.images[0] = word, [picture]
        dialog.checked = {0}
        dialog.attach_all()
        self.assertIn("Прикріплено: Word — 1, зображень — 1", self.toasts[-1])
        self.assertTrue(self.app.sync_requests)
        self.assertEqual(self.warnings, [])

    def test_local_draft_that_classroom_does_not_show_is_flagged_after_a_sync(self):
        lesson = self.day_lessons()[0]
        course = lesson.course_title
        self.app.state.setdefault("drafts", {})[lesson.unique_key] = {"id": "77", "created_at": 1.0}
        self.app.state.setdefault("course_ids", {})[course] = "555"
        self.app.datevar.set(f"{lesson.day[8:10]}.{lesson.day[5:7]}.{lesson.day[:4]}")
        self.app.update_day()
        before = [self.app.grid.item(i, "values")[-1] for i in self.app.grid.get_children()]
        self.assertIn("Створено в Google", before[0])
        self.assertNotIn("не знайдено", before[0])                                         # ще не синхронізували
        self.app.remote_classroom_entries["555"] = []
        self.app._last_sync_ts = time.time()
        self.app.update_day()
        after = [self.app.grid.item(i, "values")[-1] for i in self.app.grid.get_children()]
        self.assertIn("у Classroom не знайдено", after[0])
        self.app.remote_classroom_entries["555"] = [{"id": "77", "state": "DRAFT", "type": "Матеріал", "title": "x"}]
        self.app.update_day()
        self.assertIn("чернетка в Classroom", self.app.grid.item(self.app.grid.get_children()[0], "values")[-1])

    def test_attach_all_never_touches_a_lesson_that_already_has_a_classroom_draft(self):
        lessons = self.day_lessons()[:2]
        self.app.state.setdefault("drafts", {})[lessons[0].unique_key] = {"id": "9", "created_at": 1.0}
        dialog = chatgpt_ui.ChatGPTDayDialog(self.app, lessons)
        word, picture = self.lecture(lessons[0])
        dialog.docx[0], dialog.images[0] = word, [picture]
        dialog.checked = {0}
        dialog.attach_all()
        self.assertNotIn(lessons[0].unique_key, self.app.state.get("files", {}))        # чернетку не змінюємо
        self.assertTrue(any("уже має чернетку" in w for w in self.warnings), self.warnings)
        self.assertTrue(self.app.sync_requests)

    def test_work_survives_closing_and_reopening_the_program(self):
        """Те, що вчитель зробив, після закриття й нового запуску на місці (редактор + файли + чернетки)."""
        from classroom_assistant import editor_core
        lesson = self.day_lessons()[0]
        word, picture = self.lecture(lesson)
        self.app.import_lecture_files([word, picture])
        self.app.state.setdefault("drafts", {})[lesson.unique_key] = {"id": "1", "created_at": time.time()}
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(editor_core, "DATA", self.root_dir / "data"))
            stack.enter_context(mock.patch.object(editor_core, "ROOT", self.root_dir))
            stack.enter_context(mock.patch.object(editor_ui.messagebox, "askyesno", return_value=True))
            stack.enter_context(mock.patch.object(editor_ui.messagebox, "showinfo"))
            editor = editor_ui.SchoolEditor(self.app, self.app.update_day)
            editor.autosave_enabled = True
            editor.cfg["holidays"].append({"start": "2027-04-05", "end": "2027-04-06"})   # зміна без «Зберегти»
            self.assertTrue(editor.has_unsaved())
            self.app.on_close()                                                           # закрили програму
        with mock.patch.object(gui.messagebox, "showinfo"):
            reopened = gui.MainApp()
        self.addCleanup(reopened.destroy)
        self.assertIn({"start": "2027-04-05", "end": "2027-04-06"}, reopened.cfg["holidays"])
        self.assertTrue(reopened.state["files"][lesson.unique_key]["complete"])
        self.assertIn(lesson.unique_key, reopened.state["drafts"])
        self.assertEqual(len(attachments.files_for_lesson(reopened.state, lesson.unique_key)), 1)

    def test_closing_the_program_saves_editor_changes_and_state(self):
        saved = []
        fake_editor = mock.Mock()
        fake_editor.winfo_exists.return_value = True
        fake_editor.has_unsaved.return_value = True
        fake_editor.save.side_effect = lambda **kw: saved.append(kw) or True
        type(fake_editor).__name__ = "SchoolEditor"
        with mock.patch.object(self.app, "winfo_children", return_value=[fake_editor]), \
                mock.patch.object(gui, "save_state") as save_state, mock.patch.object(self.app, "destroy") as destroy:
            self.app.on_close()
        self.assertEqual(saved, [{"confirm": False, "close": False, "silent": True}])
        save_state.assert_called()
        destroy.assert_called_once()

    def test_closing_with_unsavable_editor_asks_instead_of_losing_work(self):
        fake_editor = mock.Mock()
        fake_editor.winfo_exists.return_value = True
        fake_editor.has_unsaved.return_value = True
        fake_editor.save.return_value = False
        type(fake_editor).__name__ = "SchoolEditor"
        with mock.patch.object(self.app, "winfo_children", return_value=[fake_editor]), \
                mock.patch.object(gui.messagebox, "askyesno", return_value=False) as ask, \
                mock.patch.object(self.app, "destroy") as destroy:
            self.app.on_close()
        self.assertTrue(ask.called)
        destroy.assert_not_called()                                                       # не закрилась, робота ціла


class EditorAutosaveTests(SharedCourseEditorTests):
    def test_autosave_writes_changes_without_closing_and_without_questions(self):
        editor = self.combined()
        editor.autosave_enabled = True
        editor.cfg["days"]["0"][0] = ["9-Б Право", "9-Б ГО"]
        self.assertTrue(editor.has_unsaved())
        with mock.patch.object(editor_ui, "persist", return_value="копія") as persist, \
                mock.patch.object(editor_ui, "save_state"), \
                mock.patch.object(editor_ui.messagebox, "askyesno", side_effect=AssertionError("запитання!")), \
                mock.patch.object(editor_ui.messagebox, "showinfo", side_effect=AssertionError("вікно!")):
            self.root.request_sync = mock.Mock()
            self.assertTrue(editor.save(confirm=False, close=False, silent=True))
        persist.assert_called_once()
        self.assertTrue(editor.winfo_exists())                                             # редактор лишився відкритим
        self.assertFalse(editor.has_unsaved())
        self.root.request_sync.assert_called_once()                                        # Classroom перевіряється

    def test_autosave_is_silent_when_data_has_errors(self):
        editor = self.combined()
        editor.cfg["course_map"]["9-Б ГО"]["plan"] = "плану-не-існує"           # справжня помилка в даних
        editor.autosave_enabled = True
        with mock.patch.object(editor_ui.messagebox, "showerror", side_effect=AssertionError("вікно!")), \
                mock.patch.object(editor_ui, "persist") as persist:
            self.assertFalse(editor.save(confirm=False, close=False, silent=True))
        persist.assert_not_called()

    def test_autosave_is_off_in_tests_so_real_data_is_never_touched(self):
        editor = self.combined()
        self.assertFalse(editor.autosave_enabled)


if __name__ == "__main__":
    unittest.main()
