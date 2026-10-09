"""5.6: перегляд вкладення подвійним клацом і заставка з логотипом (лампочка блимає, потім горить рівно)."""
import tempfile
import time
import tkinter as tk
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from PIL import Image

from test_update36 import TkCase
from test_update44 import TempProgram

from classroom_assistant import attachment_preview as preview
from classroom_assistant import gui, materials, splash


def pump(widget, n=15):
    for _ in range(n):
        widget.update()
        time.sleep(0.01)


def double_click(widget, x=40, y=30):
    """Справжній подвійний клац: два швидкі натискання (Tk сам розпізнає <Double-Button-1>)."""
    for stamp in (1000, 1060):
        widget.event_generate("<ButtonPress-1>", x=x, y=y, time=stamp)
        widget.event_generate("<ButtonRelease-1>", x=x, y=y, time=stamp + 20)


def picture(folder, name="Інфографіка.png", size=(1200, 800)):
    path = Path(folder) / name
    Image.new("RGB", size, (190, 120, 60)).save(path)
    return path


class PreviewTests(TempProgram):
    def setUp(self):
        super().setUp()
        self.app.geometry("1400x900")
        self.folder = Path(self.tmp.name)

    def windows(self):
        return [w for w in self.app.winfo_children() if w.winfo_class() == "Toplevel" and w.title().startswith("Перегляд")]

    def test_the_kinds_of_files(self):
        self.assertEqual(preview.kind_of("a.PNG"), "image")
        self.assertEqual(preview.kind_of("a.jpeg"), "image")
        self.assertEqual(preview.kind_of("лекція.docx"), "word")
        self.assertEqual(preview.kind_of("книга.pdf"), "other")
        self.assertEqual(preview.kind_of("без_розширення"), "other")

    def test_a_big_picture_opens_fitted_to_the_window_and_can_be_zoomed(self):
        window = preview.show_preview(self.app, picture(self.folder), "Інфографіка.png")
        pump(self.app, 20)
        view = window.view
        self.assertLess(view.scale, 1.0)                                              # велику картинку вписано у вікно
        self.assertTrue(view.fit_mode)
        self.assertLessEqual(view.size[0] * view.scale, view.canvas.winfo_width() + 2)
        self.assertLessEqual(view.size[1] * view.scale, view.canvas.winfo_height() + 2)
        self.assertEqual(len([i for i in view.canvas.find_all() if view.canvas.type(i) == "image"]), 1)
        before = view.scale
        view.zoom(preview.STEP)
        self.assertAlmostEqual(view.scale, before * preview.STEP)
        self.assertFalse(view.fit_mode)
        view.actual_size()
        self.assertEqual(view.scale, 1.0)
        view.fit()
        self.assertTrue(view.fit_mode)
        self.assertAlmostEqual(view.scale, before, places=2)
        self.assertIn("%", window.zoom_text.get())
        self.assertIn("1200×800", window.zoom_text.get())

    def test_a_small_picture_is_never_blown_up_and_zoom_has_limits(self):
        window = preview.show_preview(self.app, picture(self.folder, "мала.png", (120, 80)))
        pump(self.app, 20)
        self.assertEqual(window.view.scale, 1.0)
        for _ in range(40):
            window.view.zoom(preview.STEP)
        self.assertEqual(window.view.scale, preview.MAX_SCALE)
        for _ in range(80):
            window.view.zoom(1 / preview.STEP)
        self.assertEqual(window.view.scale, preview.MIN_SCALE)

    def test_the_mouse_wheel_zooms_and_a_double_click_toggles_fit_and_actual_size(self):
        window = preview.show_preview(self.app, picture(self.folder))
        pump(self.app, 20)
        view = window.view
        start = view.scale
        view.canvas.event_generate("<Button-4>", x=50, y=50)
        self.assertGreater(view.scale, start)
        view.canvas.event_generate("<Button-5>", x=50, y=50)
        self.assertAlmostEqual(view.scale, start, places=2)
        view.fit()
        double_click(view.canvas)
        self.assertEqual(view.scale, 1.0)
        view.canvas.event_generate("<ButtonPress-1>", x=5, y=5, time=9000)                # пауза: це вже новий клац, не подвійний
        view.canvas.event_generate("<ButtonRelease-1>", x=5, y=5, time=9010)
        double_click(view.canvas)
        self.assertTrue(view.fit_mode)

    def test_a_second_double_click_on_the_same_file_does_not_open_a_second_window(self):
        path = picture(self.folder)
        first = preview.show_preview(self.app, path)
        second = preview.show_preview(self.app, path)
        self.assertIs(first, second)
        self.assertEqual(len(self.windows()), 1)
        other = preview.show_preview(self.app, picture(self.folder, "друга.png"))
        self.assertIsNot(first, other)
        self.assertEqual(len(self.windows()), 2)

    def test_a_word_document_shows_headings_paragraphs_and_tables(self):
        document = Document()
        document.add_heading("Тема уроку", 1)
        document.add_paragraph("Перший абзац лекції.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Дата", "1648"
        path = self.folder / "Лекція.docx"
        document.save(path)
        window = preview.show_preview(self.app, path)
        text = window.text.get("1.0", "end")
        for expected in ("Тема уроку", "Перший абзац лекції.", "Дата", "1648"):
            self.assertIn(expected, text)
        self.assertEqual(str(window.text.cget("state")), "disabled")                   # лише перегляд: змінити не можна

    def test_an_empty_word_document_explains_itself(self):
        path = self.folder / "порожній.docx"
        Document().save(path)
        window = preview.show_preview(self.app, path)
        self.assertIn("немає тексту", window.text.get("1.0", "end"))

    def test_other_files_show_details_and_can_be_opened_in_the_default_program(self):
        path = self.folder / "довідник.pdf"
        path.write_bytes(b"%PDF-1.4 x" * 100)
        window = preview.show_preview(self.app, path)
        self.assertEqual(window.kind, "other")
        labels = " ".join(str(w.cget("text")) for w in self._walk(window) if w.winfo_class() == "TLabel")
        self.assertIn("довідник.pdf", labels)
        self.assertIn("PDF", labels)
        button = next(w for w in self._walk(window) if w.winfo_class() == "TButton"
                      and str(w.cget("text")) == "Відкрити у програмі за замовчуванням")
        with mock.patch.object(materials, "open_file") as opener:
            button.invoke()
        opener.assert_called_once_with(path)

    def test_a_missing_or_broken_file_gives_a_friendly_message_and_no_leftover_window(self):
        with mock.patch.object(preview.messagebox, "showwarning") as warn:
            self.assertIsNone(preview.show_preview(self.app, self.folder / "немає.png"))
        self.assertIn("не знайдено", warn.call_args.args[1])
        broken = self.folder / "зламана.png"
        broken.write_bytes("це не картинка".encode("utf-8"))
        with mock.patch.object(preview.messagebox, "showwarning") as warn:
            self.assertIsNone(preview.show_preview(self.app, broken))
        self.assertIn("Не вдалося показати файл", warn.call_args.args[1])
        self.assertEqual(self.windows(), [])

    def test_double_click_on_an_attachment_card_opens_the_preview_but_the_cross_still_removes_it(self):
        self.app.datevar.set("05.10.2026")
        self.app.update_day()
        self.app.grid.selection_set("0")
        self.app.update()
        lesson = self.app.selected()
        path = picture(self.folder)
        self.app.state.setdefault("attachments", {})[lesson.unique_key] = [{"path": str(path), "name": "Інфографіка.png"}]
        self.app.refresh_attachment_previews()
        self.app.update()
        cards = [w for w in self._walk(self.app.attachment_bar) if w.winfo_class() == "TFrame" and w.winfo_children()
                 and any(c.winfo_class() == "TButton" for c in w.winfo_children())]
        self.assertEqual(len(cards), 1)
        calls = []
        self.app.preview_attachment = lambda entry: calls.append(entry["path"])
        stamp = 100000
        for widget in [cards[0]] + [c for c in cards[0].winfo_children() if c.winfo_class() == "TLabel"]:
            stamp += 5000
            for shift in (0, 60):
                widget.event_generate("<ButtonPress-1>", x=5, y=5, time=stamp + shift)
                widget.event_generate("<ButtonRelease-1>", x=5, y=5, time=stamp + shift + 20)
            self.assertEqual(str(widget.cget("cursor")), "hand2")
        self.assertEqual(calls, [str(path)] * 3)                                        # картка, мініатюра й назва
        cross = next(c for c in cards[0].winfo_children() if c.winfo_class() == "TButton")
        for shift in (0, 60):
            cross.event_generate("<ButtonPress-1>", x=5, y=5, time=stamp + 9000 + shift)
        self.assertEqual(len(calls), 3)                                                 # по хрестику перегляду немає
        self.app.preview_attachment = type(self.app).preview_attachment.__get__(self.app)
        window = self.app.preview_attachment({"path": str(path), "name": "Інфографіка.png"})
        self.assertEqual(window.kind, "image")

    @staticmethod
    def _walk(widget):
        out, stack = [], [widget]
        while stack:
            current = stack.pop()
            stack.extend(current.winfo_children())
            out.append(current)
        return out


class FakeClock:
    def __init__(self):
        self.now = 0.0

    def __call__(self):
        return self.now

    def sleep(self, seconds):
        self.now += seconds


class SplashTests(TkCase):
    def make(self, **kwargs):
        clock = FakeClock()
        shown = splash.Splash(self.root, clock=clock, sleep=clock.sleep, size=300, **kwargs)
        self.addCleanup(lambda: shown.close())
        return shown, clock

    def led_brightness(self, image):
        cx, cy = round(0.852 * image.width), round(0.452 * image.height)
        box = image.crop((cx - 6, cy - 6, cx + 6, cy + 6)).convert("L")
        return sum(box.getdata()) / 144

    def test_the_lamp_goes_from_dark_to_bright_and_the_steady_glow_is_brightest(self):
        frames, steady = splash.render_frames(300)
        levels = [self.led_brightness(image) for _, image in frames]
        self.assertEqual(levels, sorted(levels))                                       # чим вищий рівень, тим яскравіша лампочка
        self.assertLess(levels[0], 60)                                                 # вимкнена — темна
        self.assertGreater(levels[-1], 150)                                            # увімкнена — яскрава
        self.assertGreaterEqual(self.led_brightness(steady), levels[-1])
        self.assertEqual(len({image.tobytes() for _, image in frames}), len(frames))

    def test_while_loading_the_lamp_blinks_a_lot_and_never_shows_the_steady_glow(self):
        shown, clock = self.make()
        seen = []
        for _ in range(125):                                                           # 2 секунди кадрами по 16 мс
            clock.now += 0.016
            shown.pump()
            if not seen or seen[-1] != shown.level:
                seen.append(shown.level)
        self.assertNotIn("steady", seen)
        self.assertGreaterEqual(len(seen), 18)                                          # близько 12 змін на секунду
        self.assertTrue(all(a != b for a, b in zip(seen, seen[1:])))                    # щоразу інша яскравість
        self.assertGreaterEqual(len(set(seen)), 4)                                      # а не два стани

    def test_when_loading_ends_early_the_lamp_burns_steadily_until_the_third_second(self):
        shown, clock = self.make()
        clock.now = 1.2
        shown.finish()
        self.assertGreaterEqual(clock.now, 3.0)
        self.assertLess(clock.now, 3.1)                                                 # рівно близько 3 секунд
        self.assertEqual(shown.level, "steady")
        self.assertEqual(shown.loaded_at, 1.2)
        self.assertTrue(shown.closed)

    def test_slow_loading_still_gets_a_visible_steady_phase_before_the_program_opens(self):
        shown, clock = self.make()
        clock.now = 2.9
        shown.finish()
        self.assertGreaterEqual(clock.now, 2.9 + splash.STEADY_MIN)                     # хоча б 0,6 с рівного світла
        clock2 = FakeClock()
        other = splash.Splash(self.root, clock=clock2, sleep=clock2.sleep, size=300)
        clock2.now = 6.0
        other.finish()
        self.assertGreaterEqual(clock2.now, 6.0 + splash.STEADY_MIN)

    def test_the_main_window_is_hidden_during_the_splash_and_shown_after(self):
        self.root.deiconify()                                                            # у тестах вікно спершу сховане: показуємо як у програмі
        self.root.update()
        shown, clock = self.make()
        hidden = (str(self.root.state()) == "withdrawn") or float(self.root.attributes("-alpha")) == 0.0
        self.assertTrue(hidden)
        clock.now = 1.0
        shown.finish()
        visible = (str(self.root.state()) != "withdrawn") and float(self.root.attributes("-alpha")) == 1.0
        self.assertTrue(visible)

    def test_it_never_runs_the_programs_timers_while_it_is_up(self):
        fired = []
        self.root.after(0, lambda: fired.append(1))
        shown, clock = self.make()
        clock.now = 1.0
        shown.finish()
        self.assertEqual(fired, [])                                                     # нічого не з'являється поверх заставки
        self.root.update()
        self.assertEqual(fired, [1])                                                    # а після неї таймери працюють як завжди

    def test_the_splash_window_is_centred_borderless_and_on_top(self):
        shown, _ = self.make()
        self.assertTrue(shown.win.overrideredirect())
        self.assertEqual(shown.win.attributes("-topmost") in (1, True), True)
        width, height = shown.win.winfo_reqwidth(), shown.win.winfo_reqheight()
        self.assertGreaterEqual(height, width)                                          # логотип і рядок «Завантаження…»

    def test_closing_is_safe_to_repeat(self):
        shown, clock = self.make()
        shown.close()
        shown.close()
        shown.finish()
        shown.pump()
        self.assertTrue(shown.closed)


class SplashInStartupTests(TempProgram):
    def test_the_program_pumps_the_splash_while_loading_and_finishes_it_at_the_very_end(self):
        events = []

        class Recorder:
            def __init__(self, root, **kwargs):
                events.append("created")

            def pump(self):
                events.append("pump")

            def finish(self):
                events.append("finish")

            def close(self):
                events.append("close")
        with mock.patch.object(gui, "Splash", Recorder):
            second = gui.MainApp(show_splash=True)
            self.addCleanup(second.destroy)
        self.assertEqual(events[0], "created")
        self.assertEqual(events[-1], "finish")                                           # лише після повного завантаження
        self.assertEqual(events.count("finish"), 1)
        self.assertGreaterEqual(events.count("pump"), 4)
        self.assertEqual(second._splash, None)

    def test_without_the_flag_or_with_the_opt_out_there_is_no_splash(self):
        with mock.patch.object(gui, "Splash", side_effect=AssertionError("заставки бути не має")):
            self.assertIsNone(self.app._splash)
            with mock.patch.dict("os.environ", {"POMICHNYK_NO_SPLASH": "1"}):
                third = gui.MainApp(show_splash=True)
                self.addCleanup(third.destroy)
            self.assertIsNone(third._splash)

    def test_a_splash_that_fails_never_prevents_the_program_from_starting(self):
        with mock.patch.object(gui, "Splash", side_effect=RuntimeError("нема екрана")):
            fourth = gui.MainApp(show_splash=True)
            self.addCleanup(fourth.destroy)
        self.assertIsNone(fourth._splash)
        self.assertTrue(fourth.winfo_exists())

    def test_launch_asks_for_the_splash(self):
        with mock.patch.object(gui, "MainApp") as app:
            gui.launch()
        app.assert_called_once_with(show_splash=True)
        app.return_value.mainloop.assert_called_once()


if __name__ == "__main__":
    unittest.main()
