"""3.5: навантаження у форматі вчителя, імена Word, зразки, дані, вікна, гарячі клавіші."""
import copy
import tempfile
import tkinter as tk
import unittest
import zipfile
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from docx import Document

from classroom_assistant import editor_ui, hotkeys, material_library, samples_ui
from classroom_assistant.chatgpt_bridge import (best_lesson_for_file, build_prompt,
                                                infographic_name, lesson_code, match_score)
from classroom_assistant.editor_core import (docx_table_rows, extract_lessons, guess_columns,
                                             import_source_dates)
from classroom_assistant.engine import day_lessons, lesson_base_name, read_json
from classroom_assistant.schedule_io import decode_schedule, resolve_stream, table_rows
from classroom_assistant.window_ui import fit_work_window
from classroom_assistant.workload_io import (export_workload_docx, export_workload_jpeg,
                                             pair_text, parse_workload_title, title_parts)


def sample_config():
    config = samples_ui.sample_workload_config()
    streams = {x for pairs in config["days"].values() for pair in pairs for x in pair if x}
    config["course_map"] = {name: {"plan": "x", "course_title": name} for name in streams}
    return config


class WorkloadFormatTests(unittest.TestCase):
    def test_pair_text_matches_teacher_notation(self):
        self.assertEqual(pair_text("8-Б ІУ", "8-Б ГО"), "8-Б ІУ / ГО")
        self.assertEqual(pair_text(None, "9-Г Право"), "/ 9-Г Право")
        self.assertEqual(pair_text("11 ІУ стандарт", None), "11 ІУ стандарт /")
        self.assertEqual(pair_text("8-Б ВІ", "8-Б ВІ"), "8-Б ВІ")
        self.assertEqual(pair_text(None, None), "")
        # різні потоки одного класу без літери не скорочуються
        self.assertEqual(pair_text("11 ІУ профіль", "11 ІУ стандарт"), "11 ІУ профіль / 11 ІУ стандарт")

    def test_title_is_like_the_sample(self):
        text = "".join(part[0] for part in title_parts(sample_config()))
        self.assertTrue(text.startswith("НАВАНТАЖЕННЯ (Прізвище І.О.) //  канікули  26.10-01.11, 24.12-10.01, 22.03-28.03 // 000000"))

    def test_word_roundtrip_through_the_importer(self):
        config = sample_config()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "n.docx"
            export_workload_docx(config, path)
            rows = table_rows(path)
            schedule, unknown, found = decode_schedule(rows, config)
            self.assertEqual(unknown, [])
            self.assertEqual(found, list(range(1, 9)))
            for day in config["days"]:
                self.assertEqual(schedule[day], [list(x) for x in config["days"][day]])
            self.assertEqual(parse_workload_title(rows),
                             {"workload_teacher": "Прізвище І.О.", "workload_code": "000000"})
            table = Document(path).tables[0]
            self.assertEqual(len(table.columns), 6)
            red_meal = [r for r in table.rows if "ХАРЧУВАННЯ" in r.cells[0].text]
            self.assertEqual(len(red_meal), 1)
            self.assertEqual(str(red_meal[0].cells[0].paragraphs[0].runs[0].font.color.rgb), "FF0000")

    def test_picture_is_created(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "n.jpg"
            export_workload_jpeg(sample_config(), path)
            with Image.open(path) as picture:
                self.assertEqual(picture.format, "JPEG")
                self.assertGreater(picture.size[0], 1300)

    def test_class_without_letter_variant_is_found_by_unique_prefix(self):
        streams = ["5-Г історія", "8-Б ІУ", "8-Б ГО"]
        self.assertEqual(resolve_stream("5-Г ІУ", streams), "5-Г історія")
        self.assertTrue(resolve_stream("8-Б ВІ", streams).startswith("? "))


class NamingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lesson = day_lessons("2026-10-01", read_json("Налаштування.json"))[0]

    def test_name_says_everything(self):
        name = lesson_base_name(self.lesson)
        self.assertEqual(name, f"{self.lesson.stream}, Урок 01.10 — {self.lesson.topic}".rstrip(" ."))
        self.assertTrue(name.startswith("11 ІУ профіль, Урок 01.10 — "))             # клас і дата — на початку
        self.assertNotIn(":", name)
        self.assertNotIn("#", name)

    def test_long_and_dirty_topics_are_safe(self):
        dirty = replace(self.lesson, topic='Тема: «Що/де?» ' * 20)
        name = lesson_base_name(dirty)
        self.assertLessEqual(len(name), 120)
        self.assertFalse(set('<>:"/\\|?*') & set(name))

    def test_prompt_asks_for_files_and_no_code_in_word(self):
        prompt = build_prompt(self.lesson)
        self.assertIn(lesson_base_name(self.lesson) + ".docx", prompt)
        self.assertIn(infographic_name(self.lesson) + ".png", prompt)
        self.assertIn("У документі НЕ пиши слів «Код уроку»", prompt)
        self.assertIn("Times New Roman", prompt)
        self.assertIn(f"КОД УРОКУ: {lesson_code(self.lesson)}", prompt)

    def test_files_are_matched_to_lessons(self):
        lessons = day_lessons("2026-10-01", read_json("Налаштування.json"))[:3]
        for lesson in lessons:
            name = lesson_base_name(lesson) + ".docx"
            self.assertGreaterEqual(match_score(lesson, name), 5)
            found = best_lesson_for_file(lessons, infographic_name(lesson) + ".png")
            if found is not None:
                self.assertEqual(found.topic, lesson.topic)
        self.assertIsNone(best_lesson_for_file(lessons, "щось.docx"))

    def test_parallel_copy_fixes_date_in_strip_header_and_homework(self):
        lesson = replace(day_lessons("2026-10-01", read_json("Налаштування.json"))[0],
                         day="2026-10-09", lesson_number=9, homework="параграф 9")
        with tempfile.TemporaryDirectory() as folder:
            master = Path(folder) / "master.docx"
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "ІСТОРІЯ УКРАЇНИ • 8 клас • 02.10.2026"
            doc.add_paragraph("Урок № 7 • 02.10.2026")
            for number in range(8):
                doc.add_paragraph(f"Розділ {number}. " + "Змістовний абзац лекції про епоху. " * 4)
            doc.save(master)
            target = Path(folder) / "out" / "copy.docx"
            with mock.patch.object(material_library, "word_path", lambda _l: target):
                result = material_library.copy_for_lesson(master, lesson)
            copy_doc = Document(result)
            paragraphs = [p.text for p in copy_doc.paragraphs if p.text.strip()]
            self.assertEqual(paragraphs[0], "Урок № 9 • 09.10.2026")
            self.assertFalse(any(x.startswith("Д/з") for x in paragraphs))          # домашнього завдання в Word немає
            self.assertIn("09.10.2026", copy_doc.sections[0].header.paragraphs[0].text)
            self.assertNotIn("02.10.2026", copy_doc.sections[0].header.paragraphs[0].text)


class SamplesAndDataTests(unittest.TestCase):
    def test_sample_ktp_is_read_by_the_program(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "ktp.docx"
            samples_ui.build_sample_ktp(path)
            rows = docx_table_rows(path)
            header = next(r for r in rows[:8] if any("тема" in c.casefold() for c in r))
            topic, homework, number = guess_columns(header)
            lessons = import_source_dates(rows, extract_lessons(rows, topic, homework, number),
                                          topic, number, 2026)
        self.assertEqual(len(lessons), 5)
        self.assertEqual(lessons[3]["source_number"], "4-5")
        self.assertEqual(lessons[0]["source_dates"], {"8-Б": "2026-09-04", "8-В": "2026-09-04",
                                                      "8-Г": "2026-09-03"})
        self.assertEqual(lessons[4]["homework"], "")

    def test_prompts_describe_the_formats(self):
        self.assertIn("«8-Б ІУ / ГО»", samples_ui.WORKLOAD_PROMPT)
        self.assertIn("ХАРЧУВАННЯ У ЇДАЛЬНІ", samples_ui.WORKLOAD_PROMPT)
        self.assertIn("«№ з/п» | «Дата» | «Тема уроку» | «Домашнє завдання»", samples_ui.KTP_PROMPT)

    def test_data_export_skips_secrets_and_code(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "root"
            for relative, text in (("data/Налаштування.json", "{}"), ("data/google_token.json", "SECRET"),
                                   ("data/google_credentials.json", "SECRET"),
                                   ("Готові Word/2026-10-01/a.docx", "x"),
                                   ("classroom_assistant/a.py", "code"), ("tests/t.py", "code"),
                                   ("Вкладення Classroom/img.png", "x")):
                target = root / relative
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(text, encoding="utf-8")
            archive = Path(folder) / "all.zip"
            count = samples_ui.export_all_data(archive, root)
            with zipfile.ZipFile(archive) as z:
                names = set(z.namelist())
        self.assertEqual(count, 3)
        self.assertIn("data/Налаштування.json", names)
        self.assertIn("Готові Word/2026-10-01/a.docx", names)
        self.assertFalse(any("token" in n or "credentials" in n or n.endswith(".py") for n in names))


class WindowAndKeysTests(unittest.TestCase):
    class FakeWindow:
        def __init__(self, screen, request):
            self.screen, self.request = screen, request
            self.geometry_value = self.minimum = None
        def after_idle(self, function): function()
        def winfo_exists(self): return True
        def update_idletasks(self): pass
        def winfo_screenwidth(self): return self.screen[0]
        def winfo_screenheight(self): return self.screen[1]
        def winfo_reqwidth(self): return self.request[0]
        def winfo_reqheight(self): return self.request[1]
        def minsize(self, w, h): self.minimum = (w, h)
        def geometry(self, value): self.geometry_value = value

    def size(self, screen, request, kind="normal"):
        window = self.FakeWindow(screen, request)
        fit_work_window(window, kind)
        w, h = window.geometry_value.split("+")[0].split("x")
        return int(w), int(h), window.minimum

    def test_window_is_not_full_screen_but_not_cropped(self):
        w, h, minimum = self.size((1920, 1080), (900, 600))
        self.assertLess(w, 1920 * 0.8)
        self.assertLess(h, 1080 * 0.85)
        self.assertGreaterEqual(w, minimum[0])
        w, h, _ = self.size((1920, 1080), (1500, 900))        # вміст великий — вікно зростає
        self.assertGreaterEqual(w, 1500)
        self.assertGreaterEqual(h, 900)

    def test_small_laptop_screen_fits(self):
        w, h, _ = self.size((1366, 768), (2000, 1200))
        self.assertLessEqual(w, 1366)
        self.assertLessEqual(h, 768)

    def test_cyrillic_layout_keys_are_recognised(self):
        event = SimpleNamespace(keysym="Cyrillic_em", keycode=0)
        self.assertEqual(hotkeys.physical_letter(event), ("v", False))
        latin = SimpleNamespace(keysym="v", keycode=0)
        self.assertEqual(hotkeys.physical_letter(latin), ("v", True))
        with mock.patch.object(hotkeys.sys, "platform", "win32"):
            ukrainian = SimpleNamespace(keysym="Cyrillic_em", keycode=86)
            self.assertEqual(hotkeys.physical_letter(ukrainian), ("v", False))


class EditorClassListTests(unittest.TestCase):
    def setUp(self):
        try:
            try:
                from tkinterdnd2 import TkinterDnD
                self.root = TkinterDnD.Tk()
            except ImportError:
                self.root = tk.Tk()
        except tk.TclError:
            self.skipTest("немає дисплея")
        self.root.withdraw()
        # Жодних модальних вікон під час тестів: вони зависли б без людини.
        for name, value in (("showinfo", None), ("showwarning", None), ("showerror", None),
                            ("askyesno", True), ("askyesnocancel", False)):
            patcher = mock.patch.object(editor_ui.messagebox, name, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.root.state = {}
        self.root.google_courses = []
        self.root.update_day = lambda: None
        self.editor = None

    def tearDown(self):
        if self.editor is not None:
            try:
                self.editor.destroy()
            except tk.TclError:
                pass
        self.root.destroy()

    def make(self):
        self.editor = editor_ui.SchoolEditor(self.root, lambda: None)
        return self.editor

    def test_list_follows_classroom_order(self):
        names = ["9-Б ІУ", "5-Г історія", "8-Б ІУ", "5-Б"]
        self.root.google_courses = [{"name": n, "id": str(i)} for i, n in enumerate(names)]
        editor = self.make()
        labels = [e["label"] for e in editor.list_entries]
        self.assertEqual(labels[:3], names[:3])
        self.assertEqual(editor.list_entries[3]["course"], "5-Б")
        self.assertIsNone(editor.list_entries[3]["plan"])
        self.assertTrue(labels[3].endswith("без КТП"))

    def test_saved_order_is_used_before_sync(self):
        self.root.state = {"classroom_order": ["10-Б ГО", "8-Б ІУ"]}
        editor = self.make()
        self.assertEqual([e["course"] for e in editor.list_entries[:2]], ["10-Б ГО", "8-Б ІУ"])

    def test_dropped_file_is_imported_without_asking_for_it_again(self):
        """Файл 8 ІУ, кинутий на порожній курс «5-Б», автоматично йде до класів 8 ІУ (без вікон і запитань)."""
        self.root.google_courses = [{"name": "5-Б", "id": "1"}]
        editor = self.make()
        editor.planlist.selection_clear(0, "end")
        editor.planlist.selection_set(0)
        editor.plan_selected()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "ktp.docx"
            samples_ui.build_sample_ktp(path)
            with mock.patch.object(editor_ui.filedialog, "askopenfilename",
                                   side_effect=AssertionError("вікно вибору файлу не має відкриватись")), \
                    mock.patch.object(editor_ui.messagebox, "askyesno",
                                      side_effect=AssertionError("запитань не має бути")):
                editor.import_plan(path=path)
        self.assertNotIn("5-Б", editor.cfg["course_map"])                 # зайвий тимчасовий потік прибрано
        for stream in ("8-Б ІУ", "8-В ІУ", "8-Г ІУ"):
            plan = editor.cfg["course_map"][stream]["plan"]
            self.assertEqual(len(editor.plans[plan]["lessons"]), 5, stream)


    def test_close_asks_and_can_save(self):
        editor = self.make()
        editor.teacher_var.set("Новий І.О.")
        with mock.patch.object(editor_ui.messagebox, "askyesnocancel", return_value=None):
            editor.close()
        self.assertTrue(editor.winfo_exists())
        with mock.patch.object(editor_ui.messagebox, "askyesnocancel", return_value=True), \
                mock.patch.object(editor, "save", return_value=True) as save:
            editor.close()
        save.assert_called_once_with(confirm=False)

    def test_close_without_changes_does_not_ask(self):
        editor = self.make()
        with mock.patch.object(editor_ui.messagebox, "askyesnocancel",
                               side_effect=AssertionError("не повинно питати")):
            editor.close()


if __name__ == "__main__":
    unittest.main()
