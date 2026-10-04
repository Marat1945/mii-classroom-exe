"""4.4: програма не вважає «створеною» чернетку, якої вже немає в Classroom (звірка після синхронізації)."""
import shutil
import tempfile
import time
import tkinter as tk
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest import mock

from test_update36 import TkCase

from docx import Document

from test_update42 import make_lecture, make_picture
from classroom_assistant import (attachments, chatgpt_ui, draft_batch, engine, google_client, gui,
                                 material_library)
from classroom_assistant.engine import lesson_base_name
from classroom_assistant.engine import build_calendar

REPO_DATA = Path(__file__).resolve().parents[1] / "data"


class TempProgram(TkCase):
    """Головне вікно на справжніх даних у тимчасовій папці (нічого в репозиторії не змінюється)."""
    NEEDS_ROOT = False

    def setUp(self):
        super().setUp()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        root = Path(self.tmp.name) / "root"
        (root / "data").mkdir(parents=True)
        for name in ("Налаштування.json", "Календарні плани.json"):
            shutil.copy(REPO_DATA / name, root / "data" / name)
        stack = ExitStack()
        for patcher in (mock.patch.object(engine, "DATA", root / "data"), mock.patch.object(engine, "ROOT", root),
                        mock.patch.object(gui, "DATA", root / "data"), mock.patch.object(gui, "ROOT", root),
                        mock.patch.object(material_library, "LIBRARY_ROOT", root / "Бібліотека"),
                        mock.patch.object(material_library, "INDEX", root / "Бібліотека" / "Індекс.json"),
                        mock.patch.object(attachments, "ATTACHMENT_ROOT", root / "Вкладення")):
            stack.enter_context(patcher)
        self.addCleanup(stack.close)
        self.toasts = []
        stack.enter_context(mock.patch.object(gui, "show_toast", side_effect=lambda p, text, *a, **k: self.toasts.append(text)))
        self.app = gui.MainApp()

        def close():
            try:
                self.app.destroy()
            except tk.TclError:
                pass
        self.addCleanup(close)
        day = [x for x in build_calendar(self.app.cfg) if x.day == "2026-10-05" and x.status == "готово"]
        self.lessons = day[:3]
        self.course_ids = {x.course_title: str(1000 + i) for i, x in enumerate(self.lessons)}
        self.app.state["course_ids"] = dict(self.course_ids)
        self.app._sync_started = time.time()

    def result(self, entries=None, truncated=None, only=None):
        ids = self.course_ids if only is None else {k: v for k, v in self.course_ids.items() if k in only}
        return {"courses": [{"name": t, "id": i} for t, i in ids.items()], "mapped": dict(ids),
                "entries": {i: [] for i in ids.values()} if entries is None else entries,
                "errors": [], "unmapped": [], "truncated": truncated or {}}

    def draft(self, lesson, **extra):
        record = {"id": f"d-{lesson.stream}", "created_at": 1.0}
        record.update(extra)
        self.app.state.setdefault("drafts", {})[lesson.unique_key] = record
        return record

    def statuses(self):
        self.app.datevar.set("05.10.2026")
        self.app.update_day()
        return {self.app.rows[int(i)].stream: self.app.grid.item(i, "values")[-1] for i in self.app.grid.get_children()}



class ReconcileTests(TempProgram):
    def test_a_draft_deleted_in_classroom_stops_being_reported_as_created(self):
        a, b, c = self.lessons
        for lesson in (a, b, c):
            self.draft(lesson)
        self.app._sync_done(self.result(), False)                              # Classroom прочитано: чернеток там немає
        self.assertEqual(self.app.state["drafts"], {})                          # усе, що видалили, прибрано
        self.assertEqual(sorted(x["id"] for x in self.app.state["drafts_removed"]),
                         sorted(f"d-{x.stream}" for x in (a, b, c)))
        self.assertIn("прибрано зі стану: 3", self.toasts[-1])
        for stream, status in self.statuses().items():
            self.assertFalse(status.startswith("Створено в Google"), (stream, status))
            self.assertNotIn("не знайдено", status)

    def test_after_the_cleanup_the_lesson_can_get_a_new_draft(self):
        lesson = self.lessons[0]
        self.draft(lesson)
        row = lesson
        _, before = draft_batch.plan_day_drafts([row], self.app.state, lambda x: "текст")
        self.assertEqual([reason for _, reason in before], ["чернетку вже створено"])        # саме це й блокувало
        self.app._sync_done(self.result(), False)
        ready, skipped = draft_batch.plan_day_drafts([row], self.app.state, lambda x: "текст")
        self.assertEqual(len(ready), 1)
        self.assertEqual(skipped, [])

    def test_only_the_missing_ones_are_removed(self):
        a, b, c = self.lessons
        for lesson in (a, b, c):
            self.draft(lesson)
        still_there = {self.course_ids[b.course_title]: [{"id": f"d-{b.stream}", "state": "DRAFT", "type": "Матеріал",
                                                          "title": "x", "updated": ""}]}
        entries = {i: [] for i in self.course_ids.values()}
        entries.update(still_there)
        self.app._sync_done(self.result(entries=entries), False)
        self.assertEqual(list(self.app.state["drafts"]), [b.unique_key])                   # у Classroom лишилась — лишилась й тут
        self.assertIn("чернетка в Classroom", self.statuses()[b.stream])

    def test_nothing_is_removed_when_it_cannot_be_verified(self):
        a, b, c = self.lessons
        self.draft(a, created_at=time.time())                                              # свіжа: Classroom ще не встиг
        self.draft(b)                                                                       # курс позначено неповним
        record = self.draft(c)
        record.pop("id")                                                                    # без id перевірити неможливо
        truncated = {self.course_ids[b.course_title]: ["Матеріал"]}
        self.app._sync_done(self.result(truncated=truncated), False)
        self.assertEqual(set(self.app.state["drafts"]), {a.unique_key, b.unique_key, c.unique_key})
        self.assertNotIn("drafts_removed", self.app.state)
        self.assertIn("перевіряю Classroom", self.statuses()[a.stream])                    # свіжу не лякаємо «не знайдено»

    def test_a_course_that_failed_to_load_keeps_its_records(self):
        a, b, _ = self.lessons
        self.draft(a)
        self.draft(b)
        result = self.result(only={a.course_title})                                        # курс b у цій синхронізації не читався
        result["entries"] = {self.course_ids[a.course_title]: []}
        result["errors"] = [f"{b.course_title}: помилка мережі"]
        self.app._sync_done(result, False)
        self.assertEqual(list(self.app.state["drafts"]), [b.unique_key])

    def test_the_cleanup_is_recorded_and_not_undone_by_the_back_arrow(self):
        lesson = self.lessons[0]
        self.draft(lesson)
        self.app._sync_done(self.result(), False)
        self.assertIn("drafts_removed", gui.NON_UNDOABLE)
        self.assertNotIn("drafts_removed", self.app._history_snapshot())
        self.assertEqual(self.app.state["drafts_removed"][0]["reason"], "немає в Classroom")

    def test_returning_to_the_program_refreshes_classroom_but_not_more_than_once_a_minute(self):
        calls = []
        self.app.request_sync = lambda delay=900: calls.append(delay)
        with mock.patch.object(google_client, "token_ready", return_value=True):
            self.app._last_sync_ts = 0
            self.app._on_focus_in(type("E", (), {"widget": self.app})())
            self.assertEqual(len(calls), 1)
            self.app._last_sync_ts = time.time()
            self.app._on_focus_in(type("E", (), {"widget": self.app})())               # щойно синхронізували
            self.assertEqual(len(calls), 1)
            self.app._last_sync_ts = 0
            self.app._on_focus_in(type("E", (), {"widget": self.app.grid})())          # фокус дрібного елемента — ні
            self.assertEqual(len(calls), 1)
        with mock.patch.object(google_client, "token_ready", return_value=False):
            self.app._on_focus_in(type("E", (), {"widget": self.app})())               # Google не підключено — тихо
            self.assertEqual(len(calls), 1)


class SyncReportsTruncationTests(unittest.TestCase):
    def test_sync_everything_returns_which_courses_were_cut_short(self):
        posts = {"1": {"items": [{"id": "x"}], "truncated": ["Матеріал"]}, "2": {"items": [], "truncated": []}}
        with mock.patch.object(google_client, "list_teacher_courses",
                               return_value=[{"id": "1", "name": "8-Б ІУ"}, {"id": "2", "name": "8-В ІУ"}]), \
                mock.patch.object(google_client, "services", return_value=(object(), object())), \
                mock.patch.object(google_client, "list_classroom_posts", side_effect=lambda cid, service=None: posts[cid]):
            result = google_client.sync_everything(["8-Б ІУ", "8-В ІУ"], {})
        self.assertEqual(result["truncated"], {"1": ["Матеріал"]})
        self.assertEqual(set(result["entries"]), {"1", "2"})


if __name__ == "__main__":
    unittest.main()


class GoogleButtonTests(TempProgram):
    """Кнопка «Підключити Google» завжди відповідає: коротка підказка без «ОК»."""

    def test_label_shows_whether_google_is_connected(self):
        with mock.patch.object(google_client, "token_ready", return_value=True):
            self.app._refresh_google_button()
            self.assertEqual(str(self.app.google_button.cget("text")), "✓ Google підключено")
        with mock.patch.object(google_client, "token_ready", return_value=False):
            self.app._refresh_google_button()
            self.assertEqual(str(self.app.google_button.cget("text")), "Підключити Google")

    def test_header_is_laid_out_again_when_the_label_changes_width(self):
        row = self.app.google_button.master
        calls = []
        original = row._layout
        row._layout = lambda *a: (calls.append(1), original(*a))[1]
        with mock.patch.object(google_client, "token_ready", return_value=True):
            self.app._refresh_google_button()
            self.app.update()
        self.assertTrue(calls)

    def test_click_when_connected_answers_at_once_and_after_the_update(self):
        started = []
        self.app.sync_classroom = lambda **kw: started.append(kw)
        with mock.patch.object(google_client, "token_ready", return_value=True), \
                mock.patch.object(gui.messagebox, "showinfo", side_effect=AssertionError("вікно з «ОК»!")):
            self.app.connect_google()
            self.assertIn("Google уже підключено", self.toasts[-1])               # відповідь одразу
            self.assertEqual(started, [{"interactive": True, "announce": True}])
            self.app._sync_done(self.result(), True, True)                        # оновлення завершилось
        self.assertIn("✓ Classroom оновлено: курсів — 3, записів — 0", self.toasts[-1])

    def test_click_during_a_running_update_says_so_instead_of_silence(self):
        started = []
        self.app.sync_classroom = lambda **kw: started.append(kw)
        self.app._sync_running = True
        with mock.patch.object(google_client, "token_ready", return_value=True):
            self.app.connect_google()
        self.assertIn("уже триває", self.toasts[-1])
        self.assertEqual(started, [])

    def test_real_problems_still_need_reading_but_missing_courses_are_only_a_warning_toast(self):
        result = self.result()
        result["errors"] = ["8-Б ІУ: мережа"]
        with mock.patch.object(gui.messagebox, "showinfo") as modal:
            self.app._sync_done(result, True, True)
        modal.assert_called_once()                                                # помилка читання — вікно
        result = self.result()
        result["unmapped"] = ["10-Б ГО"]
        with mock.patch.object(gui.messagebox, "showinfo", side_effect=AssertionError("вікно!")):
            self.app._sync_done(result, True, True)
        self.assertIn("не зіставлено курсів: 1", self.toasts[-1])

    def test_not_connected_opens_the_step_by_step_wizard(self):
        with mock.patch.object(google_client, "token_ready", return_value=False), \
                mock.patch("classroom_assistant.google_setup_ui.show_google_wizard") as wizard:
            self.app.connect_google()
        wizard.assert_called_once_with(self.app)


class DayRegenerationTests(TempProgram):
    """Кнопка «Лекції GPT на весь день» виконує завдання для ВСІХ уроків дня, навіть якщо Word уже є."""

    def day_rows(self, text="06.10.2026"):
        self.app.datevar.set(text)
        self.app.update_day()
        return list(self.app.rows)

    def give_everyone_a_word(self, rows):
        folder = Path(self.tmp.name) / "наявні"
        folder.mkdir(exist_ok=True)
        for row in rows:
            word = make_lecture(folder / f"{row.stream} {row.period}.docx", f"{row.stream}, Урок")
            self.app.state.setdefault("files", {})[row.unique_key] = {
                "path": str(word), "validated": True, "complete": True, "attached_at": 1.0}

    def groups(self, rows):
        seen, result = set(), []
        for row in rows:
            key = (material_library.signature_key(row, self.app.cfg), row.lesson_number, row.topic.casefold().strip())
            if key not in seen:
                seen.add(key)
                result.append(row)
        return result

    def test_every_lesson_of_the_day_is_prepared_even_when_all_already_have_word(self):
        rows = self.day_rows()
        self.assertGreaterEqual(len(rows), 5)
        self.give_everyone_a_word(rows)
        opened = []
        with mock.patch("classroom_assistant.chatgpt_ui.open_day_dialog", side_effect=lambda app, lessons: opened.append(lessons)), \
                mock.patch.object(gui.messagebox, "askyesno", side_effect=AssertionError("зайве запитання!")):
            self.app.chatgpt_batch()
        self.assertEqual(len(opened), 1)                                         # вікно відкрилось, а не «усе вже є»
        expected = self.groups(rows)
        self.assertEqual([x.unique_key for x in opened[0]], [x.unique_key for x in expected])
        ktp = [x for x in opened[0] if x.stream.endswith("ВІ") and x.stream.startswith("8-")]
        self.assertEqual(len(ktp), 1)                                            # 8-Б, 8-В, 8-Г ВІ — одна задача
        self.assertFalse(any("вже мають Word" in text for text in self.toasts))

    def test_the_dialog_tells_old_word_from_new_and_new_files_replace_old_ones(self):
        rows = self.day_rows()
        self.give_everyone_a_word(rows)
        lessons = self.groups(rows)[:3]
        dialog = chatgpt_ui.ChatGPTDayDialog(self.app, lessons)
        self.addCleanup(lambda: dialog.winfo_exists() and dialog.destroy())
        self.assertEqual([str(dialog.table.item(str(i), "values")[5]) for i in range(3)], ["є", "є", "є"])
        self.assertIn("«є» — Word уже був", dialog.note.cget("text"))
        self.assertIn("Word — 0", dialog.note.cget("text"))                        # старі не видають себе за нові
        self.assertEqual(dialog.checked, {0, 1, 2})                                # усі позначені за замовчуванням
        old = self.app.state["files"][lessons[0].unique_key]["path"]
        label = chatgpt_ui._label_for(self.app, lessons[0])
        name = lesson_base_name(lessons[0], label=label)
        new_word = make_lecture(Path(self.tmp.name) / (name + ".docx"), f"{label}, Урок")
        document = Document(new_word)
        document.add_paragraph("Новий розділ. " + "Додаток до лекції. " * 30)       # інший вміст → інший хеш
        document.save(new_word)
        picture = make_picture(Path(self.tmp.name) / (name + ".png"))
        dialog.docx[0], dialog.images[0] = new_word, [picture]
        dialog.refresh()
        self.assertEqual(str(dialog.table.item("0", "values")[5]), "✅")
        self.assertIn("Word — 1 · зображень — 1", dialog.note.cget("text"))
        dialog.checked = {0}
        dialog.attach_all()
        entry = self.app.state["files"][lessons[0].unique_key]
        self.assertGreater(entry["attached_at"], 1.0)                             # замінено новим
        texts = [p.text for p in Document(entry["path"]).paragraphs if p.text.strip()]
        self.assertTrue(any(x.startswith("Новий розділ.") for x in texts))

    def test_paid_generation_of_explicitly_chosen_lessons_ignores_existing_word_but_default_saves_money(self):
        rows = self.day_rows()
        self.give_everyone_a_word(rows)
        asked, jobs, infos = [], [], []
        self.app.worker = lambda task, done: jobs.append(task)
        with mock.patch.object(gui.messagebox, "askyesno", side_effect=lambda t, m, **k: asked.append(m) or True), \
                mock.patch.object(gui.messagebox, "showinfo", side_effect=lambda t, m, **k: infos.append(m)):
            self.app.batch_ai(only=rows)                                           # без force — економія
            self.assertEqual(jobs, [])
            self.assertTrue(any("вже мають Word" in m for m in infos))
            self.app.batch_ai(only=rows, force=True)                               # явний вибір у вікні
        self.assertEqual(len(jobs), 1)
        self.assertIn("ПЛАТНИХ AI-запитів", asked[-1])
        self.assertIn("буде замінено", asked[-1])


class ClassroomTypeTests(TempProgram):
    """Лекція → «Матеріал»; практична/лабораторна/контрольна/проєктна робота, оцінювання → «Завдання»."""

    TASKS = ["Практична робота. Історичне значення Французької революції.", "Практичне завдання", "Практичне заняття",
             "Навчальний проект. Стадії боротьби за незалежність: від УНР до УПА.", "Навчальний проєкт", "Проєктна робота",
             "Лабораторна робота", "Контрольна робота", "Урок контролю", "Тематичне оцінювання", "Самостійна робота",
             "Перевірна робота", "Діагностична робота", "Підсумкова робота", "Тематична атестація"]
    LECTURES = ["Суспільство. Що об’єднує людей у суспільство.", "Урок узагальнення з теми", "Протести проти політики уряду",
                "Встановлення контролю над територією", "Проект Конституції Пилипа Орлика", "Церковні унії в Україні",
                "Контекст розвитку культури", "Протестантизм і Реформація"]

    def test_which_topics_count_as_practical_or_control(self):
        for topic in self.TASKS:
            self.assertTrue(engine.is_task_lesson(topic), topic)
        for topic in self.LECTURES:
            self.assertFalse(engine.is_task_lesson(topic), topic)

    def test_the_prompt_for_a_project_or_lab_asks_for_analytical_tasks_too(self):
        from dataclasses import replace
        from classroom_assistant.chatgpt_bridge import build_prompt
        base = self.lessons[0]
        for topic in ("Навчальний проект. Стадії боротьби", "Лабораторна робота"):
            text = build_prompt(replace(base, topic=topic))
            self.assertIn("НЕ БІЛЬШЕ ніж 3 аналітичними завданнями", text, topic)
            self.assertIn("(учні здають відповідь)", text)

    def run_single_draft(self, topic):
        from dataclasses import replace
        lesson = replace(self.lessons[0], topic=topic)
        self.app.rows = [lesson]
        calls = []
        self.app.state["course_ids"][lesson.course_title] = "555"
        self.app.state.pop("drafts", None)
        self.app.selected = lambda: lesson
        self.app.effective_lesson = lambda row: row
        self.app.worker = lambda task, done: done(task())
        self.app.update_day = lambda: None
        self.app.desc.delete("1.0", "end")
        self.app.desc.insert("1.0", "Текст для учнів")
        asked = []
        record = {"id": "x1", "kind": "?", "state": "DRAFT", "created_at": time.time()}
        with mock.patch.object(gui.messagebox, "askyesno", side_effect=lambda t, m, **k: asked.append(m) or True), \
                mock.patch.object(google_client, "create_draft", side_effect=lambda *a, **k: calls.append(a) or dict(record)):
            self.app.draft()
        return calls, asked

    def test_single_draft_type_follows_the_topic(self):
        calls, asked = self.run_single_draft("Практична робота. Історичне значення Французької революції.")
        self.assertTrue(calls[0][4])                                              # assignment=True
        self.assertIn("Тип: ЗАВДАННЯ (учні зможуть здати відповідь)", asked[0])
        self.assertIn("Завдання", self.toasts[-1])
        calls, asked = self.run_single_draft("Суспільство. Що об’єднує людей у суспільство.")
        self.assertFalse(calls[0][4])                                             # матеріал
        self.assertIn("Тип: Матеріал", asked[0])

    def test_the_checkbox_forces_everything_to_be_an_assignment(self):
        self.app.assignment.set(True)
        calls, _ = self.run_single_draft("Суспільство. Що об’єднує людей у суспільство.")
        self.assertTrue(calls[0][4])

    def test_batch_creates_lectures_as_materials_and_practical_works_as_assignments(self):
        rows = self.app.rows if self.app.rows else []
        self.app.datevar.set("05.10.2026")
        self.app.update_day()
        rows = list(self.app.rows)
        for row in rows:
            self.app.state["course_ids"][row.course_title] = str(abs(hash(row.course_title)) % 10 ** 6)
        calls = []

        def fake_create(course_id, title, description, docx, assignment, attachments):
            calls.append((title, assignment))
            return {"id": f"d{len(calls)}", "kind": "ASSIGNMENT" if assignment else "MATERIAL", "state": "DRAFT"}
        with mock.patch.object(google_client, "create_draft", side_effect=fake_create), \
                mock.patch.object(gui.messagebox, "askyesno", return_value=True), \
                mock.patch.object(gui.messagebox, "showinfo"), mock.patch.object(gui.messagebox, "showwarning"):
            self.app.request_sync = lambda *a, **k: None
            self.app.batch_drafts()
            deadline = time.time() + 8
            while self.app._batch_drafts_running and time.time() < deadline:
                self.app.update()
                time.sleep(0.02)
        self.assertEqual(len(calls), len(rows))
        for (title, assignment), row in zip(calls, rows):
            self.assertEqual(assignment, engine.is_task_lesson(row.topic), row.topic)
        self.assertTrue(any(a for _, a in calls) and any(not a for _, a in calls))      # у цей день є і те, і те
        records = self.app.state["drafts"]
        self.assertEqual(len(records), len(rows))
        self.assertTrue(all(r.get("created_at") for r in records.values()))            # кожна чернетка має час створення

    def test_a_just_created_draft_is_not_taken_for_a_deleted_one_by_the_next_sync(self):
        lesson = self.lessons[0]
        self.app.state.setdefault("drafts", {})[lesson.unique_key] = {"id": "new1", "created_at": time.time()}
        self.app._sync_started = time.time()
        self.app._sync_done(self.result(), False)                                  # Classroom ще не віддає свіжу чернетку
        self.assertIn(lesson.unique_key, self.app.state["drafts"])

    def test_create_draft_uses_the_right_endpoint_and_stamps_the_time(self):
        classroom, drive = mock.MagicMock(), mock.MagicMock()
        classroom.courses().courseWork().create().execute.return_value = {"id": "w1"}
        classroom.courses().courseWorkMaterials().create().execute.return_value = {"id": "m1"}
        with mock.patch.object(google_client, "services", return_value=(classroom, drive)):
            task = google_client.create_draft("1", "Урок", "опис", None, True, None)
            lecture = google_client.create_draft("1", "Урок", "опис", None, False, None)
        self.assertEqual((task["kind"], task["id"]), ("ASSIGNMENT", "w1"))
        self.assertEqual((lecture["kind"], lecture["id"]), ("MATERIAL", "m1"))
        self.assertTrue(task["created_at"] and lecture["created_at"])
        body = classroom.courses().courseWork().create.call_args.kwargs["body"]
        self.assertEqual((body["workType"], body["state"]), ("ASSIGNMENT", "DRAFT"))     # лише чернетка, не публікація


class InfographicAndHelpTests(unittest.TestCase):
    def lesson(self, topic):
        from dataclasses import replace
        base = engine.day_lessons("2026-10-01", engine.read_json("Налаштування.json"))[0]
        return replace(base, topic=topic)

    def test_infographic_requirements_for_a_lecture(self):
        from classroom_assistant.chatgpt_bridge import build_prompt
        text = build_prompt(self.lesson("Особливості розвитку культури"))
        for needle in ("ВИМОГИ ДО ІНФОГРАФІКИ", "1536×1024", "ОРИГІНАЛЬНА для цього уроку", "ЖИВА ІЛЮСТРОВАНА",
                       "КОМПОЗИЦІЯ — «", "ТОЧНІСТЬ — найважливіше", "вичитай КОЖНЕ слово українською",
                       "переймай ЛИШЕ стиль", "не вигадуй фактів"):
            self.assertIn(needle, text, needle)
        self.assertNotIn("«Мета роботи»", text)

    def test_infographic_for_a_practical_work_follows_the_sample_structure(self):
        from classroom_assistant.chatgpt_bridge import build_prompt
        text = build_prompt(self.lesson("Практична робота. Аналіз творів"))
        for needle in ("«Мета роботи»", "«Що повторити»", "«Обсяг роботи»", "«Важливо!»", "«Що прикріпити в Classroom?»",
                       "натиснути «Здати»", "ТОЧНІСТЬ — найважливіше"):
            self.assertIn(needle, text, needle)
        self.assertNotIn("«Ключові дати»", text)

    def test_prompt_never_asks_for_homework_in_the_word(self):
        from classroom_assistant.chatgpt_bridge import build_prompt
        text = build_prompt(self.lesson("Тема"))
        self.assertIn("домашнє завдання (воно є лише в повідомленні Classroom", text)
        self.assertNotIn("їх додасть програма", text)

    def test_help_has_no_update_pages_and_explains_the_oauth_json_for_beginners(self):
        from classroom_assistant.help_ui import PAGES
        self.assertFalse([k for k in PAGES if k.startswith(("Оновлення", "Новинки"))])
        google = PAGES["2. Google OAuth JSON"]
        for needle in ("БЕЗ JSON-ФАЙЛУ GOOGLE НЕ ПУСТИТЬ ПРОГРАМУ", "класи й курси з Classroom НЕ підтягнуться",
                       "школа з охоронцем", "ЧИ ЦЕ ПАРОЛЬ? Ні", "ОДИН раз на комп'ютер", "Download JSON", "Desktop app"):
            self.assertIn(needle, google, needle)
        self.assertEqual(len(PAGES), 8)
