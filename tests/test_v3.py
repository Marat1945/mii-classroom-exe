"""Regression tests for universal schedule, local help and Word library."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock
from docx import Document
from classroom_assistant.engine import build_calendar,read_json
from classroom_assistant import material_library as lib
from classroom_assistant.schedule_io import decode_schedule,export_schedule_docx,parse_cell,table_rows
from classroom_assistant.help_ui import PAGES
from classroom_assistant.gui import topic_template,render_template,parse_date,display_date

class Version3Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg=read_json("Налаштування.json")
        cls.calendar=build_calendar(cls.cfg)

    def test_calendar_holidays_kept(self):
        from classroom_assistant.engine import day_lessons,week_phase
        self.assertEqual(day_lessons("2027-03-22",self.cfg),[])
        self.assertEqual(week_phase(__import__("datetime").date(2027,3,29),self.cfg),"чисельник")

    def test_hours_8_grade_and_mismatch_9_grade(self):
        self.assertEqual(lib.schedule_weekly_hours(self.cfg,'8-Б ІУ'),1.5)
        self.assertEqual(lib.schedule_weekly_hours(self.cfg,'8-В ІУ'),1.5)
        self.assertEqual(lib.schedule_weekly_hours(self.cfg,'8-Г ІУ'),1.5)
        self.assertEqual(lib.schedule_weekly_hours(self.cfg,'9-Б ІУ'),2)
        self.assertEqual(lib.schedule_weekly_hours(self.cfg,'9-Г ІУ'),1.5)

    def test_parallel_not_combine_unequal_hours(self):
        item=next(i for i in self.calendar if i.stream=='9-Б ІУ' and i.lesson_number==2)
        matches=lib.parallel_matches(item,self.cfg,calendar=self.calendar)
        self.assertEqual({i.stream for i in matches},{'9-Б ІУ'})

    def test_parallel_three_eighth_classes(self):
        item=next(i for i in self.calendar if i.stream=='8-Б ІУ' and i.lesson_number==7)
        matches=lib.parallel_matches(item,self.cfg,calendar=self.calendar)
        self.assertEqual({x.stream for x in matches},{'8-Б ІУ','8-В ІУ','8-Г ІУ'})

    def test_parallel_not_combine_different_program(self):
        item=next(i for i in self.calendar if i.stream=='8-Б ІУ' and i.lesson_number==7)
        changed=copy.deepcopy(self.cfg)
        changed['course_map']['8-В ІУ']['curriculum_id']='інша навчальна програма'
        matches=lib.parallel_matches(item,changed,calendar=self.calendar)
        self.assertEqual({x.stream for x in matches},{'8-Б ІУ','8-Г ІУ'})

    def test_library_dedup_and_docx_per_date(self):
        lesson=next(i for i in self.calendar if i.stream=='8-Б ІУ' and i.lesson_number==7)
        matches=lib.parallel_matches(lesson,self.cfg,calendar=self.calendar)
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);fil=root/'lecture.docx'
            d=Document();d.add_paragraph('Урок 01.10 — Інша тема.')
            for n in range(12):
                d.add_paragraph("Докладне пояснення церковних подій і зв'язків XVII століття. "*12)
            d.add_paragraph('Д/з: старе завдання')
            d.save(fil)
            with mock.patch.object(lib,'LIBRARY_ROOT',root/'library'),mock.patch.object(lib,'INDEX',root/'library'/'Індекс.json'),mock.patch.object(lib,'ROOT',root):
                item=lib.add_document(fil,lesson,self.cfg)
                second=lib.add_document(fil,lesson,self.cfg)
                self.assertEqual(item['id'],second['id'])
                self.assertEqual(len(lib.load_index()['items']),1)
                state={}
                attached=lib.attach_document(item,matches,state)
                self.assertEqual(len(attached),3)
                for x in matches:
                    path=Path(state['files'][x.unique_key]['path'])
                    self.assertTrue(path.exists())
                    txt='\n'.join(p.text for p in Document(path).paragraphs)
                    self.assertIn('Урок '+x.day[8:10]+'.'+x.day[5:7]+' — '+x.topic,txt)
                    self.assertIn('Д/з: '+x.homework,txt)
                    self.assertTrue(state['files'][x.unique_key]['validated'])
                self.assertEqual(lib.attach_document(item,matches,state),[ ])
                self.assertEqual(len(lib.attach_document(item,matches,state,replace_existing=True)),3)
                # Existing docs with drafts are never rewritten
                state['drafts']={lesson.unique_key: {'id':'google-draft'}}
                self.assertEqual(lib.attach_document(item,[lesson],state),[])

    def test_manual_slot_slash(self):
        opts=self.cfg['course_map']
        self.assertEqual(parse_cell('8-Б ІУ / ГО',opts),['8-Б ІУ','8-Б ГО'])
        self.assertEqual(parse_cell('/ 9-Г Право',opts),[None,'9-Г Право'])
        self.assertEqual(parse_cell('11 ІУ стандарт /',opts),['11 ІУ стандарт',None])
        self.assertEqual(parse_cell('8-Б ВІ',opts),['8-Б ВІ','8-Б ВІ'])

    def test_roundtrip_docx_schedule(self):
        with tempfile.TemporaryDirectory() as d:
            f=Path(d)/'schedule.docx'
            export_schedule_docx(self.cfg,f,show_bells=True,show_meal=True)
            self.assertTrue(f.exists())
            schedule,unknown,numbers=decode_schedule(table_rows(f),self.cfg)
            self.assertFalse(unknown,unknown)
            self.assertEqual(numbers,[1,2,3,4,5,6,7])
            self.assertEqual(schedule,self.cfg['days'])
            text='\n'.join(cell.text for t in Document(f).tables for row in t.rows for cell in row.cells)
            self.assertIn("ХАРЧУВАННЯ У ЇДАЛЬНІ",text)

    def test_schedule_aliases(self):
        cfg=copy.deepcopy(self.cfg)
        cfg['schedule_aliases']={'5-г іу':'5-Г історія'}
        rows=[['№','Понеділок','Вівторок','Середа','Четвер','П’ятниця'],
              ['4','5-Г ІУ','','','','']]
        schedule,missing,number=decode_schedule(rows,cfg)
        self.assertFalse(missing)
        self.assertEqual(schedule['0'][3],['5-Г історія','5-Г історія'])

    def test_manual_import_csv(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'rozklad.csv'
            p.write_text('№;Понеділок;Вівторок;Середа;Четвер;П’ятниця\n1;8-Б ІУ / ГО;8-Б ВІ;11-А ВІ;11 ІУ профіль;8-В ІУ\n2;/ 9-Г Право;;;;11 ІУ стандарт /\n',encoding='utf8')
            schedule,missing,nums=decode_schedule(table_rows(p),self.cfg)
            self.assertFalse(missing)
            self.assertEqual(schedule['0'][0],['8-Б ІУ','8-Б ГО'])
            self.assertEqual(schedule['0'][1],[None,'9-Г Право'])
            self.assertEqual(schedule['4'][1],['11 ІУ стандарт',None])

    def test_help_contains_steps_and_security(self):
        entire='\n'.join(PAGES.values())
        for word in ('ДОВІДКА','Google Classroom API','Google Drive API',
                     'Desktop app','Download JSON','Test users',
                     'Create new secret key','Вставити з буфера','Auto-recharge',
                     'ЧЕРНЕТКУ','не зупиняється'):
            self.assertIn(word,entire)

    def test_description_template_and_date(self):
        lesson=next(i for i in self.calendar if i.stream=='8-Б ІУ' and i.lesson_number==7)
        txt=f"Урок {lesson.day[8:10]}.{lesson.day[5:7]} — {lesson.topic}\nД/з: {lesson.homework}"
        templ=topic_template(txt,lesson)
        for t in ('{date}','{topic}','{homework}'):
            self.assertIn(t,templ)
        self.assertEqual(render_template(templ,lesson),txt)
        self.assertEqual(parse_date("02.11.2026").isoformat(),"2026-11-02")
        self.assertEqual(display_date("2026-11-02"),"02.11.2026")

if __name__=='__main__':unittest.main()
