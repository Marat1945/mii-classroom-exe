"""3.8: КТП за змістом файлу, видалення КТП, курси, канікули, календар, Документи, збереження без КТП."""
import json
import tempfile
import tkinter as tk
from datetime import date
from pathlib import Path
from unittest import mock

from docx import Document

from test_update36 import BaseRoot, TkCase

from classroom_assistant import (course_match, data_tools, datepicker, editor_ui, engine, gui, ktp_detect,
                                 locations, samples_ui, workload_io)
from classroom_assistant.editor_core import validate_working
from classroom_assistant.schedule_io import decode_schedule, table_rows

REAL_FILES = {
    "Календарне_ВІ_9-Б_9-Г_2026-2027_дати_і_ДЗ.docx": (9, {"VI"}, ["9-Б ВІ", "9-Г ВІ"]),
    "Календарне_8_клас_ІУ_2026-2027_дати_і_ДЗ(1).docx": (8, {"IU"}, ["8-Б ІУ", "8-В ІУ", "8-Г ІУ"]),
    "Календарне_ГО_9_2026-2027__дати_і_ДЗ.docx": (9, {"GO"}, ["9-Б ГО", "9-Г ГО"]),
    "Календарне_Правознавство_9_2026-2027_9-Б_9-Г.docx": (9, {"LAW"}, ["9-Б Право", "9-Г Право"]),
    "Календарне планування - історія України 10 ПРОФІЛЬ (3 години =105 год.).docx": (10, {"IU"}, ["10 ІУ профіль"]),
    "Календарне планування - історія України 11 клас 1,5 год.docx": (11, {"IU"}, ["11 ІУ стандарт"]),
    "Календарне планування - історія України 11 клас 3 год (з моїми змінами).doc": (11, {"IU"}, ["11 ІУ профіль"]),
    "Календарне планування - Всесвітня історія 11 клас.docx": (11, {"VI"}, ["11-А ВІ", "11-В ВІ"]),
    "Календарно-тематичне планування - Громадянська освіта 10 клас.docx": (10, {"GO"}, ["10-Б ГО"]),
    "Календарне_5_клас_з_домашніми_завданнями_і_датами(1).docx": (5, set(), ["5-Г ІУ"]),
}
STREAMS = ["5-Г ІУ", "8-Б ІУ", "8-В ІУ", "8-Г ІУ", "9-Б ІУ", "9-Г ІУ", "9-Б ВІ", "9-Г ВІ", "9-Б Право", "9-Г Право",
           "9-Б ГО", "9-Г ГО", "10 ІУ профіль", "10-Б ГО", "11 ІУ профіль", "11 ІУ стандарт", "11-А ВІ", "11-В ВІ"]


def make_ktp(path, header, dates, rows=5):
    doc = Document()
    doc.add_paragraph(header)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("№ з/п", "Дата", "Тема уроку", "Домашнє завдання")):
        cell.text = text
    for number in range(1, rows + 1):
        cells = table.add_row().cells
        cells[0].text = str(number)
        cells[1].text = " ".join(f"{cls} {10 + number:02d}.09." for cls in dates)
        cells[2].text = f"Тема {number}. {header}"
        cells[3].text = f"§ {number}"
    doc.save(str(path))
    return Path(path)


def config_with(streams, titles=None):
    config = data_tools.blank_config()
    plans = {}
    for stream in streams:
        plans[stream] = {"filename": "Очікує файл КТП", "lessons": [], "needs_review": True}
        config["course_map"][stream] = {"plan": stream, "course_title": (titles or {}).get(stream, stream)}
    return config, plans


class DetectionTests(TkCase):
    NEEDS_ROOT = False

    def test_real_file_names_point_to_the_right_classes(self):
        for name, (grade, subjects, expected) in REAL_FILES.items():
            with self.subTest(name=name[:50]):
                detection = ktp_detect.detect(name)
                self.assertEqual(detection.grade, grade)
                self.assertEqual(set(detection.subjects), subjects)
                self.assertEqual(ktp_detect.targets(detection, STREAMS), expected)

    def test_world_history_file_never_matches_ukrainian_history_class(self):
        found = ktp_detect.targets(ktp_detect.detect("Календарне_ВІ_9-Б_9-Г_2026-2027_дати_і_ДЗ.docx"), STREAMS)
        self.assertNotIn("9-Б ІУ", found)

    def test_header_text_and_dates_help_when_the_name_is_vague(self):
        detection = ktp_detect.detect("ктп.docx", "Календарно-тематичне планування. Історія України, 8 клас",
                                      [{"source_dates": {"8-Б": "2026-09-04", "8-Г": "2026-09-03"}}])
        self.assertEqual((detection.grade, set(detection.subjects)), (8, {"IU"}))
        self.assertEqual(ktp_detect.targets(detection, STREAMS), ["8-Б ІУ", "8-Г ІУ"])       # 8-В немає в датах

    def test_profile_and_standard_are_not_guessed(self):
        found = ktp_detect.targets(ktp_detect.detect("Історія України 11 клас.docx"), STREAMS)
        self.assertTrue(ktp_detect.level_conflict(found))


class EditorKtpFlowTests(TkCase):
    def editor(self, streams, titles=None, courses=()):
        config, plans = config_with(streams, titles)
        self.root.state = {}
        self.root.google_courses = [{"name": n, "id": str(i)} for i, n in enumerate(courses)]
        self.root.update_day = lambda: None
        with mock.patch.object(editor_ui, "deep_copy_data", return_value=(config, plans)):
            editor = editor_ui.SchoolEditor(self.root, lambda: None)
        editor.deiconify()
        editor.update()
        self.addCleanup(lambda: self._close(editor))
        return editor

    @staticmethod
    def _close(editor):
        try:
            editor.destroy()
        except tk.TclError:
            pass

    @staticmethod
    def commit(editor):
        """Імпорт автоматичний: вікна перегляду немає (лишилось для старих викликів)."""
        for w in editor.winfo_children():
            if isinstance(w, editor_ui.ImportPreview):
                w.commit()

    def select(self, editor, course):
        index = next(i for i, e in enumerate(editor.list_entries) if e["course"] == course)
        editor.planlist.selection_clear(0, "end")
        editor.planlist.selection_set(index)
        editor.plan_selected()

    def test_world_history_dropped_on_ukrainian_history_class_goes_where_it_belongs(self):
        editor = self.editor(["9-Б ІУ", "9-Б ВІ", "9-Г ВІ"])
        self.select(editor, "9-Б ІУ")
        with tempfile.TemporaryDirectory() as folder:
            ktp = make_ktp(Path(folder) / "Календарне_ВІ_9-Б_9-Г_2026-2027.docx",
                           "Календарне планування. Всесвітня історія, 9 клас", ["9-Б", "9-Г"])
            with mock.patch.object(editor_ui.messagebox, "askyesnocancel",
                                   side_effect=AssertionError("запитань не має бути")), \
                    mock.patch.object(editor_ui.messagebox, "askyesno",
                                      side_effect=AssertionError("запитань не має бути")):
                editor.import_plan(path=ktp)
        cmap = editor.cfg["course_map"]
        self.assertEqual(cmap["9-Б ВІ"]["plan"], cmap["9-Г ВІ"]["plan"])               # спільний КТП
        self.assertEqual(len(editor.plans[cmap["9-Б ВІ"]["plan"]]["lessons"]), 5)
        self.assertEqual(editor.plans[cmap["9-Б ІУ"]["plan"]]["lessons"], [])           # ІУ не зачеплено
        note = editor.history_note.cget("text")
        self.assertIn("9-Б ВІ, 9-Г ВІ", note)
        self.assertIn("а не до «9-Б ІУ»", note)

    def test_import_can_be_undone_with_back(self):
        editor = self.editor(["9-Б ІУ", "9-Б ВІ"])
        before = json.dumps([editor.cfg, editor.plans], sort_keys=True)
        with tempfile.TemporaryDirectory() as folder:
            ktp = make_ktp(Path(folder) / "Календарне_ВІ_9-Б.docx", "Всесвітня історія, 9 клас", ["9-Б"])
            editor.import_plan(path=ktp)
            self.assertNotEqual(json.dumps([editor.cfg, editor.plans], sort_keys=True), before)
            editor.undo()
        self.assertEqual(json.dumps([editor.cfg, editor.plans], sort_keys=True), before)

    def test_dropped_beside_any_class_the_program_decides_by_content(self):
        editor = self.editor(["9-Б ІУ", "9-Б ВІ", "9-Г ВІ"])
        editor.planlist.selection_clear(0, "end")
        editor.current_entry, editor.current_plan = None, None
        with tempfile.TemporaryDirectory() as folder:
            ktp = make_ktp(Path(folder) / "Календарне_ВІ_9-Б_9-Г.docx", "Всесвітня історія, 9 клас", ["9-Б", "9-Г"])
            editor.import_plan(path=ktp)
            self.commit(editor)
        cmap = editor.cfg["course_map"]
        self.assertEqual(len(editor.plans[cmap["9-Г ВІ"]["plan"]]["lessons"]), 5)
        self.assertEqual(editor.plans[cmap["9-Б ІУ"]["plan"]]["lessons"], [])

    def test_course_without_subjects_gets_a_stream_made_from_the_file(self):
        editor = self.editor(["8-Б ІУ"], courses=["8-Б ІУ", "8-Д"])
        self.select(editor, "8-Д")
        with tempfile.TemporaryDirectory() as folder:
            ktp = make_ktp(Path(folder) / "Календарне_8_клас_ІУ.docx", "Історія України, 8 клас", ["8-Д"])
            editor.import_plan(path=ktp)
            self.commit(editor)
        self.assertIn("8-Д ІУ", editor.cfg["course_map"])
        self.assertEqual(editor.cfg["course_map"]["8-Д ІУ"]["course_title"], "8-Д")
        self.assertEqual(len(editor.plans[editor.cfg["course_map"]["8-Д ІУ"]["plan"]]["lessons"]), 5)

    def test_delete_ktp_keeps_class_and_unshares_parallels(self):
        editor = self.editor(["8-Б ІУ", "8-В ІУ"])
        with tempfile.TemporaryDirectory() as folder:
            ktp = make_ktp(Path(folder) / "Календарне_8_клас_ІУ.docx", "Історія України, 8 клас", ["8-Б", "8-В"])
            self.select(editor, "8-Б ІУ")
            editor.import_plan(path=ktp)
            self.commit(editor)
        cmap = editor.cfg["course_map"]
        self.assertEqual(cmap["8-Б ІУ"]["plan"], cmap["8-В ІУ"]["plan"])
        shared = cmap["8-В ІУ"]["plan"]
        self.select(editor, "8-Б ІУ")
        editor.delete_ktp()
        self.assertEqual(len(editor.plans[shared]["lessons"]), 5)                       # 8-В зберіг КТП
        self.assertNotEqual(cmap["8-Б ІУ"]["plan"], shared)
        self.assertEqual(editor.plans[cmap["8-Б ІУ"]["plan"]]["lessons"], [])
        self.assertIn("8-Б ІУ", cmap)                                                    # клас лишився
        self.select(editor, "8-В ІУ")
        editor.delete_ktp()
        self.assertEqual(editor.plans[shared]["lessons"], [])

    def test_menu_has_delete_edit_and_add_course(self):
        editor = self.editor(["8-Б ІУ"])
        popped = []
        with mock.patch.object(tk.Menu, "tk_popup", lambda menu, x, y, entry="": popped.append(menu)):
            bx, by, _w, _h = editor.planlist.bbox(0)
            editor._plan_context_menu(type("E", (), {"x": bx + 5, "y": by + 3, "x_root": 0, "y_root": 0})())
            on_item = [popped[-1].entrycget(i, "label") for i in range(popped[-1].index("end") + 1)
                       if popped[-1].type(i) == "command"]
            editor._plan_context_menu(type("E", (), {"x": 5, "y": 400, "x_root": 0, "y_root": 0})())
            empty = [popped[-1].entrycget(i, "label") for i in range(popped[-1].index("end") + 1)
                     if popped[-1].type(i) == "command"]
        self.assertTrue(any("Редагувати курс" in x for x in on_item))
        self.assertTrue(any("Додати новий курс" in x for x in on_item))
        self.assertTrue(any("Видалити КТП" in x for x in on_item))
        self.assertTrue(any("Додати новий курс" in x for x in empty))
        self.assertFalse(any("Видалити" in x for x in empty))

    def test_combined_course_is_edited_and_new_course_added(self):
        editor = self.editor(["9-Б Право"], titles={"9-Б Право": "9-Б Право + ГО"}, courses=["9-Б Право + ГО", "9-В ІУ"])
        self.select(editor, "9-Б Право + ГО")
        editor.edit_course(False)
        dialog = editor.winfo_children()[-1]
        api = dialog.api
        api.subject.set("ГО")
        api.add()                                                                         # «9-Б ГО» у тому ж курсі
        self.assertEqual(api.streams, ["9-Б Право", "9-Б ГО"])
        api.ok()
        cmap = editor.cfg["course_map"]
        self.assertEqual(cmap["9-Б ГО"]["course_title"], "9-Б Право + ГО")
        rows = [e for e in editor.list_entries if e["course"] == "9-Б Право + ГО"]
        self.assertEqual([e["streams"] for e in rows], [["9-Б ГО"], ["9-Б Право"]])     # окремий рядок на предмет
        self.assertTrue(all(e["course_streams"] == ["9-Б ГО", "9-Б Право"] for e in rows))
        self.select(editor, "9-В ІУ")
        editor.edit_course(True)
        dialog = editor.winfo_children()[-1]
        dialog.api.title.set("9-В ІУ")
        dialog.api.cls.set("9-В")
        dialog.api.add()
        dialog.api.ok()
        self.assertIn("9-В ІУ", editor.cfg["course_map"])

    def test_classes_tab_explains_links_and_auto_link_fixes_names(self):
        editor = self.editor(["8-Б ВІ", "5-Г ІУ"], courses=["8-Б ВI", "5-Г історія"])      # латинська I у Classroom
        titles = {s: i["course_title"] for s, i in editor.cfg["course_map"].items()}
        self.assertEqual(titles, {"8-Б ВІ": "8-Б ВI", "5-Г ІУ": "5-Г історія"})            # підставлено при відкритті
        row = [editor.streamtree.item(i, "values") for i in editor.streamtree.get_children()]
        self.assertTrue(all("Classroom ✓" in v[3] and "без КТП" in v[3] for v in row), row)
        self.assertEqual(editor.auto_link(), {})                                          # вдруге нічого не змінює

    def test_saving_without_ktp_is_allowed(self):
        editor = self.editor(["8-Б ІУ", "8-В ІУ"])
        editor.cfg["days"]["0"][0] = ["8-Б ІУ", "8-В ІУ"]
        self.assertTrue(any(e.startswith("Порожній КТП") for e in validate_working(editor.cfg, editor.plans)))
        errors = []
        with mock.patch.object(editor_ui, "persist", return_value="backup") as persist, \
                mock.patch.object(editor_ui, "save_state"), \
                mock.patch.object(editor_ui.messagebox, "showerror", side_effect=lambda *a, **k: errors.append(a)):
            editor.parent.update_day = lambda: None
            editor.save(confirm=False)
        self.assertEqual(errors, [])
        persist.assert_called_once()

    def test_empty_ktp_is_shown_as_not_loaded(self):
        config, plans = config_with(["8-Б ІУ"])
        config["days"]["0"][0] = ["8-Б ІУ", "8-Б ІУ"]
        lessons = engine.build_calendar(config, plans)
        self.assertTrue(lessons)
        self.assertEqual({x.status for x in lessons}, {"КТП не завантажено"})
        self.assertIn("КТП ще не завантажено", lessons[0].topic)


class HolidaysAndCalendarTests(TkCase):
    def test_holidays_roundtrip_through_the_workload_title(self):
        config = samples_ui.sample_workload_config()
        config["course_map"] = {}
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "n.docx"
            workload_io.export_workload_docx(config, path)
            rows = table_rows(path)
        info = workload_io.parse_workload_title(rows, "2026-09-01")
        self.assertEqual(info["holidays"], [{"start": "2026-10-26", "end": "2026-11-01"},
                                            {"start": "2026-12-24", "end": "2027-01-10"},
                                            {"start": "2027-03-22", "end": "2027-03-28"}])
        self.assertIn("канікули  26.10-01.11, 24.12-10.01, 22.03-28.03",
                      "".join(part[0] for part in workload_io.title_parts(config)))

    def test_title_without_holidays_has_none(self):
        self.assertEqual(workload_io.holidays_from_title("НАВАНТАЖЕННЯ (Іваненко І.І.) // 123456", "2026-09-01"), [])

    def test_import_offers_to_take_holidays_from_the_file(self):
        config, plans = config_with([])
        self.root.state, self.root.google_courses, self.root.update_day = {}, [], lambda: None
        with mock.patch.object(editor_ui, "deep_copy_data", return_value=(config, plans)):
            editor = editor_ui.SchoolEditor(self.root, lambda: None)
        editor.update()
        def close():
            try:
                editor.destroy()
            except tk.TclError:
                pass
        self.addCleanup(close)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "n.docx"
            samples_ui.build_sample_workload(path)
            editor.import_schedule(path=path)
        self.assertEqual([(h["start"], h["end"]) for h in editor.cfg["holidays"]],
                         [("2026-10-26", "2026-11-01"), ("2026-12-24", "2027-01-10"), ("2027-03-22", "2027-03-28")])

    def test_calendar_opens_on_the_working_month_and_blocks_earlier_days(self):
        dialog = datepicker.CalendarDialog(self.root, None, default=date(2027, 3, 1), min_date=date(2027, 3, 22))
        self.assertEqual((dialog.year, dialog.month), (2027, 3))
        dialog.choose(date(2027, 3, 10))
        self.assertIsNone(dialog.result)
        dialog.choose(date(2027, 3, 28))
        self.assertEqual(dialog.result, date(2027, 3, 28))

    def test_date_field_asks_for_default_and_minimum_lazily(self):
        variable = tk.StringVar()
        seen = {}
        def fake(parent, text, weekday, title, default=None, min_date=None):
            seen.update(default=default, min_date=min_date)
            return "28.03.2027"
        field = datepicker.DateField(self.root, variable, default=lambda: date(2027, 3, 22),
                                     min_date=lambda: date(2027, 3, 22))
        with mock.patch.object(datepicker, "pick_date", fake):
            field.open()
        self.assertEqual(seen, {"default": date(2027, 3, 22), "min_date": date(2027, 3, 22)})
        self.assertEqual(variable.get(), "28.03.2027")


class LocationAndMainWindowTests(TkCase):
    NEEDS_ROOT = False

    def test_data_moves_to_documents_once_and_old_copy_stays(self):
        with tempfile.TemporaryDirectory() as folder:
            documents, legacy = Path(folder) / "Documents", Path(folder) / "AppData" / "old"
            (legacy / "data").mkdir(parents=True)
            (legacy / "data" / "Налаштування.json").write_text('{"мої": 1}', encoding="utf-8")
            (legacy / "Готові Word").mkdir()
            (legacy / "Готові Word" / "a.docx").write_text("w", encoding="utf-8")
            root, where = locations.resolve_root(documents, legacy)
            self.assertEqual((root, where), (documents / locations.APP_FOLDER, "migrated"))
            self.assertEqual((root / "data" / "Налаштування.json").read_text("utf-8"), '{"мої": 1}')
            self.assertTrue((root / "Готові Word" / "a.docx").exists())
            self.assertTrue((legacy / "data" / "Налаштування.json").exists())            # старе не видалено
            (root / "data" / "Налаштування.json").write_text('{"нове": 2}', encoding="utf-8")
            self.assertEqual(locations.resolve_root(documents, legacy), (root, "existing"))
            self.assertEqual((root / "data" / "Налаштування.json").read_text("utf-8"), '{"нове": 2}')

    def test_fresh_install_and_failed_copy(self):
        with tempfile.TemporaryDirectory() as folder:
            documents = Path(folder) / "Documents"
            root, where = locations.resolve_root(documents, Path(folder) / "none")
            self.assertEqual((root, where), (documents / locations.APP_FOLDER, "new"))
            self.assertTrue(root.is_dir())
        with tempfile.TemporaryDirectory() as folder:
            legacy = Path(folder) / "old"
            (legacy / "data").mkdir(parents=True)
            (legacy / "data" / "Налаштування.json").write_text("{}", encoding="utf-8")
            with mock.patch.object(locations.shutil, "copytree", side_effect=OSError("диск повний")):
                self.assertEqual(locations.resolve_root(Path(folder) / "D", legacy), (legacy, "legacy"))

    def test_window_has_no_version_in_the_title(self):
        app = gui.MainApp()
        try:
            self.assertEqual(app.title(), "Помічник учителя Classroom")
        finally:
            app.destroy()

    def test_sync_links_course_names_by_meaning_and_remembers_ids(self):
        app = gui.MainApp()
        try:
            app.cfg["course_map"] = {"8-Б ВІ": {"plan": "p", "course_title": "8-Б ВІ"}}
            app.state["course_ids"] = {}
            with mock.patch.object(gui, "write_json") as write:
                app._link_courses_automatically([{"name": "8-Б ВI", "id": 5}])
            self.assertEqual(app.cfg["course_map"]["8-Б ВІ"]["course_title"], "8-Б ВI")
            self.assertEqual(app.state["course_ids"], {"8-Б ВI": "5"})
            write.assert_called_once()
        finally:
            app.destroy()

    def test_missing_ktp_message_is_clear(self):
        lesson = type("L", (), {"status": "КТП не завантажено"})()
        self.assertIn("перетягніть", gui.MainApp._not_ready_text(lesson))


class SampleTabTests(TkCase):
    def test_schedule_sample_tab_has_description_picture_and_ready_files(self):
        win = samples_ui.show_samples(self.root)
        try:
            win.update()
            notebook = next(w for w in win.winfo_children() if w.winfo_class() == "TNotebook")
            tabs = [notebook.tab(i, "text") for i in range(notebook.index("end"))]
            self.assertIn("Зразок розкладу", tabs)
            self.assertIsNotNone(getattr(win, "_sample_image", None))
            self.assertIn("ХАРЧУВАННЯ У ЇДАЛЬНІ", samples_ui.SCHEDULE_DESCRIPTION)
            self.assertIn("канікули", samples_ui.SCHEDULE_DESCRIPTION)
        finally:
            win.destroy()

    def test_ready_samples_are_valid_and_readable_by_the_program(self):
        config = samples_ui.sample_workload_config()
        config["course_map"] = {n: {} for n in {x for ps in config["days"].values() for p in ps for x in p if x}}
        with tempfile.TemporaryDirectory() as folder:
            jpg = samples_ui.build_sample_workload_jpeg(Path(folder) / "s.jpg")
            self.assertGreater(jpg.stat().st_size, 20000)
            csv_path = samples_ui.build_sample_workload_csv(Path(folder) / "s.csv")
            schedule, unknown, found = decode_schedule(table_rows(csv_path), config)
        self.assertEqual(unknown, [])
        self.assertEqual(schedule["0"][0], ["8-Б ІУ", "8-Б ГО"])
