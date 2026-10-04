"""4.2: Word і зображення пачкою за назвою «9-Б ВІ, Урок 05.10 — Тема»; предмети спільного курсу; чисельник/знаменник."""
import json
import shutil
import tempfile
import tkinter as tk
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest import mock

from docx import Document
from PIL import Image

from test_update36 import TkCase
from test_update38 import config_with

from classroom_assistant import (attachments, data_tools, editor_core, editor_ui, engine, file_match, gui,
                                 material_library)
from classroom_assistant.chatgpt_bridge import build_prompt, infographic_name
from classroom_assistant.engine import Lesson, build_calendar, lesson_base_name, read_json

REPO_DATA = Path(__file__).resolve().parents[1] / "data"


def L(day, period, stream, topic, number=1):
    return Lesson(day, 0, period, "08:30", "09:15", stream, stream, "p", number, topic, "§ 1", "ktp.docx")


class FileMatchTests(unittest.TestCase):
    def test_name_formats(self):
        parsed = file_match.parse_name("9-Б ВІ, Урок 05.10 — Практична робота. Історичне значення Французької революції")
        self.assertEqual((parsed.stream, parsed.day, parsed.month), ("9-Б ВІ", 5, 10))
        self.assertTrue(parsed.topic.startswith("Практична робота."))
        self.assertEqual(file_match.parse_file("/x/9-Б ВІ, Урок 05.10 — Тема (1).docx").topic, "Тема")   # браузер додав (1)
        old = file_match.parse_name("Урок № 8, 02.10.2026 Культура")
        self.assertEqual((old.stream, old.number, old.day, old.month), (None, 8, 2, 10))
        self.assertEqual(file_match.parse_name("Урок 28.09 — Тема.").stream, None)
        self.assertIsNone(file_match.parse_name("просто якась назва"))

    def test_class_plus_date_pick_exactly_one_lesson(self):
        lessons = [L("2026-10-05", 1, "8-Б ІУ", "Церковне життя"), L("2026-10-05", 2, "8-В ІУ", "Церковне життя"),
                   L("2026-10-06", 1, "8-Б ІУ", "Інша тема")]
        found = file_match.find_lesson(lessons, file_match.parse_name("8-В ІУ, Урок 05.10 — Церковне життя"))
        self.assertEqual((found.stream, found.day), ("8-В ІУ", "2026-10-05"))
        self.assertIsNone(file_match.find_lesson(lessons, file_match.parse_name("8-Г ІУ, Урок 05.10 — Тема")))
        self.assertIsNone(file_match.find_lesson(lessons, file_match.parse_name("8-Б ІУ, Урок 09.10 — Тема")))

    def test_spelling_of_the_class_does_not_matter(self):
        lessons = [L("2026-10-05", 1, "8-Б ВІ", "Тема")]
        self.assertIsNotNone(file_match.find_lesson(lessons, file_match.parse_name("8-Б ВI, Урок 05.10 — Тема")))   # латинська I

    def test_no_class_in_the_name_is_never_guessed_between_parallels(self):
        lessons = [L("2026-10-05", 1, "8-Б ІУ", "Церковне життя"), L("2026-10-05", 2, "8-В ІУ", "Церковне життя")]
        self.assertIsNone(file_match.find_lesson(lessons, file_match.parse_name("Урок 05.10 — Церковне життя")))
        alone = [L("2026-10-05", 1, "8-Б ІУ", "Церковне життя"), L("2026-10-05", 2, "9-Б ВІ", "Зовсім інше")]
        self.assertEqual(file_match.find_lesson(alone, file_match.parse_name("Урок 05.10 — Церковне життя")).stream, "8-Б ІУ")

    def test_punctuation_removed_from_the_file_name_does_not_break_the_match(self):
        lessons = [L("2026-09-01", 6, "10-Б ГО", "Що таке ідентичність? Види ідентичності."),
                   L("2026-09-01", 7, "10-Б ГО", "Самореалізація людини, її розвиток.")]
        name = lesson_base_name(lessons[0])
        self.assertNotIn("?", name)
        self.assertEqual(file_match.find_lesson(lessons, file_match.parse_name(name)).period, 6)
        self.assertEqual(file_match.find_lesson(lessons, file_match.parse_name(lesson_base_name(lessons[1]))).period, 7)

    def test_double_lesson_is_told_apart_by_topic(self):
        lessons = [L("2026-10-05", 1, "8-Б ІУ", "Перша тема дня"), L("2026-10-05", 2, "8-Б ІУ", "Друга тема дня")]
        self.assertEqual(file_match.find_lesson(lessons, file_match.parse_name("8-Б ІУ, Урок 05.10 — Друга тема дня")).period, 2)

    def test_renamed_download_is_recognised_by_its_first_line(self):
        parsed = file_match.parse_file("Новий документ (3).docx", "9-Б ВІ, Урок 05.10 — Тема")
        self.assertEqual((parsed.stream, parsed.day), ("9-Б ВІ", 5))
        with tempfile.TemporaryDirectory() as folder:
            doc = Document()
            doc.add_paragraph("")
            doc.add_paragraph("9-Б ВІ, Урок 05.10 — Тема")
            path = Path(folder) / "щось.docx"
            doc.save(path)
            self.assertEqual(file_match.first_line_of_docx(path), "9-Б ВІ, Урок 05.10 — Тема")


class NamingAndPromptTests(unittest.TestCase):
    def test_names_start_with_class_and_date_and_are_file_safe(self):
        lesson = L("2026-10-05", 1, "9-Б ВІ", "Практична робота. Історичне значення Французької революції.")
        self.assertEqual(lesson_base_name(lesson), "9-Б ВІ, Урок 05.10 — Практична робота. Історичне значення Французької революції")
        dirty = L("2026-10-05", 1, "9-Б ВІ", 'Що таке «право»: "так" чи ні? ' * 8)
        name = lesson_base_name(dirty)
        self.assertLessEqual(len(name), 120)
        self.assertFalse(set('<>:"/\\|?*') & set(name))
        self.assertEqual(file_match.parse_name(name).stream, "9-Б ВІ")

    def test_word_and_picture_share_one_name_and_prompt_says_so(self):
        lesson = L("2026-10-05", 1, "9-Б ВІ", "Практична робота")
        self.assertEqual(infographic_name(lesson), lesson_base_name(lesson))
        prompt = build_prompt(lesson)
        self.assertIn("«9-Б ВІ, Урок 05.10 — Практична робота.docx»", prompt)
        self.assertIn("«9-Б ВІ, Урок 05.10 — Практична робота.png»", prompt)
        self.assertIn("«9-Б ВІ, Урок 05.10 — Практична робота.» — жирний", prompt)         # перший рядок документа
        self.assertNotIn("Інфографіка.", prompt)


def make_lecture(path, first_line):
    doc = Document()
    doc.add_paragraph(first_line)
    for n in range(9):
        doc.add_paragraph(f"Розділ {n}. " + "Змістовний абзац лекції про цю тему. " * 8)
    doc.save(str(path))
    return Path(path)


def make_picture(path):
    Image.new("RGB", (60, 40), "#cfe8ff").save(path)
    return Path(path)


class BulkLectureDropTests(TkCase):
    NEEDS_ROOT = False

    def setUp(self):
        super().setUp()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root_dir = Path(self.tmp.name) / "root"
        (self.root_dir / "data").mkdir(parents=True)
        for name in ("Налаштування.json", "Календарні плани.json"):
            shutil.copy(REPO_DATA / name, self.root_dir / "data" / name)
        stack = ExitStack()
        for patcher in (mock.patch.object(engine, "DATA", self.root_dir / "data"),
                        mock.patch.object(engine, "ROOT", self.root_dir),
                        mock.patch.object(gui, "DATA", self.root_dir / "data"),
                        mock.patch.object(material_library, "LIBRARY_ROOT", self.root_dir / "Бібліотека"),
                        mock.patch.object(material_library, "INDEX", self.root_dir / "Бібліотека" / "Індекс.json"),
                        mock.patch.object(attachments, "ATTACHMENT_ROOT", self.root_dir / "Вкладення"),
                        mock.patch.object(gui.messagebox, "askyesno", side_effect=AssertionError("запитання!")),
                        mock.patch.object(gui.messagebox, "askyesnocancel", side_effect=AssertionError("запитання!"))):
            stack.enter_context(patcher)
        self.addCleanup(stack.close)
        self.shown = []
        for name in ("showinfo", "showwarning", "showerror"):
            stack.enter_context(mock.patch.object(gui.messagebox, name,
                                                  side_effect=lambda *a, **k: self.shown.append(a[1])))
        self.app = gui.MainApp()

        def close():
            try:
                self.app.destroy()
            except tk.TclError:
                pass
        self.addCleanup(close)
        self.lessons = [x for x in build_calendar(self.app.cfg) if x.status == "готово"]
        self.files = Path(self.tmp.name) / "завантажено"
        self.files.mkdir()

    def lesson(self, stream, index=0):
        return [x for x in self.lessons if x.stream == stream][index]

    def lecture(self, lesson):
        name = lesson_base_name(lesson)
        return (make_lecture(self.files / (name + ".docx"), f"{lesson.stream}, Урок {lesson.day[8:10]}.{lesson.day[5:7]} — {lesson.topic}"),
                make_picture(self.files / (name + ".png")))

    def test_many_words_and_pictures_find_their_own_lessons_in_one_drop(self):
        a, b, c = self.lesson("8-Б ІУ"), self.lesson("9-Б ВІ"), self.lesson("10-Б ГО")
        paths = []
        for lesson in (a, b, c):
            paths += list(self.lecture(lesson))
        stray = make_lecture(self.files / "щось інше.docx", "Просто текст без класу й дати")
        self.app.import_lecture_files(paths + [stray])
        for lesson in (a, b, c):
            entry = self.app.state["files"][lesson.unique_key]
            self.assertTrue(entry["complete"], lesson.stream)
            texts = [p.text for p in Document(entry["path"]).paragraphs if p.text.strip()]
            label = engine.same_day_label(lesson, material_library.parallel_matches(
                lesson, self.app.cfg, calendar=self.lessons))
            self.assertEqual(texts[0], f"{label}, Урок {lesson.day[8:10]}.{lesson.day[5:7]} — {lesson.topic}")
            self.assertFalse(any(x.startswith("Д/з") for x in texts))
            self.assertEqual(len(files_for(self.app.state, lesson)), 1, lesson.stream)         # картинка — у вкладеннях
        report = self.shown[-1]
        self.assertIn("Розпізнано уроків: 3", report)
        self.assertIn("щось інше.docx — не вдалося визначити урок", report)
        self.assertIn("У Classroom нічого не створено", report)

    def test_parallel_classes_get_the_lecture_too(self):
        a = self.lesson("8-Б ІУ", 2)
        word, picture = self.lecture(a)
        self.app.import_lecture_files([word, picture])
        spread = [x for x in self.lessons if x.stream in ("8-В ІУ", "8-Г ІУ") and x.lesson_number == a.lesson_number]
        got = [x for x in spread if self.app.state["files"].get(x.unique_key, {}).get("complete")]
        self.assertTrue(got, "паралелі мали отримати цей Word")
        self.assertIn("також для:", self.shown[-1] if self.shown else self.app.foot.cget("text"))
        for x in got:
            self.assertEqual(len(files_for(self.app.state, x)), 1)

    def test_picture_without_word_and_existing_draft_and_newest_word_wins(self):
        a, b = self.lesson("9-Б ВІ", 3), self.lesson("10-Б ГО", 3)
        picture = make_picture(self.files / (lesson_base_name(a) + ".png"))
        self.app.import_lecture_files([picture])
        self.assertEqual(len(files_for(self.app.state, a)), 1)
        self.assertNotIn(a.unique_key, self.app.state.get("files", {}))
        self.app.state.setdefault("drafts", {})[b.unique_key] = {"id": "1"}
        word, _ = self.lecture(b)
        self.app.import_lecture_files([word])
        self.assertIn("чернетка вже створена", self.shown[-1])
        self.assertNotIn(b.unique_key, self.app.state.get("files", {}))
        first = make_lecture(self.files / "9-Б ВІ, Урок старий.docx", "x")
        c = self.lesson("8-Г ВІ", 1)
        old = make_lecture(self.files / (lesson_base_name(c) + ".docx"), f"{c.stream}, Урок {c.day[8:10]}.{c.day[5:7]} — {c.topic}")
        new = make_lecture(self.files / (lesson_base_name(c) + " (1).docx"), f"{c.stream}, Урок {c.day[8:10]}.{c.day[5:7]} — {c.topic}")
        import os
        os.utime(old, (1000, 1000))
        self.app.import_lecture_files([old, new])
        self.assertIn("з кількох Word взято найновіший", self.app.foot.cget("text"))     # один урок — без вікна
        self.assertTrue(self.app.state["files"][c.unique_key]["complete"])

    def test_dropping_on_a_row_still_goes_by_name_and_keeps_old_behaviour_for_the_rest(self):
        a = self.lesson("9-Б ВІ", 2)
        word, picture = self.lecture(a)
        self.app.datevar.set("02.10.2026")
        self.app.update_day()
        self.assertTrue(self.app.rows)
        extra = self.files / "довідка.pdf"
        extra.write_bytes(b"%PDF-1.4 test")
        data = " ".join("{" + str(p) + "}" for p in (word, picture, extra))
        with mock.patch.object(self.app.grid, "identify_row", return_value="0"):
            self.app._drop_on_lesson(type("E", (), {"data": data})())
        self.assertTrue(self.app.state["files"][a.unique_key]["complete"])            # за назвою, а не за рядком
        row_lesson = self.app.rows[0]
        names = [x["name"] for x in self.app.state.get("attachments", {}).get(row_lesson.unique_key, [])]
        self.assertTrue(any("довідка" in n for n in names), names)                    # нерозпізнане — до рядка

    def test_the_whole_window_accepts_files_and_shows_a_hint(self):
        for widget in (self.app, self.app.foot, self.app.undo_button, self.app.date_entry):
            self.assertTrue(self.app.tk.call("bind", str(widget), "<<DropTargetTypes>>"), widget)
        self.app._show_window_hint()
        self.app.update()
        self.assertEqual(self.app.window_hint.winfo_manager(), "place")
        self.app._hide_window_hint()
        self.assertEqual(self.app.window_hint.winfo_manager(), "")
        a = self.lesson("10-Б ГО", 4)
        word, picture = self.lecture(a)
        data = "{" + str(word) + "} {" + str(picture) + "}"
        self.app._drop_anywhere(type("E", (), {"data": data})())
        self.assertTrue(self.app.state["files"][a.unique_key]["complete"])

    def test_button_is_in_the_menu_and_toolbar(self):
        texts = []
        stack = [self.app]
        while stack:
            widget = stack.pop()
            stack.extend(widget.winfo_children())
            if widget.winfo_class() == "TButton":
                texts.append(str(widget.cget("text")))
        self.assertIn("Word і картинки пачкою…", texts)


def files_for(state, lesson):
    return attachments.files_for_lesson(state, lesson.unique_key)


class MessageAndTasksTests(unittest.TestCase):
    def lesson(self, topic):
        from dataclasses import replace
        base = engine.day_lessons("2026-10-01", read_json("Налаштування.json"))[0]
        return replace(base, topic=topic)

    def test_standard_message_uses_asterisks_instead_of_numbers(self):
        text = engine.html_classroom_text(self.lesson("Особливості розвитку культури"))
        for line in ("* Перегляньте прикріплене відео до теми.",
                     "* Уважно опрацюйте прикріплений матеріал уроку.",
                     "* Наприкінці Word-документа знайдіть «ПЛАН-КОНСПЕКТ УРОКУ ДЛЯ ЗАПИСУ В ЗОШИТ»."):
            self.assertIn(line, text)
        for old in ("1. Перегляньте", "2. Уважно", "3. Наприкінці", "4. ✍"):
            self.assertNotIn(old, text)
        self.assertNotIn("Перегляньте прикріплене відео", engine.html_classroom_text(self.lesson("Тема"), video=False))

    def test_practical_and_control_lessons_ask_for_up_to_three_analytical_tasks(self):
        for topic in ("Практична робота. Історичне значення Французької революції", "Практичне заняття",
                      "Тематичне оцінювання", "Контрольна робота"):
            prompt = build_prompt(self.lesson(topic))
            self.assertIn("НЕ БІЛЬШЕ ніж 3 аналітичними завданнями", prompt, topic)
            self.assertIn("в інтернеті", prompt)
            self.assertIn("прикріпити фото", prompt)
            self.assertIn("прикріпити до завдання в Google Classroom", prompt)
            self.assertIn("ДЛЯ ЦЬОГО УРОКУ", prompt)
            self.assertNotIn("8–12 змістовних тематичних розділів", prompt, topic)         # без суперечності з завданнями
            self.assertIn("коротке повторення теорії", prompt)
            text = engine.html_classroom_text(self.lesson(topic))
            self.assertIn("* Виконайте завдання (до 3) з Word-документа", text)
            self.assertIn("прикріпити до завдання в Classroom", text)
            self.assertNotIn("знайдіть «ПЛАН-КОНСПЕКТ", text)
        for topic in ("Особливості розвитку культури", "Урок узагальнення з теми"):
            self.assertNotIn("НЕ БІЛЬШЕ ніж 3", build_prompt(self.lesson(topic)), topic)
            self.assertIn("8–12 змістовних тематичних розділів", build_prompt(self.lesson(topic)), topic)
            self.assertIn("знайдіть «ПЛАН-КОНСПЕКТ", engine.html_classroom_text(self.lesson(topic)))


class SharedCourseEditorTests(TkCase):
    def editor(self, streams, titles=None, courses=()):
        config, plans = config_with(streams, titles)
        self.root.state, self.root.update_day = {}, lambda: None
        self.root.google_courses = [{"name": n, "id": str(i)} for i, n in enumerate(courses)]
        with mock.patch.object(editor_ui, "deep_copy_data", return_value=(config, plans)):
            editor = editor_ui.SchoolEditor(self.root, lambda: None)
        editor.deiconify()
        editor.update()

        def close():
            try:
                editor.destroy()
            except tk.TclError:
                pass
        self.addCleanup(close)
        return editor

    def combined(self):
        return self.editor(["9-Б Право", "9-Б ГО"], {"9-Б Право": "9-Б Право + ГО", "9-Б ГО": "9-Б Право + ГО"},
                           courses=["9-Б Право + ГО"])

    def select(self, editor, stream):
        index = next(i for i, e in enumerate(editor.list_entries) if stream in e["streams"])
        editor.planlist.selection_clear(0, "end")
        editor.planlist.selection_set(index)
        editor.plan_selected()

    def test_each_subject_of_a_shared_course_is_its_own_row(self):
        editor = self.combined()
        self.assertEqual([e["label"] for e in editor.list_entries],
                         ["9-Б Право + ГО  ▸ ГО   — без КТП", "9-Б Право + ГО  ▸ Право   — без КТП"])

    def test_rename_moves_schedule_and_pending_saved_lessons(self):
        editor = self.combined()
        editor.cfg["days"]["0"][5] = ["9-Б Право", "9-Б ГО"]
        editor.rename_stream("9-Б Право", "9-Б Правознавство")
        self.assertEqual(editor.cfg["days"]["0"][5], ["9-Б Правознавство", "9-Б ГО"])
        self.assertIn("9-Б Правознавство", editor.cfg["course_map"])
        self.assertNotIn("9-Б Право", editor.cfg["course_map"])
        self.assertEqual(editor.cfg["pending_stream_renames"], {"9-Б Право": "9-Б Правознавство"})
        editor.rename_stream("9-Б Правознавство", "9-Б Правознавство-2")
        self.assertEqual(editor.cfg["pending_stream_renames"], {"9-Б Право": "9-Б Правознавство-2"})   # ланцюжок
        editor.rename_stream("9-Б Правознавство-2", "9-Б Право")
        self.assertNotIn("pending_stream_renames", editor.cfg)                                       # повернули як було
        with self.assertRaises(ValueError):
            editor.rename_stream("9-Б Право", "9-Б ГО")                                              # зайнято
        with self.assertRaises(ValueError):
            editor.rename_stream("9-Б Право", "x|y")

    def test_saving_a_rename_keeps_word_attachments_and_drafts(self):
        editor = self.combined()
        old, new = "9-Б Право", "9-Б Правознавство"
        editor.parent.state.update({
            "files": {f"2026-10-05|3|{old}": {"path": "w.docx"}, "2026-10-05|3|9-Б ГО": {"path": "g.docx"}},
            "drafts": {f"2026-10-05|3|{old}": {"id": "7"}},
            "attachments": {f"2026-10-12|3|{old}": [{"name": "a.png"}]}})
        editor.rename_stream(old, new)
        with mock.patch.object(editor_ui, "persist", return_value="копія"), mock.patch.object(editor_ui, "save_state"):
            editor.parent.update_day = lambda: None
            editor.save(confirm=False)
        state = self.root.state
        self.assertIn(f"2026-10-05|3|{new}", state["files"])
        self.assertIn("2026-10-05|3|9-Б ГО", state["files"])                      # інший предмет не зачеплено
        self.assertEqual(state["drafts"], {f"2026-10-05|3|{new}": {"id": "7"}})   # чернетки не губляться
        self.assertEqual(list(state["attachments"]), [f"2026-10-12|3|{new}"])

    def test_migration_helper_touches_only_matching_lesson_keys(self):
        state = {"files": {"2026-10-05|1|А": 1, "2026-10-05|1|Б": 2}, "other": {"А": 3}, "text": "А"}
        self.assertEqual(editor_core.migrate_stream_keys(state, {"А": "В"}), 1)
        self.assertEqual(state, {"files": {"2026-10-05|1|В": 1, "2026-10-05|1|Б": 2}, "other": {"А": 3}, "text": "А"})

    def test_f2_and_menu_rename_a_subject_straight_from_the_list(self):
        editor = self.combined()
        self.assertTrue(editor.planlist.bind("<F2>"))
        self.select(editor, "9-Б Право")
        with mock.patch.object(editor, "_ask_text", return_value="9-Б Правознавство"):
            editor.rename_subject()
        self.assertIn("9-Б Правознавство", editor.cfg["course_map"])
        popped = []
        with mock.patch.object(tk.Menu, "tk_popup", lambda menu, x, y, entry="": popped.append(menu)):
            box = editor.planlist.bbox(0)
            editor._plan_context_menu(type("E", (), {"x": box[0] + 5, "y": box[1] + 3, "x_root": 0, "y_root": 0})())
        labels = [popped[-1].entrycget(i, "label") for i in range(popped[-1].index("end") + 1)
                  if popped[-1].type(i) == "command"]
        self.assertTrue(any("Перейменувати предмет" in x for x in labels), labels)
        self.assertTrue(any("чисельник / знаменник" in x for x in labels), labels)

    def test_rename_button_inside_the_course_dialog(self):
        editor = self.combined()
        self.select(editor, "9-Б Право")
        editor.edit_course(False)
        dialog = editor.winfo_children()[-1]
        dialog.api.box.selection_set(dialog.api.streams.index("9-Б Право"))
        with mock.patch.object(editor, "_ask_text", return_value="9-Б Правознавство"):
            dialog.api.rename()
        self.assertEqual(sorted(dialog.api.streams), ["9-Б ГО", "9-Б Правознавство"])
        dialog.api.ok()                                                           # не видаляє «старий» потік
        self.assertEqual(sorted(editor.cfg["course_map"]), ["9-Б ГО", "9-Б Правознавство"])

    def test_numerator_and_denominator_schedule_for_alternating_subjects(self):
        editor = self.combined()
        self.select(editor, "9-Б Право")
        editor.edit_course_schedule()
        dialog = editor.winfo_children()[-1]
        api = dialog.api
        self.assertEqual(sorted(api.subjects), ["9-Б ГО", "9-Б Право"])
        api.day.set("Середа")
        api.number.set("6")
        api.top.set("9-Б Право")
        api.bottom.set("9-Б ГО")
        api.assign()
        self.assertEqual(editor.cfg["days"]["2"][5], ["9-Б Право", "9-Б ГО"])
        self.assertIn("Середа, урок 6:  чисельник — 9-Б Право;  знаменник — 9-Б ГО", api.box.get(0))
        api.day.set("П’ятниця")
        api.number.set("2")
        api.top.set("9-Б ГО")
        api.bottom.set("— немає —")                                              # раз на два тижні
        api.assign()
        self.assertEqual(editor.cfg["days"]["4"][1], ["9-Б ГО", None])
        api.box.selection_set(0)
        api.remove()
        self.assertEqual(editor.cfg["days"]["2"][5], [None, None])
        self.assertEqual(len(api.places), 1)

    def test_a_foreign_lesson_in_the_slot_needs_confirmation(self):
        editor = self.editor(["9-Б Право", "9-Б ГО", "8-Б ІУ"], {"9-Б Право": "9-Б Право + ГО", "9-Б ГО": "9-Б Право + ГО"},
                             courses=["9-Б Право + ГО", "8-Б ІУ"])
        editor.cfg["days"]["0"][0] = ["8-Б ІУ", None]
        self.select(editor, "9-Б ГО")
        editor.edit_course_schedule()
        api = editor.winfo_children()[-1].api
        api.top.set("9-Б Право")
        api.bottom.set("9-Б ГО")
        with mock.patch.object(editor_ui.messagebox, "askyesno", return_value=False) as ask:
            api.assign()
        self.assertTrue(ask.called)
        self.assertEqual(editor.cfg["days"]["0"][0], ["8-Б ІУ", None])               # не затерли чужий урок

    def test_delete_course_removes_every_subject_of_it(self):
        editor = self.combined()
        self.select(editor, "9-Б ГО")
        editor.delete_course()
        self.assertEqual(editor.cfg["course_map"], {})


if __name__ == "__main__":
    unittest.main()
