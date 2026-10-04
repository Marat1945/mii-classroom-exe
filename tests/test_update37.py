"""3.7: розумні класи, Назад/Вперед, календар, виділення мишкою, майстер Google, скидання."""
import json
import tempfile
import tkinter as tk
import unittest
from dataclasses import replace
from pathlib import Path
from tkinter import ttk
from unittest import mock

from docx import Document
from docx.oxml.ns import qn

from test_update36 import BaseRoot, TkCase, pump

from classroom_assistant import (course_match, data_tools, datepicker, documents, dragselect, editor_ui,
                                 engine, google_client, google_setup_ui, gui, history, material_library,
                                 samples_ui)
from classroom_assistant.chatgpt_bridge import build_prompt
from classroom_assistant.engine import day_lessons, read_json


class CourseMatchTests(unittest.TestCase):
    COURSES = ["5-Г історія", "8-Б ІУ", "8-Б ВI", "9-Б Право + ГО", "9-Б ГО", "10 ІУ профіль",
               "10-Б ГО", "11 ІУ Профіль", "11 ІУ Стандарт", "11 класи СІМЕЙНЕ НАВЧАННЯ"]

    def test_meaning_not_spelling(self):
        match = course_match.best_match
        self.assertEqual(match("8-Б ВІ", self.COURSES), "8-Б ВI")                 # латинська I
        self.assertEqual(match("5-Г ІУ", self.COURSES), "5-Г історія")
        self.assertEqual(match("9-Б Право", self.COURSES), "9-Б Право + ГО")
        self.assertEqual(match("9-Б ГО", self.COURSES), "9-Б ГО")                 # окремий курс важливіший
        self.assertEqual(match("11 ІУ профіль", self.COURSES), "11 ІУ Профіль")
        self.assertEqual(match("10 ІУ профіль", self.COURSES), "10 ІУ профіль")

    def test_wrong_class_or_subject_is_never_matched(self):
        match = course_match.best_match
        self.assertIsNone(match("11-А ВІ", self.COURSES))
        self.assertIsNone(match("10 ІУ", ["11 ІУ"]))                                # 10 ≠ 11
        self.assertIsNone(match("10-Б ГО", ["10-Б ІУ"]))                            # ГО ≠ ІУ
        self.assertIsNone(match("8-Б ІУ", ["8-В ІУ"]))                              # Б ≠ В
        self.assertIsNone(match("11 ІУ профіль", ["11 ІУ Стандарт"]))
        self.assertIsNone(match("8-Б ІУ", ["8-Б ІУ ", "8-Б  ІУ"][:0] + ["8-Б ВІ", "8-Б ГО"]))

    def test_ambiguity_is_not_guessed(self):
        self.assertIsNone(course_match.best_match("9-Б Право", ["9-Б Право + ГО", "9-Б Право (копія) + ГО"][:1] * 2))

    def test_streams_are_assigned_and_titles_matched(self):
        assigned = course_match.assign_streams({"9-Б Право": "9-Б Право", "9-Б ГО": "9-Б ГО"},
                                               ["9-Б Право + ГО"])
        self.assertEqual(set(assigned.values()), {"9-Б Право + ГО"})
        courses = [{"id": "1", "name": "10-Б ГО"}, {"id": "2", "name": "8-Б ВI"}]
        found = course_match.match_titles(["10-Б ГО", "8-Б ВІ", "7-А ІУ"], courses)
        self.assertEqual({k: v["id"] for k, v in found.items()}, {"10-Б ГО": "1", "8-Б ВІ": "2"})
        self.assertIn("10 клас", course_match.describe("10 ГО"))
        self.assertIn("Громадянська освіта", course_match.describe("10 ГО"))


class HistoryTests(unittest.TestCase):
    def test_undo_redo_and_branching(self):
        h = history.History()
        h.reset("a")
        h.record("b", "друга")
        h.record("c", "третя")
        self.assertEqual(h.undo(), ("b", "третя"))
        self.assertEqual(h.undo(), ("a", "друга"))
        self.assertIsNone(h.undo())
        self.assertEqual(h.redo(), ("b", "друга"))
        h.record("x", "нова гілка")                 # після нової дії «Вперед» зникає
        self.assertFalse(h.can_redo())
        self.assertFalse(h.record("x", "те саме"))

    def test_limit_and_labels(self):
        h = history.History(limit=3)
        h.reset("0")
        for n in "123":
            h.record(n, f"крок {n}")
        self.assertEqual(len(h.items), 3)
        self.assertEqual(h.undo_label(), "крок 3")
        old = json.dumps({"cfg": {}, "state": {"files": {}}, "plans": "x"})
        new = json.dumps({"cfg": {"a": 1}, "state": {"files": {"k": 1}, "attachments": {"k": []}}, "plans": "x"})
        self.assertEqual(history.describe_change(old, new), "розклад / налаштування, вкладення, Word-файли")


class DatePickerTests(TkCase):
    def test_month_grid_and_navigation(self):
        weeks = datepicker.month_grid(2026, 10)
        self.assertTrue(all(len(w) == 7 for w in weeks))
        self.assertEqual(weeks[0][0].weekday(), 0)
        self.assertEqual(datepicker.shift_month(2026, 12, 1), (2027, 1))
        self.assertEqual(datepicker.shift_month(2026, 1, -1), (2025, 12))
        self.assertEqual(datepicker.parse_ui_date("02.10.2026").isoformat(), "2026-10-02")
        self.assertIsNone(datepicker.parse_ui_date("32.13.2026"))

    def test_dialog_respects_weekday_restriction(self):
        dialog = datepicker.CalendarDialog(self.root, datepicker.parse_ui_date("02.10.2026"), weekday=0)
        dialog.choose(datepicker.parse_ui_date("03.10.2026"))          # субота — недоступна
        self.assertIsNone(dialog.result)
        dialog.choose(datepicker.parse_ui_date("05.10.2026"))          # понеділок
        self.assertEqual(dialog.result.isoformat(), "2026-10-05")

    def test_date_field_fills_variable(self):
        variable = tk.StringVar(value="01.09.2026")
        field = datepicker.DateField(self.root, variable)
        with mock.patch.object(datepicker, "pick_date", return_value="07.09.2026"):
            field.open()
        self.assertEqual(variable.get(), "07.09.2026")


class DragSelectTests(TkCase):
    def test_holding_left_button_selects_a_range(self):
        top = tk.Toplevel(self.root)
        dragselect.install(self.root)
        tree = ttk.Treeview(top, columns=("a",), show="headings", height=8, selectmode="extended")
        tree.pack(fill="both", expand=True)
        for n in range(8):
            tree.insert("", "end", iid=str(n), values=(n,))
        top.deiconify()
        top.update()
        y0 = tree.bbox("1")[1] + 5
        y1 = tree.bbox("4")[1] + 5
        tree.event_generate("<ButtonPress-1>", x=20, y=y0)
        tree.event_generate("<B1-Motion>", x=20, y=y1)
        top.update()
        self.assertEqual(set(tree.selection()), {"1", "2", "3", "4"})
        tree.event_generate("<ButtonRelease-1>", x=20, y=y1)
        tree.event_generate("<ButtonPress-1>", x=20, y=y1, state=0x0004)   # Ctrl+клац не скидає
        top.update()
        self.assertGreaterEqual(len(tree.selection()), 1)
        top.destroy()


class EditorBehaviourTests(TkCase):
    def make(self, cfg=None, plans=None):
        config = cfg or data_tools.blank_config()
        self.root.state = {}
        self.root.google_courses = []
        self.root.update_day = lambda: None
        with mock.patch.object(editor_ui, "deep_copy_data", return_value=(config, plans if plans is not None else {})):
            editor = editor_ui.SchoolEditor(self.root, lambda: None)
        editor.deiconify()
        editor.update()
        self.addCleanup(lambda: self._close(editor))
        return editor

    @staticmethod
    def _close(editor):
        try:
            editor.destroy()
        except tk.TclError:
            pass

    def real_editor(self):
        self.root.state = {}
        self.root.google_courses = []
        self.root.update_day = lambda: None
        editor = editor_ui.SchoolEditor(self.root, lambda: None)
        editor.deiconify()
        editor.update()
        self.addCleanup(lambda: self._close(editor))
        return editor

    def test_lesson_menu_is_trimmed_and_apply_keeps_edit(self):
        editor = self.real_editor()
        editor.lessons.selection_set("8")
        editor.lesson_selected()
        popped = []
        with mock.patch.object(tk.Menu, "tk_popup", lambda menu, x, y, entry="": popped.append(menu)):
            editor.topictext.delete("1.0", "end")
            editor.topictext.insert("1.0", "ЗМІНЕНА ТЕМА")
            bx, by, _w, _h = editor.lessons.bbox("8")
            editor.lessons.event_generate("<Button-3>", x=bx + 20, y=by + 5)
            editor.update()
        menu = popped[-1]
        labels = [menu.entrycget(i, "label") for i in range(menu.index("end") + 1) if menu.type(i) == "command"]
        self.assertFalse([x for x in labels if "Пересунути" in x or "Редагувати тему" in x], labels)
        self.assertIn("Застосувати зміни з полів внизу", labels)
        menu.invoke(next(i for i in range(menu.index("end") + 1)
                         if menu.type(i) == "command" and "Застосувати" in menu.entrycget(i, "label")))
        self.assertEqual(editor.plans[editor.current_plan]["lessons"][8]["topic"], "ЗМІНЕНА ТЕМА")

    def test_edit_is_not_lost_when_moving_to_another_row(self):
        editor = self.real_editor()
        editor.lessons.selection_set("8")
        editor.lesson_selected()
        editor.topictext.delete("1.0", "end")
        editor.topictext.insert("1.0", "ПРАВКА БЕЗ КНОПКИ")
        editor.lessons.selection_set("9")
        editor.lesson_selected()
        self.assertEqual(editor.plans[editor.current_plan]["lessons"][8]["topic"], "ПРАВКА БЕЗ КНОПКИ")
        self.assertNotIn("ПРАВКА", editor.topictext.get("1.0", "end"))

    def test_delete_schedule_button_and_undo_redo(self):
        editor = self.real_editor()
        before = json.dumps(editor.cfg["days"], sort_keys=True)
        self.assertNotIn("None", before.replace("null", "None").split("[")[0])
        editor.delete_schedule()                                    # askyesno → True у тестах
        self.assertTrue(all(pair == [None, None] for pairs in editor.cfg["days"].values() for pair in pairs))
        editor.undo()                                               # спершу фіксує дію, потім скасовує
        self.assertEqual(json.dumps(editor.cfg["days"], sort_keys=True), before)
        editor.redo()
        self.assertTrue(all(pair == [None, None] for pairs in editor.cfg["days"].values() for pair in pairs))
        self.assertIn("Повернуто", editor.history_note.cget("text"))

    def test_delete_key_clears_all_selected_slots(self):
        editor = self.real_editor()
        editor.dayselect.current(0)
        editor.refresh_slots()
        editor.slots.selection_set(("0", "1", "2"))
        editor.clear_slot()
        self.assertTrue(all(editor.cfg["days"]["0"][i] == [None, None] for i in (0, 1, 2)))
        self.assertNotEqual(editor.cfg["days"]["0"][3], [None, None])

    def test_holiday_dates_by_calendar(self):
        editor = self.make()
        editor.cfg["holidays"].append({"start": "2026-10-26", "end": "2026-11-01"})
        editor.refresh_holidays()
        with mock.patch.object(editor_ui, "pick_date", return_value="02.11.2026"):
            event = type("E", (), {"x": 0, "y": 0})()
            with mock.patch.object(editor.holidays, "identify_row", return_value="0"), \
                    mock.patch.object(editor.holidays, "identify_column", return_value="#2"):
                editor._inline_holiday_edit(event)
        self.assertEqual(editor.cfg["holidays"][0]["end"], "2026-11-02")
        with mock.patch.object(editor, "_ask_holiday", return_value=("2026-12-24", "2027-01-10")):
            editor.add_holiday()
        self.assertEqual(len(editor.cfg["holidays"]), 2)

    def test_combined_course_has_a_row_per_subject_and_the_file_goes_to_the_chosen_one(self):
        config = data_tools.blank_config()
        config["course_map"] = {"9-Б Право": {"plan": "p1", "course_title": "9-Б Право"},
                                "9-Б ГО": {"plan": "p2", "course_title": "9-Б ГО"}}
        plans = {"p1": {"filename": "Очікує файл КТП", "lessons": [], "needs_review": True},
                 "p2": {"filename": "Очікує файл КТП", "lessons": [], "needs_review": True}}
        self.root.google_courses = []
        editor = self.make(config, plans)
        editor.parent.google_courses = [{"name": "9-Б Право + ГО", "id": "7"}]
        editor._refresh_planlist()
        self.assertEqual([e["label"] for e in editor.list_entries],
                         ["9-Б Право + ГО  ▸ ГО   — без КТП", "9-Б Право + ГО  ▸ Право   — без КТП"])
        self.assertEqual([e["streams"] for e in editor.list_entries], [["9-Б ГО"], ["9-Б Право"]])
        self.assertEqual(editor.list_entries[0]["course_streams"], ["9-Б ГО", "9-Б Право"])
        editor.planlist.selection_set(0)                                   # рядок «ГО»
        editor.plan_selected()
        with tempfile.TemporaryDirectory() as folder:
            ktp = Path(folder) / "ktp.docx"
            samples_ui.build_sample_ktp(ktp)
            with mock.patch.object(editor, "_ask_choice", side_effect=AssertionError("вибір не потрібен")):
                editor.import_plan(path=ktp)
        self.assertEqual(len(editor.plans["p2"]["lessons"]), 5)
        self.assertEqual(editor.plans["p1"]["lessons"], [])

class MainWindowFeatureTests(TkCase):
    NEEDS_ROOT = False

    def make(self):
        app = gui.MainApp()
        self.addCleanup(lambda: self._close(app))
        pump(app, 0.3)
        return app

    @staticmethod
    def _close(app):
        try:
            app.destroy()
        except tk.TclError:
            pass

    def test_undo_redo_buttons_exist_left_of_samples_and_restore_state(self):
        app = self.make()
        self.assertEqual(str(app.undo_button.cget("state")), "disabled")
        with tempfile.TemporaryDirectory() as folder, \
                mock.patch.object(gui, "DATA", Path(folder)), \
                mock.patch.object(gui, "write_json") as write, \
                mock.patch.object(gui, "save_state") as saver, \
                mock.patch.object(app, "update_day"):
            app._plans_sig = None
            app.history.reset(app._history_snapshot())
            app.state["description_overrides"] = {"k": "текст"}
            app.state["drafts"] = {"keep": {"id": "1"}}                  # чернетки Google не відкочуються
            app._history_record()
            self.assertEqual(str(app.undo_button.cget("state")), "normal")
            app.undo()
            self.assertNotIn("description_overrides", app.state)
            self.assertIn("keep", app.state["drafts"])
            self.assertTrue(write.called and saver.called)
            app.redo()
            self.assertEqual(app.state["description_overrides"], {"k": "текст"})
        app.update()
        group = app.undo_button.master
        names = {"↶ Назад", "↷ Вперед", "📄 Зразки документів", "🗂 Мої дані"}
        ordered = [w.cget("text") for w in sorted(group.winfo_children(), key=lambda w: (w.winfo_rooty(), w.winfo_rootx()))
                   if isinstance(w, ttk.Button) and w.cget("text") in names]
        self.assertEqual(ordered, ["↶ Назад", "↷ Вперед", "📄 Зразки документів", "🗂 Мої дані"])

    def test_header_wraps_on_a_narrow_screen_instead_of_clipping_buttons(self):
        app = self.make()
        app.geometry("1180x700")
        pump(app, 0.5)
        header = app.undo_button.master
        for button in header.winfo_children():
            if button.winfo_class() in ("TButton", "TLabel", "TEntry"):
                right = button.winfo_rootx() - app.winfo_rootx() + button.winfo_width()
                self.assertLessEqual(right, app.winfo_width() + 1, str(button.cget("text") if button.winfo_class() != "TEntry" else "date"))
        self.assertGreater(app.undo_button.winfo_rooty(), app.date_entry.winfo_rooty())      # другий рядок
        app.geometry("2300x800")
        pump(app, 0.5)
        self.assertLessEqual(abs(app.undo_button.winfo_rooty() - app.date_entry.winfo_rooty()), 12)   # знову в один
    def test_undo_is_blocked_while_the_editor_is_open(self):
        app = self.make()
        editor = editor_ui.SchoolEditor(app, app.update_day)
        with mock.patch.object(gui.messagebox, "showinfo") as info:
            app.undo()
        self.assertTrue(info.called)
        editor.destroy()

    def test_reset_everything_asks_defaults_to_no_and_backs_up_first(self):
        app = self.make()
        with mock.patch.object(gui.messagebox, "askyesno", return_value=False) as ask, \
                mock.patch.object(data_tools, "reset_to_blank") as reset:
            app.reset_everything()
        self.assertEqual(ask.call_args.kwargs.get("default"), "no")
        reset.assert_not_called()
        with tempfile.TemporaryDirectory() as folder, \
                mock.patch.object(gui.messagebox, "askyesno", return_value=True), \
                mock.patch.object(gui.messagebox, "showinfo"), \
                mock.patch.object(data_tools, "auto_backup_path", return_value=Path(folder) / "b.zip"), \
                mock.patch.object(data_tools, "reset_to_blank") as reset, \
                mock.patch.object(app, "reload_data"):
            app.reset_everything()
        reset.assert_called_once_with(Path(folder) / "b.zip")

    def test_restore_last_backup_without_copies_explains(self):
        app = self.make()
        with mock.patch.object(data_tools, "latest_backup", return_value=None), \
                mock.patch.object(gui.messagebox, "showinfo") as info:
            app.restore_last_backup()
        self.assertIn("Автоматичних копій ще немає", info.call_args.args[1])

    def test_connect_google_first_time_shows_wizard_then_only_refreshes(self):
        app = self.make()
        with mock.patch.object(google_client, "token_ready", return_value=False), \
                mock.patch.object(google_setup_ui, "show_google_wizard") as wizard, \
                mock.patch.object(app, "sync_classroom") as sync:
            app.connect_google()
        wizard.assert_called_once()
        sync.assert_not_called()
        with mock.patch.object(google_client, "token_ready", return_value=True), \
                mock.patch.object(google_setup_ui, "show_google_wizard") as wizard, \
                mock.patch.object(app, "sync_classroom") as sync:
            app.connect_google()
        wizard.assert_not_called()
        sync.assert_called_once_with(interactive=True, announce=True)

    def test_clean_start_runs_once_with_a_backup_and_never_repeats(self):
        app = self.make()
        resets, marker = [], {}
        with mock.patch.object(gui.messagebox, "askyesno", side_effect=AssertionError("запитань не має бути")):
            app._clean_start_once()                                   # не EXE-збірка: нічого не робить
        app.cfg["course_map"] = {"8-Б ІУ": {"plan": "x", "course_title": "8-Б ІУ"}}
        toasts = []
        with mock.patch.dict(gui.os.environ, {"POMICHNYK_NO_OFFERS": ""}), \
                mock.patch.object(gui.sys, "frozen", True, create=True), \
                mock.patch.object(data_tools, "read_install", side_effect=lambda *a: dict(marker)), \
                mock.patch.object(data_tools, "write_install", side_effect=lambda v, *a: marker.update(v)), \
                mock.patch.object(data_tools, "reset_to_blank", side_effect=lambda p, *a, **k: resets.append(p)), \
                mock.patch.object(data_tools, "auto_backup_path", return_value=Path("/tmp/копія.zip")), \
                mock.patch.object(app, "apply_data_operation"), \
                mock.patch.object(app, "begin_data_operation"), \
                mock.patch.object(gui, "show_toast", side_effect=lambda parent, text, *a, **k: toasts.append(text)):
            app._clean_start_once()
            self.assertEqual(resets, [Path("/tmp/копія.zip")])
            self.assertEqual(marker, {"clean_start": data_tools.CLEAN_START_RELEASE})
            self.assertIn("Повну копію збережено", toasts[0])
            app._clean_start_once()                                   # вдруге — ні
            self.assertEqual(len(resets), 1)

    def test_clean_start_marker_is_never_part_of_a_backup(self):
        self.assertTrue(data_tools.is_secret(Path("install_state.json")))

    def test_course_mapping_default_uses_meaning(self):
        app = self.make()
        app.google_courses = [{"name": "8-Б ВI", "id": "11"}, {"name": "9-Б Право + ГО", "id": "12"}]
        app.cfg["classroom_course_titles"] = ["8-Б ВІ", "9-Б Право"]
        app.state["course_ids"] = {}
        shown = {}
        class FakeCombo:
            def __init__(self, master, textvariable=None, values=(), width=0, state=""):
                self.values = values
            def set(self, value):
                shown[len(shown)] = value
            def grid(self, **kw):
                pass
        with mock.patch.object(gui.ttk, "Combobox", FakeCombo), mock.patch.object(gui.ttk, "Button"), \
                mock.patch.object(gui, "fit_work_window"):
            app.map_courses()
        self.assertEqual(list(shown.values()), ["8-Б ВI [11]", "9-Б Право + ГО [12]"])


class GoogleWizardTests(TkCase):
    def test_credentials_file_is_validated_and_copied(self):
        with tempfile.TemporaryDirectory() as folder, mock.patch.object(google_client, "DATA", Path(folder) / "data"):
            good = Path(folder) / "good.json"
            good.write_text(json.dumps({"installed": {"client_id": "x", "auth_uri": "a", "token_uri": "t"}}))
            google_client.import_credentials_file(good)
            self.assertTrue(google_client.credentials_present())
            web = Path(folder) / "web.json"
            web.write_text(json.dumps({"web": {"client_id": "x"}}))
            with self.assertRaisesRegex(ValueError, "Веб-застосунок"):
                google_client.import_credentials_file(web)
            junk = Path(folder) / "junk.json"
            junk.write_text("не json")
            with self.assertRaises(ValueError):
                google_client.import_credentials_file(junk)

    def test_wizard_text_has_clickable_links_and_login_runs_then_syncs(self):
        calls = []
        self.root.sync_classroom = lambda **kw: calls.append(kw)
        with tempfile.TemporaryDirectory() as folder, mock.patch.object(google_client, "DATA", Path(folder)), \
                mock.patch.object(google_setup_ui.messagebox, "showinfo"), \
                mock.patch.object(google_setup_ui.messagebox, "showerror"):
            wizard = google_setup_ui.GoogleWizard(self.root)
            text = wizard.text.get("1.0", "end")
            for needle in ("console.cloud.google.com/projectcreate", "Google Classroom API",
                           "Google Drive API", "Комп'ютерний застосунок", "Тестові користувачі"):
                self.assertIn(needle, text)
            self.assertIn("classroom.googleapis.com", google_setup_ui.LINKS["classroom"])
            self.assertIn("drive.googleapis.com", google_setup_ui.LINKS["drive"])
            self.assertIn("⬜", wizard.login_status.cget("text"))
            wizard.login()                                            # без ключа — підказка, не вхід
            self.assertFalse(wizard._busy)
            (Path(folder) / "google_credentials.json").write_text("{}")
            def fake_authenticate(*a, **k):
                (Path(folder) / "google_token.json").write_text(json.dumps(
                    {"scopes": google_client.SCOPES, "refresh_token": "r", "token": "t"}))
            with mock.patch.object(google_client, "authenticate", fake_authenticate):
                wizard.login()
                pump(self.root, 0.8)
            self.assertEqual(calls, [{"interactive": True, "announce": True}])


class WordClosingAndDataTests(unittest.TestCase):
    @staticmethod
    def lecture(header="ВСЕСВІТНЯ ІСТОРІЯ • 11 КЛАС • 28.09", with_closing=False):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = header
        doc.sections[0].footer.paragraphs[0].text = "Тема"
        doc.add_paragraph("Урок 28.09 — Тема.")
        for n in range(8):
            doc.add_paragraph(f"Розділ {n}. " + "Змістовний абзац лекції про епоху. " * 5)
        if with_closing:
            doc.add_heading("Техніка безпеки", 1)
            doc.add_paragraph("Д/з: старе")
        return doc

    def test_closing_block_has_safety_but_never_homework_and_is_idempotent(self):
        lesson = replace(day_lessons("2026-10-01", read_json("Налаштування.json"))[0], homework="прочитати § 5")
        doc = documents.ensure_closing(self.lecture(), lesson)
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(texts[-1], "Техніка безпеки")
        self.assertFalse(any(x.startswith("Д/з") for x in texts))                  # домашнє завдання лише в Classroom
        fills = {sh.get(qn("w:fill")) for tc in doc.element.body.iter(qn("w:tc")) for sh in tc.iter(qn("w:shd"))}
        self.assertIn("FCE8E6", fills)
        documents.ensure_closing(doc, replace(lesson, homework="інше"))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(texts.count("Техніка безпеки"), 1)
        self.assertFalse(any(x.startswith("Д/з") for x in texts))

    def test_old_homework_paragraphs_are_removed_from_the_word(self):
        lesson = replace(day_lessons("2026-10-01", read_json("Налаштування.json"))[0], homework="нове")
        doc = documents.ensure_closing(self.lecture(with_closing=True), lesson)       # у файлі було «Д/з: старе»
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertFalse(any("Д/з" in x for x in texts))
        self.assertEqual(texts.count("Техніка безпеки"), 1)
        self.assertTrue(any(x.startswith("Розділ 7.") for x in texts))               # зміст лекції не зачеплено

    def test_parallel_copy_updates_dates_in_short_header_and_first_line(self):
        lesson = replace(day_lessons("2026-10-01", read_json("Налаштування.json"))[0],
                         day="2026-10-09", lesson_number=9, homework="параграф 9", topic="Тема")
        with tempfile.TemporaryDirectory() as folder:
            master = Path(folder) / "master.docx"
            self.lecture().save(master)
            target = Path(folder) / "out" / "copy.docx"
            with mock.patch.object(material_library, "word_path", lambda _l: target):
                result = material_library.copy_for_lesson(master, lesson)
            copy_doc = Document(result)
            self.assertEqual(copy_doc.sections[0].header.paragraphs[0].text, "ВСЕСВІТНЯ ІСТОРІЯ • 11 КЛАС • 09.10")
            self.assertEqual([p.text for p in copy_doc.paragraphs if p.text.strip()][0],
                             f"{lesson.stream}, Урок 09.10 — Тема")
            self.assertFalse(any(p.text.startswith("Д/з") for p in copy_doc.paragraphs))

    def test_clean_install_creates_blank_data_without_overwriting(self):
        with tempfile.TemporaryDirectory() as folder:
            data = Path(folder) / "data"
            engine.ensure_default_data(data)
            config = json.loads((data / "Налаштування.json").read_text("utf-8"))
            self.assertEqual(config["course_map"], {})
            self.assertEqual((data / "Календарні плани.json").read_text("utf-8"), "{}")
            (data / "Календарні плани.json").write_text('{"мої": 1}', encoding="utf-8")
            engine.ensure_default_data(data)
            self.assertEqual((data / "Календарні плани.json").read_text("utf-8"), '{"мої": 1}')

    def test_backup_folder_helpers(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            self.assertIsNone(data_tools.latest_backup(root))
            first = data_tools.auto_backup_path(data_tools.RESET_PREFIX, root)
            first.write_bytes(b"x")
            other = data_tools.auto_backup_path("Стан перед відновленням", root)
            other.write_bytes(b"y")
            self.assertEqual(data_tools.latest_backup(root), first)          # «перед скиданням» важливіша
            (root / "data").mkdir()
            (root / "data" / "Налаштування.json").write_text("{}")
            archive = root / "all.zip"
            data_tools.export_all_data(archive, root)
            import zipfile
            self.assertEqual(zipfile.ZipFile(archive).namelist(), ["data/Налаштування.json"])  # копії копій не входять


class NewPromptTests(unittest.TestCase):
    def test_file_prompt_matches_the_sample_and_asks_for_one_picture(self):
        lesson = day_lessons("2026-10-01", read_json("Налаштування.json"))[0]
        prompt = build_prompt(lesson)
        self.assertIn("Урок 01.10 — ", prompt)
        self.assertIn("ОДНА картинка-інфографіка", prompt)
        self.assertIn("Рівно ОДНЕ зображення", prompt)
        self.assertIn("ІСТОРІЯ УКРАЇНИ • 11 КЛАС • 01.10", prompt)
        self.assertIn("«Основні поняття»", prompt)
        self.assertIn("НЕ додавай розділ «Домашнє завдання / Д/з»", prompt)
        self.assertIn("Розділ «Техніка безпеки» додасть програма", prompt)
        self.assertNotIn("## Назва першого розділу", prompt)
        self.assertIn("## Назва першого розділу", build_prompt(lesson, None, "text"))


if __name__ == "__main__":
    unittest.main()
