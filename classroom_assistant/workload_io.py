"""«Навантаження» вчителя: Word і JPEG у вигляді вашого зразка.

Зразок: A4 книжкова, суцільна таблиця з рамками; перший рядок — заголовок
«НАВАНТАЖЕННЯ (прізвище) // канікули … // код»; далі «№ / дні тижня»; у першій
колонці номер уроку жирним і час дрібніше; червоний рядок «ХАРЧУВАННЯ У ЇДАЛЬНІ»;
у клітинках «8-Б ІУ / ГО» (чисельник / знаменник), «/ 9-Г Право», «11 ІУ стандарт /».
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

DAYS = ("Понеділок", "Вівторок", "Середа", "Четвер", "П’ятниця")
COLUMN_TWIPS = (1142, 1980, 2066, 1894, 1980, 1980)     # як у зразку
TABLE_TWIPS = sum(COLUMN_TWIPS)
HEIGHT_TITLE, HEIGHT_HEADER, HEIGHT_ROW, HEIGHT_MEAL = 560, 890, 680, 269
CLASS_TOKEN = re.compile(r"^\d{1,2}[-–][А-ЯІЇЄҐA-Z]$")
DEFAULT_MEAL = "ХАРЧУВАННЯ У ЇДАЛЬНІ"


# ---------- спільна логіка ----------
RANGE = re.compile(r"(\d{2})\.(\d{2})\s*[-–]\s*(\d{2})\.(\d{2})")


def holidays_from_title(text, year_start) -> list:
    """«канікули 26.10-01.11, 24.12-10.01» → [{'start': ISO, 'end': ISO}], роки за навчальним роком."""
    low = str(text).casefold()
    at = low.find("канікул")
    if at < 0:
        return []
    base = int(str(year_start)[:4])
    result = []
    for d1, m1, d2, m2 in RANGE.findall(str(text)[at:]):
        try:
            def iso(day, month):
                y = base if int(month) >= 8 else base + 1
                return date(y, int(month), int(day)).isoformat()
            start, end = iso(d1, m1), iso(d2, m2)
        except ValueError:
            continue
        if end >= start:
            result.append({"start": start, "end": end})
    return result


def holidays_text(config) -> str:
    parts = []
    for item in sorted(config.get("holidays", []), key=lambda h: str(h.get("start", ""))):
        try:
            start = date.fromisoformat(item["start"])
            end = date.fromisoformat(item["end"])
        except (KeyError, ValueError, TypeError):
            continue
        parts.append(f"{start:%d.%m}-{end:%d.%m}")
    return ", ".join(parts)


def title_parts(config):
    """[(текст, жирний, pt)] — як у зразку: 16 pt назва, 10 pt решта."""
    teacher = str(config.get("workload_teacher", "")).strip()
    code = str(config.get("workload_code", "")).strip()
    parts = [("НАВАНТАЖЕННЯ" + (f" ({teacher})" if teacher else "") + " ", True, 16)]
    holidays = [x.strip() for x in holidays_text(config).split(",") if x.strip()]
    if holidays:
        parts.append(("//  канікули  ", False, 10))
        for i, text in enumerate(holidays):
            if i:
                parts.append((", ", False, 10))
            parts.append((text, True, 10))
    if code:
        parts.append((" // " if holidays else "// ", False, 10))
        parts.append((code, True, 10))
    return parts


def pair_text(first, second) -> str:
    """Чисельник/знаменник як у зразку: «8-Б ІУ / ГО», «/ 9-Г Право», «11 ІУ профіль /»."""
    first, second = first or "", second or ""
    if not first and not second:
        return ""
    if first == second:
        return first
    if first and second:
        head_a, head_b = first.split()[0], second.split()[0]
        if head_a == head_b and CLASS_TOKEN.match(head_a):
            return f"{first} / {second[len(head_b):].strip()}"
        return f"{first} / {second}"
    return f"{first} /" if first else f"/ {second}"


def period_label_times(config, index) -> str:
    try:
        start, end = config["period_times"][index]
    except (IndexError, KeyError, ValueError):
        return ""
    return f"{start}-{end}".replace(":", ".")


def row_count(config) -> int:
    return len(config.get("period_times", []))


def parse_workload_title(rows, year_start=None) -> dict:
    """З імпортованої таблиці дістає прізвище та код із першого рядка-заголовка."""
    for row in rows[:3]:
        text = " ".join(dict.fromkeys(str(c) for c in row if str(c).strip()))     # об'єднані клітинки
        if not text.upper().lstrip().startswith("НАВАНТАЖЕННЯ"):
            continue
        info = {}
        teacher = re.search(r"\(([^)]+)\)", text)
        if teacher:
            info["workload_teacher"] = teacher.group(1).strip()
        tail = re.search(r"//\s*([0-9A-Za-zА-Яа-я\-]{3,})\s*$", text)
        if tail and not re.search(r"\d{2}\.\d{2}", tail.group(1)):
            info["workload_code"] = tail.group(1)
        if year_start:
            found = holidays_from_title(text, year_start)
            if found:
                info["holidays"] = found
        return info
    return {}


# ---------- Word ----------
def export_workload_docx(config, path, show_numbers=True, show_times=True, show_meal=True):
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Twips

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    side = (11906 - TABLE_TWIPS) // 2
    section.left_margin = section.right_margin = Twips(side)
    section.top_margin = section.bottom_margin = Twips(426)
    section.header_distance = section.footer_distance = Twips(0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    table = doc.add_table(rows=0, cols=6)
    table.autofit = False
    properties = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    properties.append(borders)
    for column, width in zip(table.columns, COLUMN_TWIPS):
        column.width = Twips(width)

    def set_height(row, twips):
        row_properties = row._tr.get_or_add_trPr()
        height = OxmlElement("w:trHeight")
        height.set(qn("w:val"), str(twips))
        height.set(qn("w:hRule"), "atLeast")
        row_properties.append(height)

    def write(cell, runs, align=WD_ALIGN_PARAGRAPH.CENTER, new_paragraph=False):
        paragraph = cell.add_paragraph() if new_paragraph else cell.paragraphs[0]
        paragraph.alignment = align
        for text, bold, size, color in runs:
            run = paragraph.add_run(text)
            run.bold = bold or None
            run.font.size = Pt(size)
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        return paragraph

    def merged_row(height):
        row = table.add_row()
        set_height(row, height)
        cell = row.cells[0].merge(row.cells[5])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        return cell

    write(merged_row(HEIGHT_TITLE), [(t, b, s, None) for t, b, s in title_parts(config)])

    header = table.add_row()
    set_height(header, HEIGHT_HEADER)
    for index, label in enumerate(("№",) + DAYS):
        cell = header.cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        write(cell, [(label, True, 14, None)])

    meal_after = int(config.get("meal_break_after", 2) or 0)
    count = row_count(config)
    for number in range(1, count + 1):
        row = table.add_row()
        set_height(row, HEIGHT_ROW)
        first = row.cells[0]
        first.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if show_numbers:
            write(first, [(str(number), True, 11, None)])
        if show_times:
            write(first, [(period_label_times(config, number - 1), False, 9, None)],
                  new_paragraph=show_numbers)
        for day in range(5):
            pairs = config["days"].get(str(day), [])
            a, b = pairs[number - 1] if len(pairs) >= number else (None, None)
            cell = row.cells[day + 1]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            write(cell, [(pair_text(a, b), False, 11, None)])
        if show_meal and number == meal_after and number < count:
            write(merged_row(HEIGHT_MEAL),
                  [(config.get("meal_label") or DEFAULT_MEAL, True, 11, "FF0000")])
    for row in table.rows:                       # ширина клітинок = ширина колонок
        for cell, width in zip(row.cells, COLUMN_TWIPS):
            if cell._tc.tcPr is None or cell._tc.tcPr.find(qn("w:gridSpan")) is None:
                cell.width = Twips(width)
    doc.save(str(path))
    return Path(path)


# ---------- JPEG ----------
def _font(size, bold=False):
    from PIL import ImageFont
    names = (["C:/Windows/Fonts/timesbd.ttf",
              "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"] if bold else
             ["C:/Windows/Fonts/times.ttf",
              "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"])
    for name in names:
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def export_workload_jpeg(config, path, show_numbers=True, show_times=True, show_meal=True):
    from PIL import Image, ImageDraw
    unit = 0.1256                                  # пікселів на 1 twip (як у зразку JPG)
    px = lambda twips: int(round(twips * unit))
    pt = lambda size: int(round(size * 2.51))
    margin = 14
    widths = [px(w) for w in COLUMN_TWIPS]
    table_w = sum(widths)
    count = row_count(config)
    meal_after = int(config.get("meal_break_after", 2) or 0)
    has_meal = show_meal and 0 < meal_after < count
    height = (px(HEIGHT_TITLE) + px(HEIGHT_HEADER) + count * px(HEIGHT_ROW)
              + (px(HEIGHT_MEAL) if has_meal else 0) + 2 * margin)
    image = Image.new("RGB", (table_w + 2 * margin, height), "white")
    draw = ImageDraw.Draw(image)
    ink = "black"
    x0, y = margin, margin

    def box(x, top, w, h):
        draw.rectangle((x, top, x + w, top + h), outline=ink, width=2)

    def centered(text, font, cx, cy, fill=ink):
        draw.text((cx, cy), text, font=font, fill=fill, anchor="mm")

    def fit(text, bold, size, max_width):
        font = _font(size, bold)
        while size > 14 and draw.textlength(text, font=font) > max_width:
            size -= 1
            font = _font(size, bold)
        return font

    # заголовок: складений рядок з різними розмірами
    box(x0, y, table_w, px(HEIGHT_TITLE))
    parts = [(t, _font(pt(s), b)) for t, b, s in title_parts(config)]
    total = sum(draw.textlength(t, font=f) for t, f in parts)
    cursor = x0 + (table_w - total) / 2
    base = y + px(HEIGHT_TITLE) / 2
    for text, font in parts:
        draw.text((cursor, base), text, font=font, fill=ink, anchor="lm")
        cursor += draw.textlength(text, font=font)
    y += px(HEIGHT_TITLE)

    # шапка днів
    x = x0
    for width, label in zip(widths, ("№",) + DAYS):
        box(x, y, width, px(HEIGHT_HEADER))
        centered(label, _font(pt(14), True), x + width / 2, y + px(HEIGHT_HEADER) / 2)
        x += width
    y += px(HEIGHT_HEADER)

    for number in range(1, count + 1):
        row_h = px(HEIGHT_ROW)
        x = x0
        for column, width in enumerate(widths):
            box(x, y, width, row_h)
            cx, cy = x + width / 2, y + row_h / 2
            if column == 0:
                time_text = period_label_times(config, number - 1) if show_times else ""
                if show_numbers and time_text:
                    centered(str(number), _font(pt(11), True), cx, cy - 15)
                    centered(time_text, _font(pt(9)), cx, cy + 17)
                elif show_numbers:
                    centered(str(number), _font(pt(11), True), cx, cy)
                elif time_text:
                    centered(time_text, _font(pt(9)), cx, cy)
            else:
                pairs = config["days"].get(str(column - 1), [])
                a, b = pairs[number - 1] if len(pairs) >= number else (None, None)
                text = pair_text(a, b)
                if text:
                    centered(text, fit(text, False, pt(11), width - 12), cx, cy)
            x += width
        y += row_h
        if has_meal and number == meal_after:
            box(x0, y, table_w, px(HEIGHT_MEAL))
            centered(config.get("meal_label") or DEFAULT_MEAL, _font(pt(11), True),
                     x0 + table_w / 2, y + px(HEIGHT_MEAL) / 2, "#FF0000")
            y += px(HEIGHT_MEAL)
    image.save(str(path), "JPEG", quality=95, subsampling=0, dpi=(150, 150))
    return Path(path)
