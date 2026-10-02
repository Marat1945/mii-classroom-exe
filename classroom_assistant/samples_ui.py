"""Зразки документів, які програма розуміє, промти для ШІ та папка з даними."""
from __future__ import annotations

import os
import subprocess
import sys
import tkinter as tk
import zipfile
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText

from .engine import ROOT
from .window_ui import fit_work_window

WORKLOAD_PROMPT = """Створи документ Microsoft Word (.docx) — «Навантаження» (розклад уроків учителя).

ФОРМАТ (суворо):
• Аркуш А4, книжкова орієнтація. У документі ЛИШЕ ОДНА таблиця з 6 колонками, без інших абзаців, колонтитулів і картинок.
• Шрифт Times New Roman. Усі клітинки по центру (по горизонталі й вертикалі). Межі — чорні, тонкі (0,5 пт).
• Ширина колонок (см): 2,0 / 3,5 / 3,6 / 3,3 / 3,5 / 3,5.
• РЯДОК 1 (усі 6 клітинок об'єднано): «НАВАНТАЖЕННЯ (Прізвище І.О.) // канікули ДД.ММ-ДД.ММ, ДД.ММ-ДД.ММ // код». Слово «НАВАНТАЖЕННЯ» і прізвище — жирно 16 пт; решта — 10 пт, дати канікул і код — жирно.
• РЯДОК 2: «№», «Понеділок», «Вівторок», «Середа», «Четвер», «П’ятниця» — жирно 14 пт.
• Далі ПО ОДНОМУ РЯДКУ НА КОЖЕН УРОК (1, 2, 3 …). У першій клітинці: номер уроку жирно 11 пт, а під ним час «08.30-09.15» (9 пт, формат ГГ.ХХ-ГГ.ХХ). В інших клітинках — потік: клас і предмет, як у мене: «8-Б ІУ», «11 ІУ профіль», «9-Г Право».
• Чисельник і знаменник: якщо уроки різні, пиши «чисельник / знаменник», наприклад «8-Б ІУ / ГО» (коли клас той самий, у другій частині достатньо назви предмета). Лише знаменник: «/ 9-Г Право». Лише чисельник: «11 ІУ стандарт /». Якщо урок щотижня однаковий — пиши назву один раз: «8-Б ВІ».
• Порожня клітинка = уроку немає.
• Після 2-го уроку окремий рядок на всю ширину: «ХАРЧУВАННЯ У ЇДАЛЬНІ» — червоним жирним 11 пт.
• Нічого не вигадуй і не змінюй: використай лише ті дані, які я надаю нижче.

МОЇ ДАНІ:
[Вставте сюди свій розклад: для кожного дня тижня — потік на кожному уроці (з поділом на чисельник/знаменник), час дзвоників, канікули, прізвище та код.]
"""

KTP_PROMPT = """Створи документ Microsoft Word (.docx) — календарно-тематичне планування (КТП), яке потім завантажується в програму «Помічник учителя Classroom».

ФОРМАТ (суворо):
• Аркуш А4. У документі ОДНА таблиця з 4 колонками. Перший рядок — шапка: «№ з/п» | «Дата» | «Тема уроку» | «Домашнє завдання».
• Один урок = один рядок. У першій колонці — лише число уроку: 1, 2, 3 … (подвійний урок з однією темою можна записати «46-47»). НЕ користуйся автоматичною нумерацією Word — впиши числа текстом.
• Рядки назв розділів («Розділ 1. …») — окремим рядком на всю ширину, БЕЗ номера.
• Колонка «Дата»: для кожного класу дата проведення у форматі ДД.ММ. — наприклад «8-Б 04.09. 8-В 04.09. 8-Г 03.09.». Якщо клас один або дати однакові — достатньо «04.09.». Якщо дат ще немає — залиш порожньою.
• Колонка «Тема уроку»: повне формулювання теми без скорочень і без номерів параграфів.
• Колонка «Домашнє завдання»: коротко, як записують у Classroom: «прочитати § 5, повторити § 4». Якщо домашнього завдання немає — порожньо.
• Не додавай нічого поза таблицею (титулів, підписів, приміток). Не вигадуй тем: використай лише мій текст нижче.

МОЇ ДАНІ:
[Вставте сюди перелік тем уроків за порядком, класи, дати й домашні завдання.]
"""


def sample_workload_config() -> dict:
    row = lambda *cells: [list(c) if isinstance(c, tuple) else [c, c] for c in cells]
    days = {
        "0": row(("8-Б ІУ", "8-Б ГО"), ("8-В ІУ", "8-В ГО"), "5-А історія", "9-Б ВІ", "9-Б ВІ",
                 (None, "9-Б Право"), (None, None), (None, None)),
        "1": row("8-Б ВІ", "8-В ВІ", "8-Г ВІ", "11 ІУ профіль", "10-Б ГО", "10-Б ГО", (None, None), (None, None)),
        "2": row("11-А ВІ", "10 ІУ профіль", "9-Б ІУ", (None, None), (None, None),
                 ("9-Б Право", "9-Б ГО"), (None, None), (None, None)),
        "3": row("11 ІУ профіль", "10 ІУ профіль", "8-Г ІУ", "11 ІУ стандарт", "10-Б ВІ",
                 (None, None), ("9-Г ІУ", "9-Г ГО"), (None, None)),
        "4": row("8-В ІУ", "8-Б ІУ", "10 ІУ профіль", "11 ІУ профіль", "9-Г ІУ",
                 ("11 ІУ стандарт", None), "9-Б ІУ", (None, None)),
    }
    return {
        "days": days,
        "period_times": [["08:30", "09:15"], ["09:30", "10:15"], ["10:30", "11:15"],
                         ["11:35", "12:20"], ["12:40", "13:25"], ["13:30", "14:15"],
                         ["14:20", "15:05"], ["15:10", "15:55"]],
        "meal_break_after": 2, "meal_label": "ХАРЧУВАННЯ У ЇДАЛЬНІ",
        "holidays": [{"start": "2026-10-26", "end": "2026-11-01"},
                     {"start": "2026-12-24", "end": "2027-01-10"},
                     {"start": "2027-03-22", "end": "2027-03-28"}],
        "workload_teacher": "Прізвище І.О.", "workload_code": "000000",
    }


def build_sample_workload(path):
    from .workload_io import export_workload_docx
    return export_workload_docx(sample_workload_config(), path)


def build_sample_workload_jpeg(path):
    from .workload_io import export_workload_jpeg
    return export_workload_jpeg(sample_workload_config(), path)


def build_sample_workload_csv(path):
    """Той самий зразок у CSV (Excel) — програма читає і такий файл."""
    import csv
    import tempfile
    from .schedule_io import table_rows
    with tempfile.TemporaryDirectory() as folder:
        docx = build_sample_workload(Path(folder) / "n.docx")
        rows = table_rows(docx)
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        for row in rows:
            writer.writerow([row[0]] if len(set(row)) == 1 else row)
    return Path(path)


SCHEDULE_DESCRIPTION = """ЯК ВИГЛЯДАЄ РОЗКЛАД (навантаження), який розуміє програма:

• Аркуш А4, шрифт Times New Roman, суцільна таблиця з рамками.
• Перший рядок: «НАВАНТАЖЕННЯ (Прізвище І.О.) // канікули 26.10-01.11, 24.12-10.01, 22.03-28.03 // код». Програма сама бере з нього прізвище, код і КАНІКУЛИ.
• Другий рядок: «№», «Понеділок», «Вівторок», «Середа», «Четвер», «П’ятниця».
• Далі один рядок на один урок. У першій клітинці номер уроку жирним, а під ним час: 08.30-09.15.
• Клітинка — це потік: «8-Б ІУ», «11 ІУ профіль», «9-Г Право». Порожня клітинка — уроку немає.
• Чисельник / знаменник через косу рису: «8-Б ІУ / ГО» (у чисельник — 8-Б ІУ, у знаменник — 8-Б ГО); «/ 9-Г Право» — лише знаменник; «11 ІУ стандарт /» — лише чисельник.
• Після 2-го уроку окремий червоний рядок «ХАРЧУВАННЯ У ЇДАЛЬНІ».

ІУ — Історія України, ВІ — Всесвітня історія, ГО — Громадянська освіта, Право — Правознавство.

КУДИ ЗАВАНТАЖУВАТИ: редактор → вкладка «Розклад» → перетягніть файл на таблицю або «Завантажити навантаження / розклад». Згенерувати такий документ (Word, JPEG) з вашого розкладу можна там само."""


def build_sample_ktp(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    head = doc.add_paragraph("Календарно-тематичне планування. Історія України, 8 клас (зразок)")
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.runs[0].bold = True
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("№ з/п", "Дата", "Тема уроку", "Домашнє завдання")):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    section = table.add_row().cells[0].merge(table.rows[-1].cells[3])
    section.text = "Розділ 1. Приклад назви розділу"
    section.paragraphs[0].runs[0].bold = True
    rows = [
        ("1", "8-Б 04.09. 8-В 04.09. 8-Г 03.09.", "Перша тема уроку, записана повністю", "прочитати § 1"),
        ("2", "8-Б 07.09. 8-В 07.09. 8-Г 07.09.", "Друга тема уроку", "прочитати § 2, повторити § 1"),
        ("3", "8-Б 11.09. 8-В 11.09. 8-Г 10.09.", "Третя тема. Практичне заняття", "завдання в зошиті"),
        ("4-5", "8-Б 14.09. 8-В 14.09. 8-Г 14.09.", "Подвійний урок з однією темою", "прочитати § 4–5"),
        ("6", "8-Б 18.09. 8-В 18.09. 8-Г 17.09.", "Узагальнення з розділу 1", ""),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    for width, column in zip((Cm(1.6), Cm(5.4), Cm(7.2), Cm(4.2)), table.columns):
        for cell in column.cells:
            cell.width = width
    doc.save(str(path))
    return Path(path)


def _copy(window, text, label="Промт скопійовано"):
    window.clipboard_clear()
    window.clipboard_append(text)
    window.update()
    messagebox.showinfo("Буфер обміну", label + ". Вставте його (Ctrl+V) у свій ШІ-чат.", parent=window)


def schedule_tab(book, win):
    """Зразок розкладу: опис формату, картинка, готові файли й промт для ШІ."""
    frame = ttk.Frame(book, padding=8)
    book.add(frame, text="Зразок розкладу")
    left = ttk.Frame(frame)
    left.pack(side="left", fill="both", expand=True)
    right = ttk.Frame(frame)
    right.pack(side="left", fill="y", padx=(12, 0))
    ttk.Label(left, text=SCHEDULE_DESCRIPTION, wraplength=640, justify="left").pack(anchor="w")
    bar = ttk.Frame(left)
    bar.pack(fill="x", pady=8)

    def saver(title, builder, suffix, kinds, filename):
        def run():
            path = filedialog.asksaveasfilename(parent=win, title=title, defaultextension=suffix,
                                                initialfile=filename, filetypes=kinds)
            if not path:
                return
            try:
                builder(path)
                messagebox.showinfo("Зразок", f"Зразок збережено:\n{path}", parent=win)
            except Exception as ex:
                messagebox.showerror("Зразок", str(ex), parent=win)
        return run
    ttk.Button(bar, text="💾 Зразок Word", command=saver(
        "Зберегти зразок розкладу (Word)", build_sample_workload, ".docx",
        [("Word DOCX", "*.docx")], "Зразок_розкладу.docx")).pack(side="left")
    ttk.Button(bar, text="🖼 Зразок JPEG", command=saver(
        "Зберегти зразок розкладу (зображення)", build_sample_workload_jpeg, ".jpg",
        [("Зображення JPEG", "*.jpg")], "Зразок_розкладу.jpg")).pack(side="left", padx=6)
    ttk.Button(bar, text="📊 Зразок Excel (CSV)", command=saver(
        "Зберегти зразок розкладу (CSV)", build_sample_workload_csv, ".csv",
        [("CSV (Excel)", "*.csv")], "Зразок_розкладу.csv")).pack(side="left")
    ttk.Button(bar, text="📋 Промт для ШІ",
               command=lambda: _copy(win, WORKLOAD_PROMPT)).pack(side="left", padx=6)
    ttk.Label(left, text="Точний промт, за яким ШІ (або людина) створить такий самий документ:").pack(anchor="w")
    box = ScrolledText(left, wrap="word", font=("Segoe UI", 9), height=4)
    box.insert("1.0", WORKLOAD_PROMPT)
    box.pack(fill="both", expand=True)
    ttk.Label(right, text="Приклад (умовні дані):", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    try:
        import tempfile
        from PIL import Image
        folder = Path(tempfile.mkdtemp(prefix="pomichnyk_"))
        jpg = build_sample_workload_jpeg(folder / "sample.jpg")
        with Image.open(jpg) as picture:
            width = 470
            picture.resize((width, int(picture.height * width / picture.width))).save(folder / "sample.png")
        win._sample_image = tk.PhotoImage(file=str(folder / "sample.png"))
        ttk.Label(right, image=win._sample_image, relief="solid").pack(anchor="nw", pady=4)
    except Exception as ex:                                  # без картинки вкладка все одно працює
        ttk.Label(right, text=f"Попередній перегляд недоступний: {ex}", wraplength=420).pack(anchor="w")


def show_samples(parent):
    win = tk.Toplevel(parent)
    win.title("Зразки документів, які розуміє програма")
    fit_work_window(win, "large")
    win.transient(parent)
    win.escape_closes = True
    ttk.Label(win, padding=(12, 10), wraplength=1000, justify="left",
              text=("Завантажуйте ці документи перетягуванням мишкою або кнопками в редакторі: "
                    "«Розклад» → «Завантажити навантаження / розклад», «Календарні плани» → "
                    "перетягніть файл КТП на потрібний клас. Збережіть зразок, подивіться, як він "
                    "виглядає, а точний промт нижче допоможе створити такий самий документ "
                    "самостійно або через ШІ.")).pack(anchor="w")
    book = ttk.Notebook(win)
    book.pack(fill="both", expand=True, padx=10, pady=6)

    def tab(title, description, prompt, builder, filename):
        frame = ttk.Frame(book, padding=8)
        book.add(frame, text=title)
        ttk.Label(frame, text=description, wraplength=980, justify="left",
                  foreground="#365777").pack(anchor="w", pady=(0, 6))
        bar = ttk.Frame(frame)
        bar.pack(fill="x", pady=(0, 6))

        def save_sample():
            path = filedialog.asksaveasfilename(
                parent=win, title="Зберегти зразок", defaultextension=".docx",
                initialfile=filename, filetypes=[("Word DOCX", "*.docx")])
            if not path:
                return
            try:
                builder(path)
                messagebox.showinfo("Зразок", f"Зразок збережено:\n{path}", parent=win)
            except Exception as ex:
                messagebox.showerror("Зразок", str(ex), parent=win)
        ttk.Button(bar, text="💾 Зберегти зразок (Word)…", command=save_sample).pack(side="left")
        ttk.Button(bar, text="📋 Копіювати промт для ШІ",
                   command=lambda: _copy(win, prompt)).pack(side="left", padx=8)
        box = ScrolledText(frame, wrap="word", font=("Segoe UI", 10), height=8)
        box.insert("1.0", prompt)
        box.pack(fill="both", expand=True)

    schedule_tab(book, win)
    tab("Календарне планування (КТП)",
        "Таблиця «№ — Дата — Тема — Домашнє завдання». Програма підтягує теми, домашні "
        "завдання, назви розділів (рядки без номера пропускаються) та дати за класами.",
        KTP_PROMPT, build_sample_ktp, "Зразок_КТП.docx")
    ttk.Button(win, text="Закрити", command=win.destroy).pack(pady=8)
    return win


# ---------- «Мої дані» ----------
from .data_tools import (blank_config, data_files, export_all_data,  # noqa: E402,F401
                         is_secret, reset_to_blank, restore_from_zip)


def open_folder(path: Path):
    path.mkdir(parents=True, exist_ok=True)
    if sys.platform == "win32":
        os.startfile(str(path))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


def show_data_folder(parent):
    win = tk.Toplevel(parent)
    win.title("Мої дані")
    fit_work_window(win, "small")
    win.transient(parent)
    win.escape_closes = True
    ttk.Label(win, padding=(14, 12), wraplength=620, justify="left",
              text=("Усе, що ви вносите в програму (розклад, КТП, готові Word, вкладення, "
                    "резервні копії), зберігається в ОДНІЙ папці на цьому комп'ютері:")).pack(anchor="w")
    path_box = tk.Text(win, height=2, wrap="word", font=("Segoe UI", 10), relief="flat", bg="#EAF2FA")
    path_box.insert("1.0", str(ROOT))
    path_box.configure(state="disabled")
    path_box.pack(fill="x", padx=14)
    ttk.Label(win, padding=(14, 10), wraplength=620, justify="left", foreground="#365777",
              text=("Паролі, токени Google та OAuth-файли в копію НЕ потрапляють. "
                    "Копію можна відкрити на іншому комп'ютері або використати для інших потреб.")
              ).pack(anchor="w")
    bar = ttk.Frame(win, padding=(14, 4))
    bar.pack(fill="x")

    def export():
        target = filedialog.asksaveasfilename(
            parent=win, title="Зберегти копію всіх даних", defaultextension=".zip",
            initialfile="Помічник учителя — мої дані.zip", filetypes=[("ZIP-архів", "*.zip")])
        if not target:
            return
        try:
            count = export_all_data(target)
            messagebox.showinfo("Мої дані", f"Збережено файлів: {count}.\n{target}", parent=win)
        except Exception as ex:
            messagebox.showerror("Мої дані", str(ex), parent=win)
    ttk.Button(bar, text="📂 Відкрити папку з даними", command=lambda: open_folder(ROOT)).pack(side="left")
    ttk.Button(bar, text="⬇ Зберегти копію всіх даних (ZIP)…", command=export).pack(side="left", padx=8)

    def default_name(prefix):
        from datetime import datetime
        return f"{prefix} {datetime.now():%Y-%m-%d %H-%M}.zip"

    def after_change(message):
        reload = getattr(parent, "reload_data", None)
        if reload:
            reload()
        messagebox.showinfo("Мої дані", message, parent=win)

    def blank_start():
        if not messagebox.askyesno(
                "Почати з порожньої програми",
                "Програма стане ПОРОЖНЬОЮ: без розкладу, календарних планів, класів, готових Word "
                "і вкладень. Підключення до Google збережеться.\n\n"
                "СПЕРШУ ви збережете повну копію всіх даних у ZIP — з неї можна відновитися кнопкою "
                "«Відновити з копії». Продовжити?", parent=win):
            return
        target = filedialog.asksaveasfilename(
            parent=win, title="Куди зберегти повну копію ПЕРЕД очищенням", defaultextension=".zip",
            initialfile=default_name("Копія даних до очищення"), filetypes=[("ZIP-архів", "*.zip")])
        if not target:
            messagebox.showinfo("Мої дані", "Без копії нічого не очищено.", parent=win)
            return
        try:
            count = reset_to_blank(target)
        except Exception as ex:
            messagebox.showerror("Мої дані", f"Не вдалося: {ex}\nДані не змінено, якщо копію не створено.",
                                 parent=win)
            return
        after_change(f"Готово. Збережено файлів у копії: {count}.\nПрограма порожня: завантажте "
                     "розклад (редактор → «Розклад») і КТП (перетягніть на клас).")

    def restore():
        source = filedialog.askopenfilename(
            parent=win, title="Оберіть копію даних (ZIP)", filetypes=[("ZIP-архів", "*.zip")])
        if not source:
            return
        if not messagebox.askyesno(
                "Відновити з копії",
                "Поточні дані будуть замінені даними з копії. Перед цим програма сама збереже "
                "поточний стан у ZIP-файл. Продовжити?", parent=win):
            return
        safety = filedialog.asksaveasfilename(
            parent=win, title="Куди зберегти ПОТОЧНИЙ стан перед відновленням",
            defaultextension=".zip", initialfile=default_name("Стан перед відновленням"),
            filetypes=[("ZIP-архів", "*.zip")])
        if not safety:
            return
        try:
            count = restore_from_zip(source, safety)
        except Exception as ex:
            messagebox.showerror("Мої дані", str(ex), parent=win)
            return
        after_change(f"Відновлено файлів: {count}.")

    ttk.Button(bar, text="🧹 Почати з порожньої програми…", command=blank_start).pack(side="left", padx=8)
    ttk.Button(bar, text="♻ Відновити з копії…", command=restore).pack(side="left")
    ttk.Button(win, text="Закрити", command=win.destroy).pack(pady=10)
    return win
