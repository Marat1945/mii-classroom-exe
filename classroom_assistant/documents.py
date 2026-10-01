"""Оформлення Word: Times New Roman 14, кольори, конспект, безпека, Д/з останнім."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY=RGBColor(31,78,121)
def shade(paragraph,fill):
    pPr=paragraph._p.get_or_add_pPr()
    shd=OxmlElement("w:shd")
    shd.set(qn("w:fill"),fill)
    pPr.append(shd)

def create_word(lesson, destination, material=None):
    """material None => clearly flagged UNFINISHED outline, not an invented historical lecture."""
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
    sec.left_margin=Cm(2.1);sec.right_margin=Cm(1.7)
    normal=doc.styles["Normal"]
    normal.font.name="Times New Roman"; normal.font.size=Pt(14)
    normal.paragraph_format.first_line_indent=Cm(.9)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.09
    for key,size,color in [("Title",18,NAVY),("Heading 1",16,NAVY)]:
        st=doc.styles[key];st.font.name="Times New Roman";st.font.size=Pt(size)
        st.font.bold=True;st.font.color.rgb=color
        st.paragraph_format.first_line_indent=Cm(0)
        st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
    def p(text="",bold=False, bg=None, left=False):
        q=doc.add_paragraph()
        q.alignment=WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.JUSTIFY
        if left:q.paragraph_format.first_line_indent=Cm(0)
        if bg:shade(q,bg)
        r=q.add_run(text);r.bold=bold;r.font.name="Times New Roman";r.font.size=Pt(14)
        return q
    def head(text,bg="EFF4FA",color=NAVY):
        q=doc.add_heading(text,1);shade(q,bg)
        for r in q.runs:r.font.color.rgb=color
    p("Урок "+lesson.day[8:10]+"."+lesson.day[5:7]+" — "+lesson.topic.rstrip(".")+"。".replace("。","."),True,"EAF2F8",True)
    t=doc.add_paragraph(style="Title");t.add_run(lesson.topic)
    if material is None:
        p("РОБОЧА ЗАГОТОВКА. Повну лекцію ще не згенеровано або не перевірено вчителем. Цей файл НЕ ПУБЛІКУВАТИ учням.",True,"FCE8E6",True)
        head("Тема й навчальні орієнтири")
        p("Тема за календарним планом: "+lesson.topic)
        p("Джерело теми: "+lesson.source_file)
        head("Матеріал для пояснення")
        p("Це місце для перевіреного навчального пояснення. Для автоматичного створення повної лекції налаштуйте AI або виберіть готовий Word.")
        notebook=["Тема: "+lesson.topic,"Основні поняття, факти і причинно-наслідкові зв’язки — заповнити після підготовки лекції."]
    else:
        sections=material.get("sections",[])
        for i,item in enumerate(sections,1):
            head(f"{i}. {item.get('heading','Розділ')}")
            for para in item.get("paragraphs",[]):
                if para.strip():p(para)
        notebook=material.get("notebook",[])
        if not sections or not notebook:
            raise ValueError("Необхідні розділи лекції та план-конспект")
    head("ПЛАН-КОНСПЕКТ УРОКУ ДЛЯ ЗАПИСУ В ЗОШИТ","EAF4E4",RGBColor(53,96,37))
    p("У зошит потрібно записати тільки наведений нижче план-конспект. Усю лекцію переписувати не потрібно.",True,"EAF4E4")
    for item in notebook:
        p("• "+item)
    head("Техніка безпеки","FCE8E6",RGBColor(155,37,31))
    p("Під час повітряної тривоги перебувайте в безпечному місці. Жовта тривога також є сигналом небезпеки — не ігноруйте її. До навчання повертайтеся лише тоді, коли це безпечно.",False, "FCE8E6")
    p("Д/з: "+(lesson.homework or "Не зазначено у КТП — уточнити у вчителя."),True,"FFF2CC",True)
    path=Path(destination);path.parent.mkdir(exist_ok=True,parents=True)
    doc.save(path)
    return path
