"""Перегляд вкладення Classroom прямо в програмі (подвійний клац лівою кнопкою по вкладенню).

Зображення — з масштабуванням (колесо миші, кнопки, «По вікну», 1:1) і перетягуванням; Word — текст і таблиці;
решта файлів — відомості й кнопка «Відкрити». Одне вікно на файл: повторний подвійний клац лише виносить його наперед.
"""
from __future__ import annotations

import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk
from tkinter.scrolledtext import ScrolledText

from . import materials
from .window_ui import fit_work_window, remember_window, reuse_window

IMAGE_TYPES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")
WORD_TYPES = (".docx",)
MIN_SCALE, MAX_SCALE, STEP = 0.05, 8.0, 1.25
WORD_LIMIT = 600                                                   # абзаців у попередньому перегляді Word


def kind_of(path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in IMAGE_TYPES:
        return "image"
    if suffix in WORD_TYPES:
        return "word"
    return "other"


def human_size(size: int) -> str:
    return materials.human_size(size)


class ImageView(ttk.Frame):
    """Зображення з масштабом: «по вікну» (за замовчуванням), колесо миші, +/−, 1:1, перетягування лівою кнопкою."""

    def __init__(self, parent, path):
        super().__init__(parent)
        from PIL import Image
        with Image.open(path) as source:
            source.load()
            self.original = source.convert("RGBA") if source.mode in ("P", "LA", "PA", "RGBA") else source.convert("RGB")
        self.size = self.original.size
        self.scale = 1.0
        self.fit_mode = True
        self._photo = None
        self._job = None
        self.canvas = tk.Canvas(self, bg="#2B2822", highlightthickness=0)
        across = ttk.Scrollbar(self, orient="horizontal", command=self.canvas.xview)
        down = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(xscrollcommand=across.set, yscrollcommand=down.set)
        across.pack(side="bottom", fill="x")
        down.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.canvas.bind("<Configure>", self._on_resize)
        self.canvas.bind("<MouseWheel>", lambda e: self.zoom(STEP if e.delta > 0 else 1 / STEP, (e.x, e.y)))
        self.canvas.bind("<Button-4>", lambda e: self.zoom(STEP, (e.x, e.y)))
        self.canvas.bind("<Button-5>", lambda e: self.zoom(1 / STEP, (e.x, e.y)))
        self.canvas.bind("<ButtonPress-1>", lambda e: self.canvas.scan_mark(e.x, e.y))
        self.canvas.bind("<B1-Motion>", lambda e: self.canvas.scan_dragto(e.x, e.y, gain=1))
        self.canvas.bind("<Double-Button-1>", lambda _e: self.toggle())
        self.on_scale = None

    def _fit_scale(self) -> float:
        width, height = max(self.canvas.winfo_width(), 50), max(self.canvas.winfo_height(), 50)
        return max(MIN_SCALE, min(1.0, width / self.size[0], height / self.size[1]))     # великі зменшуємо, малі не збільшуємо

    def _on_resize(self, _event=None):
        if self.fit_mode:
            self.set_scale(self._fit_scale(), keep_fit=True)

    def set_scale(self, scale, keep_fit=False, anchor=None):
        scale = max(MIN_SCALE, min(MAX_SCALE, scale))
        if not keep_fit:
            self.fit_mode = False
        self.scale = scale
        self.render(anchor)
        if self.on_scale:
            self.on_scale(scale)

    def render(self, anchor=None):
        from PIL import Image, ImageTk
        width, height = max(1, round(self.size[0] * self.scale)), max(1, round(self.size[1] * self.scale))
        resized = self.original.resize((width, height), Image.LANCZOS if self.scale < 1 else Image.BICUBIC)
        self._photo = ImageTk.PhotoImage(resized, master=self)
        self.canvas.delete("all")
        offset_x = max(0, (self.canvas.winfo_width() - width) // 2)                       # менше за вікно — по центру
        offset_y = max(0, (self.canvas.winfo_height() - height) // 2)
        self.canvas.create_image(offset_x, offset_y, anchor="nw", image=self._photo)
        self.canvas.configure(scrollregion=(0, 0, max(width, self.canvas.winfo_width()), max(height, self.canvas.winfo_height())))

    def zoom(self, factor, anchor=None):
        self.set_scale(self.scale * factor, anchor=anchor)

    def fit(self):
        self.fit_mode = True
        self.set_scale(self._fit_scale(), keep_fit=True)

    def actual_size(self):
        self.set_scale(1.0)

    def toggle(self):
        self.actual_size() if self.fit_mode else self.fit()


def _word_text(path) -> list:
    """Абзаци й таблиці Word як [(стиль, текст)]; без запуску Word."""
    from docx import Document
    document = Document(str(path))
    items = []
    for paragraph in document.paragraphs[:WORD_LIMIT]:
        text = paragraph.text.strip()
        if text:
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            items.append(("h" if style.startswith(("heading", "заголовок", "title")) else "p", text))
    for table in document.tables[:10]:
        for row in table.rows[:60]:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                items.append(("t", "  │  ".join(cells)))
    return items


def show_preview(app, path, title=None):
    """Відкрити перегляд вкладення (одне вікно на файл). Повертає вікно або None, якщо файл не вдалося відкрити."""
    path = Path(path)
    key = "preview:" + str(path)
    existing = reuse_window(app, key)
    if existing:
        return existing
    if not path.is_file():
        messagebox.showwarning("Перегляд вкладення", f"Файл не знайдено на диску:\n{path}", parent=app)
        return None
    kind = kind_of(path)
    win = tk.Toplevel(app)
    win.title(f"Перегляд: {title or path.name}")
    fit_work_window(win, "large")
    win.transient(app)
    win.escape_closes = True
    win.kind = kind
    top = ttk.Frame(win, padding=(10, 6))
    top.pack(fill="x")
    info = tk.StringVar(master=win, value=f"{title or path.name}   ·   {human_size(path.stat().st_size)}")
    win.info = info
    ttk.Label(top, textvariable=info, font=("Segoe UI", 10, "bold")).pack(side="left")
    bottom = ttk.Frame(win, padding=(10, 6))
    bottom.pack(side="bottom", fill="x")
    ttk.Button(bottom, text="Відкрити у програмі за замовчуванням", command=lambda: _open_external(win, path)).pack(side="left")
    ttk.Button(bottom, text="Показати у папці", command=lambda: materials.reveal(path)).pack(side="left", padx=8)
    ttk.Button(bottom, text="Закрити", command=win.destroy).pack(side="right")
    try:
        if kind == "image":
            view = ImageView(win, path)
            win.view = view
            zoom_text = tk.StringVar(master=win, value="")
            win.zoom_text = zoom_text
            view.on_scale = lambda s: zoom_text.set(f"{round(s * 100)}%   ({view.size[0]}×{view.size[1]})")
            tools = ttk.Frame(top)
            tools.pack(side="right")
            ttk.Button(tools, text="−", width=3, command=lambda: view.zoom(1 / STEP)).pack(side="left")
            ttk.Label(tools, textvariable=zoom_text, width=18, anchor="center").pack(side="left", padx=6)
            ttk.Button(tools, text="+", width=3, command=lambda: view.zoom(STEP)).pack(side="left")
            ttk.Button(tools, text="По вікну", command=view.fit).pack(side="left", padx=(10, 0))
            ttk.Button(tools, text="1:1", width=5, command=view.actual_size).pack(side="left", padx=4)
            view.pack(fill="both", expand=True, padx=10, pady=4)
            win.after(60, view.fit)
        elif kind == "word":
            box = ScrolledText(win, wrap="word", font=("Segoe UI", 10), padx=14, pady=10, relief="flat")
            box.tag_configure("h", font=("Segoe UI", 12, "bold"), foreground="#1F4E79", spacing1=8, spacing3=3)
            box.tag_configure("t", font=("Consolas", 9), foreground="#444444")
            win.text = box
            for style, text in _word_text(path):
                box.insert("end", text + "\n", style if style != "p" else ())
            if not box.get("1.0", "end").strip():
                box.insert("end", "У документі немає тексту для попереднього перегляду. Відкрийте його у Word.")
            box.configure(state="disabled")
            box.pack(fill="both", expand=True, padx=10, pady=4)
        else:
            ttk.Label(win, justify="left", padding=24, font=("Segoe UI", 11),
                      text=f"Файл: {path.name}\nТип: {path.suffix.upper().lstrip('.') or 'невідомий'}\nРозмір: "
                           f"{human_size(path.stat().st_size)}\n\nДля цього типу файла попереднього перегляду в програмі немає.\n"
                           "Натисніть «Відкрити у програмі за замовчуванням».").pack(fill="both", expand=True)
    except Exception as error:                                                           # пошкоджений файл: пояснюємо, не падаємо
        win.destroy()
        messagebox.showwarning("Перегляд вкладення", f"Не вдалося показати файл «{path.name}»:\n{error}\n\n"
                               "Його можна відкрити у програмі за замовчуванням.", parent=app)
        return None
    remember_window(app, key, win)
    return win


def _open_external(win, path):
    try:
        materials.open_file(path)
    except Exception as error:
        messagebox.showerror("Перегляд вкладення", f"Не вдалося відкрити файл:\n{error}", parent=win)
