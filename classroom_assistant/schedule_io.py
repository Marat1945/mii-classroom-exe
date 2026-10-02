"""Розклад: імпорт таблиці, ручне редагування та експорт у друкований DOCX."""
from __future__ import annotations
import csv
import io
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from .editor_core import csv_rows,docx_table_rows,tidy

DAYS=("Понеділок","Вівторок","Середа","Четвер","П’ятниця")
ALTS=("пн","вт","ср","чт","пт")
NAMES=("понеділок","вівторок","середа","четвер","пятниця","п'ятниця","п’ятниця")

def _xlsx_rows(path):
    """Читання першого аркуша XLSX стандартною бібліотекою, без макросів."""
    m={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    r={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    with ZipFile(path) as z:
        strings=[]
        if 'xl/sharedStrings.xml' in z.namelist():
            ss=ET.fromstring(z.read('xl/sharedStrings.xml'))
            strings=[''.join(x.text or '' for x in si.iter('{'+m['m']+'}t'))
                     for si in ss.findall('m:si',m)]
        workbook=ET.fromstring(z.read('xl/workbook.xml'))
        sheet=workbook.find('m:sheets/m:sheet',m)
        if sheet is None:raise ValueError("У XLSX немає аркушів")
        relid=sheet.get('{'+r['r']+'}id')
        reltree=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        relnode=next((x for x in reltree if x.get('Id')==relid),None)
        if relnode is None:raise ValueError("Не вдалося прочитати аркуш XLSX")
        target=relnode.get('Target').lstrip('/')
        if not target.startswith('xl/'):target='xl/'+target
        doc=ET.fromstring(z.read(target))
        ret=[]
        for row in doc.findall('.//m:sheetData/m:row',m):
            cells={}
            for cell in row.findall('m:c',m):
                ref=cell.get('r','')
                match=re.match(r'^([A-Z]+)',ref)
                if not match:continue
                ind=0
                for ch in match.group(1):ind=ind*26+ord(ch)-64
                val=cell.find('m:v',m)
                text=(val.text or '') if val is not None else ''
                if cell.get('t')=='s':
                    text=strings[int(text)] if text else ''
                elif cell.get('t')=='inlineStr':
                    text=''.join(n.text or '' for n in cell.findall('.//m:t',m))
                cells[ind-1]=tidy(text)
            if cells:
                ret.append([cells.get(i,'') for i in range(max(cells)+1)])
        return ret

def table_rows(path):
    suffix=Path(path).suffix.lower()
    if suffix=='.csv':return csv_rows(path)
    if suffix in ('.doc','.docx'):return docx_table_rows(path)
    if suffix=='.xlsx':return _xlsx_rows(path)
    raise ValueError("Імпорт розкладу підтримує DOCX, DOC, CSV та XLSX.")

def _normalize(text):
    return tidy(text).casefold().replace('’',"'").replace('–','-').replace('—','-')

def resolve_stream(raw,stream_names,base="",aliases=None):
    raw=tidy(raw)
    if not raw or raw in ('—','-','–'):return None
    match=re.match(r'^(\d{1,2})\s*[-–]\s*([А-Яа-яA-Z])(?:\s+(.+))?$',raw)
    if base and not re.match(r'^\d',raw):
        c=re.match(r'^(\d{1,2})\s*[-–]\s*[А-Яа-яA-Z]',base)
        if c:
            raw=base.split()[0]+' '+raw
    alias=(aliases or {}).get(_normalize(raw))
    if alias in stream_names:return alias
    exact=[s for s in stream_names if _normalize(s)==_normalize(raw)]
    if len(exact)==1:return exact[0]
    # «5-Г ІУ» у вашому файлі, а в програмі єдиний потік класу — «5-Г історія».
    head=_normalize(raw).split()[0] if raw.split() else ""
    if re.match(r'^\d{1,2}-[а-яіїєґa-z]$',head):
        same=[s for s in stream_names if _normalize(s).split()[:1]==[head]]
        if len(same)==1:return same[0]
    return "? "+raw

def parse_cell(value,known_streams,aliases=None):
    text=tidy(value).replace('\n',' ')
    if not text or text in ('—','–','-'):return [None,None]
    halves=[tidy(x) for x in text.split('/',1)]
    if len(halves)==1:halves=[halves[0],halves[0]]
    a=resolve_stream(halves[0],known_streams,aliases=aliases)
    b=resolve_stream(halves[1],known_streams,base=halves[0],aliases=aliases)
    return [a,b]

def decode_schedule(rows,config):
    """Повертає всі дні без перезапису; невідомі потоки повідомляються у прев'ю."""
    header_i=None
    daycols={}
    for i,row in enumerate(rows[:18]):
        found={}
        for j,cell in enumerate(row):
            text=_normalize(cell)
            for d,token in enumerate(("понеділ","вівтор","серед","четвер","пятниц","п'ятниц")):
                if token in text:
                    found[d if d<4 else 4]=j
        if len(found)==5:
            header_i=i;daycols=found;break
    if header_i is None:
        raise ValueError("Не знайдено рядок із п'ятьма днями тижня. Потрібна таблиця: №, Понеділок, ..., П'ятниця.")
    count=len(config['period_times'])
    schedule={str(d):[[None,None] for _ in range(count)] for d in range(5)}
    unknown=set()
    found_numbers=set()
    aliases=config.get("schedule_aliases",{})
    for row in rows[header_i+1:]:
        if not row:continue
        label=_normalize(row[0])
        match=re.match(r'^(\d{1,2})(?:\b|[.)\s])',label)
        if not match:continue
        period=int(match.group(1))
        if period<1 or period>count:continue
        found_numbers.add(period)
        for day,col in daycols.items():
            value=row[col] if col<len(row) else ''
            pair=parse_cell(value,config['course_map'],aliases=aliases)
            for x in pair:
                if x and x.startswith('? '):unknown.add(x[2:])
            schedule[str(day)][period-1]=pair
    if not found_numbers:raise ValueError("У таблиці не знайдено уроків із номерами 1, 2, ...")
    return schedule,sorted(unknown),sorted(found_numbers)

def export_schedule_docx(config,path,show_bells=False,show_meal=False,show_numbers=True):
    from docx import Document
    from docx.shared import Pt,Cm,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    sect=doc.sections[0]
    sect.orientation=WD_ORIENT.LANDSCAPE
    sect.page_width=Cm(29.7);sect.page_height=Cm(21.0)
    sect.left_margin=Cm(1.3);sect.right_margin=Cm(1.3)
    sect.top_margin=Cm(1.5);sect.bottom_margin=Cm(1.5)
    default=doc.styles['Normal']
    default.font.name='Times New Roman';default.font.size=Pt(11)
    head=doc.add_paragraph()
    head.alignment=WD_ALIGN_PARAGRAPH.CENTER
    h=head.add_run("НАВАНТАЖЕННЯ · РОЗКЛАД УРОКІВ")
    h.bold=True;h.font.name='Times New Roman';h.font.size=Pt(16)
    years=f"{config['year_start'][:4]}–{config['year_end'][:4]}"
    sub=doc.add_paragraph(f"Навчальний рік: {years}    •    Чисельник / знаменник")
    sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
    table=doc.add_table(rows=1,cols=6)
    table.style='Table Grid'
    for j,value in enumerate(('№',)+DAYS):table.rows[0].cells[j].text=value
    def fill_color(cell,color):
        tcpr=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),color);tcpr.append(s)
    for cell in table.rows[0].cells:fill_color(cell,'D9EAF4')
    count=len(config['period_times'])
    meal_after=int(config.get('meal_break_after',2))
    for n in range(1,count+1):
        row=table.add_row()
        times=config["period_times"][n-1]
        number=str(n) if show_numbers else ""
        row.cells[0].text=number+(f"\n{times[0]}–{times[1]}" if show_bells else "")
        for d in range(5):
            entry=config['days'].get(str(d),[])
            a,b=entry[n-1] if len(entry)>=n else (None,None)
            if a==b:
                content=a or ''
            else:
                content=f"{a or ''} / {b or ''}"
            row.cells[d+1].text=content
        if show_meal and n==meal_after:
            meal=table.add_row().cells[0]
            merged=meal.merge(table.rows[-1].cells[5])
            merged.text=config.get('meal_label','ХАРЧУВАННЯ У ЇДАЛЬНІ')
            fill_color(merged,'FFF0D1')
            for p in merged.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.bold=True
    for index,row in enumerate(table.rows):
        row.height=Cm(1.14)
        row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell in row.cells:
            cell.vertical_alignment=1
            for para in cell.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(11)
                    if index==0:run.bold=True
    doc.save(path)
    return Path(path)

def export_schedule_jpeg(config,path,show_bells=False,show_meal=False,show_numbers=True):
    """Створити друкований розклад JPEG без зовнішніх інтернет-сервісів."""
    from PIL import Image,ImageDraw,ImageFont
    from textwrap import wrap
    from pathlib import Path
    import os

    def pick_font(size,bold=False):
        names=(
            ["C:/Windows/Fonts/arialbd.ttf","/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]
            if bold else
            ["C:/Windows/Fonts/arial.ttf","/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
        )
        for name in names:
            if Path(name).exists():
                return ImageFont.truetype(name,size)
        return ImageFont.load_default()

    width=2000
    margin=55
    label_w=190
    day_w=(width-2*margin-label_w)//5
    row_h=124 if not show_bells else 142
    title_h=185
    header_h=100
    row_count=len(config.get("period_times",[]))
    meal_after=int(config.get("meal_break_after",2))
    has_meal=(show_meal and 0<meal_after<=row_count)
    height=title_h+header_h+row_h*row_count+(62 if has_meal else 0)+margin
    img=Image.new("RGB",(width,height),"white")
    draw=ImageDraw.Draw(img)
    bold=pick_font(35,True);regular=pick_font(29)
    header=pick_font(33,True);title=pick_font(46,True)
    small=pick_font(24)
    title_txt="НАВАНТАЖЕННЯ · РОЗКЛАД УРОКІВ"
    left=(width-draw.textbbox((0,0),title_txt,font=title)[2])//2
    draw.text((left,27),title_txt,font=title,fill="#182e43")
    subtitle=f"Навчальний рік {config['year_start'][:4]}–{config['year_end'][:4]}   •   Чисельник / знаменник"
    subleft=(width-draw.textbbox((0,0),subtitle,font=regular)[2])//2
    draw.text((subleft,105),subtitle,font=regular,fill="#3b5160")
    y=title_h
    headers=("№","Понеділок","Вівторок","Середа","Четвер","П’ятниця")
    x_positions=[margin,margin+label_w]+[margin+label_w+i*day_w for i in range(1,6)]
    for i,header_txt in enumerate(headers):
        x=x_positions[i]
        w=label_w if i==0 else day_w
        draw.rectangle((x,y,x+w,y+header_h),fill="#d9eaf4",outline="#899daa",width=2)
        bounds=draw.textbbox((0,0),header_txt,font=header)
        draw.text((x+(w-(bounds[2]-bounds[0]))//2,y+30),header_txt,font=header,fill="#182e43")
    y+=header_h
    for row_num in range(1,row_count+1):
        start,end=config["period_times"][row_num-1]
        for j in range(6):
            x=x_positions[j]
            cell_w=label_w if j==0 else day_w
            draw.rectangle((x,y,x+cell_w,y+row_h),fill="white",outline="#899daa",width=2)
            if j==0:
                title_row=str(row_num) if show_numbers else ""
                time_row=f"{start}–{end}" if show_bells else ""
                if title_row:
                    draw.text((x+cell_w//2-11,y+20),title_row,font=bold,fill="#172839")
                if time_row:
                    time_font=small
                    tw=draw.textbbox((0,0),time_row,font=time_font)[2]
                    draw.text((x+(cell_w-tw)//2,y+row_h-45),time_row,font=time_font,fill="#27384a")
            else:
                pairs=config["days"].get(str(j-1),[])
                first,second=pairs[row_num-1] if len(pairs)>=row_num else (None,None)
                # The '-' placeholder is never printed for an empty week.
                content=first if first==second else (f"{first or ''} / {second or ''}" if first or second else "")
                content=str(content or "")
                lines=[]
                chunk=""
                for word in content.split():
                    candidate=(chunk+" "+word).strip()
                    if chunk and draw.textbbox((0,0),candidate,font=regular)[2]>cell_w-20:
                        lines.append(chunk);chunk=word
                    else:chunk=candidate
                if chunk:lines.append(chunk)
                for i,line in enumerate(lines[:3]):
                    w=draw.textbbox((0,0),line,font=regular)[2]
                    draw.text((x+(cell_w-w)//2,y+22+i*34),
                              line,font=regular,fill="#182e43")
        y+=row_h
        if row_num==meal_after and has_meal:
            draw.rectangle((margin,y,width-margin,y+62),fill="#fff1d8",outline="#899daa",width=2)
            label=config.get("meal_label","ХАРЧУВАННЯ У ЇДАЛЬНІ")
            txtw=draw.textbbox((0,0),label,font=bold)[2]
            draw.text(((width-txtw)//2,y+12),label,font=bold,fill="#9b4b04")
            y+=62
    img.save(path,"JPEG",quality=94,subsampling=0)
    return Path(path)
