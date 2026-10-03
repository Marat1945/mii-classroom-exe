"""Локальна бібліотека Word з обережним зіставленням предмета, КТП й годин."""
from __future__ import annotations
import hashlib
import json
import os
import re
import shutil
import tempfile
from datetime import date
from pathlib import Path
from docx import Document
from .engine import ROOT, build_calendar, safe_name, word_path

LIBRARY_ROOT = ROOT / "Бібліотека уроків"
INDEX = LIBRARY_ROOT / "Індекс.json"
BANNED = ("РОБОЧА ЗАГОТОВКА", "Цей файл НЕ ПУБЛІКУВАТИ", "ще не згенеровано")

def norm(value):
    return re.sub(r"\s+"," ",re.sub(r"[«»“”\"'.,:;!?()\[\]–—-]"," ",str(value or "").casefold())).strip()

def stream_subject(stream):
    """Повертає (клас/рік навчання, предмет) без літери класу."""
    match=re.match(r"^\s*(\d+)(?:\s*[-–]\s*[А-ЯA-Zа-яa-z])?\s+(.+?)\s*$",stream)
    if not match:
        return (stream,norm(stream))
    grade=match.group(1)
    text=match.group(2).strip()
    return grade,norm(text)

def schedule_weekly_hours(config,stream):
    """Кількість 45-хвилинних уроків за два тижні / 2."""
    count=sum(
        (pair[0]==stream)+(pair[1]==stream)
        for pairs in config["days"].values() for pair in pairs
    )
    return count/2

def signature(lesson,config):
    meta=config["course_map"][lesson.stream]
    grade,subject=stream_subject(lesson.stream)
    return {
        "grade":grade,
        "subject":subject,
        "program":str(meta.get("curriculum_id") or meta["plan"]),
        "weekly_hours":schedule_weekly_hours(config,lesson.stream),
    }

def signature_key(lesson,config):
    return json.dumps(signature(lesson,config),ensure_ascii=False,sort_keys=True)

def parallel_matches(lesson,config,plans=None,calendar=None):
    """Суворий збіг програми/годин/теми/порядкового номера.
    За різного навантаження або КТП ніякого автоматичного перенесення.
    """
    source=signature(lesson,config)
    available=calendar if calendar is not None else build_calendar(config,plans)
    return [x for x in available
            if x.status=="готово"
            and signature(x,config)==source
            and x.lesson_number==lesson.lesson_number
            and norm(x.topic)==norm(lesson.topic)]

def check_real_docx(path):
    path=Path(path)
    if path.suffix.lower()!=".docx" or not path.is_file():
        raise ValueError("Потрібен наявний документ Word формату .docx")
    try: doc=Document(path)
    except Exception as ex:raise ValueError("Word-документ не відкривається або пошкоджений") from ex
    txt="\n".join(p.text for p in doc.paragraphs)
    if any(s.casefold() in txt.casefold() for s in BANNED):
        raise ValueError("Вибрано Word-заготовку. Виберіть справжню повну лекцію.")
    if len(txt.strip())<350:
        raise ValueError("Word містить замало тексту для повноцінної лекції; перевірте файл.")
    return len(txt)

def _save_index(index):
    LIBRARY_ROOT.mkdir(parents=True,exist_ok=True)
    fd,name=tempfile.mkstemp(prefix="tmp_lib_",suffix=".json",dir=LIBRARY_ROOT)
    try:
        with os.fdopen(fd,"w",encoding="utf8") as h:
            json.dump(index,h,ensure_ascii=False,indent=2)
            h.flush();os.fsync(h.fileno())
        Path(name).replace(INDEX)
    finally:
        if Path(name).exists():Path(name).unlink()

def load_index():
    if INDEX.exists():
        return json.loads(INDEX.read_text(encoding="utf8"))
    return {"version":1,"items":[]}

def _hash(path):
    h=hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda:stream.read(65536),b""):
            h.update(block)
    return h.hexdigest()

def add_document(path,lesson,config,origin="власний Word"):
    """Копія в бібліотеці; конфіденційні файли JSON сюди не копіюються."""
    check_real_docx(path)
    index=load_index()
    sha=_hash(path)
    sig=signature(lesson,config)
    # Reuse identical existing copy for same signature and topic.
    for item in index["items"]:
        if item["sha256"]==sha and item["signature"]==sig and item["topic_norm"]==norm(lesson.topic):
            if Path(item["path"]).is_file():return item
    LIBRARY_ROOT.mkdir(parents=True,exist_ok=True)
    key=hashlib.sha256((sha+signature_key(lesson,config)+norm(lesson.topic)).encode()).hexdigest()[:18]
    target=LIBRARY_ROOT/(f"{sig['grade']} клас {key}.docx")
    shutil.copy2(path,target)
    item={"id":key,"path":str(target),"sha256":sha,"signature":sig,
          "topic":lesson.topic,"topic_norm":norm(lesson.topic),
          "lesson_number":lesson.lesson_number,
          "origin":origin,"added":date.today().isoformat()}
    index["items"].append(item)
    _save_index(index)
    return item

def find_for_lesson(lesson,config):
    sig=signature(lesson,config)
    candidates=[item for item in load_index()["items"]
                if item["signature"]==sig and item["topic_norm"]==norm(lesson.topic)
                and Path(item["path"]).is_file()]
    # If same topic multiple times in a plan, prefer matching sequence index; if still
    # ambiguous, do not silently attach the wrong item.
    exact=[i for i in candidates if i["lesson_number"]==lesson.lesson_number]
    if len(exact)==1:return exact[0]
    if len(exact)>1:
        unique={i["sha256"] for i in exact}
        return exact[0] if len(unique)==1 else None
    return candidates[0] if len(candidates)==1 else None

def _change_text(paragraph,text):
    if paragraph.runs:
        paragraph.runs[0].text=text
        for run in paragraph.runs[1:]:run.text=""
    else:paragraph.add_run(text)

def _fix_header_dates(doc,lesson):
    """У колонтитулах дата уроку стає датою цього класу: «28.09» і «28.09.2026»."""
    short=lesson.day[8:10]+"."+lesson.day[5:7]
    full=short+"."+lesson.day[:4]
    pattern=re.compile(r"\b\d{2}\.\d{2}(\.\d{4})?\b")
    for section in doc.sections:
        for part in (section.header,section.footer):
            for paragraph in part.paragraphs:
                updated=pattern.sub(lambda m:full if m.group(1) else short,paragraph.text)
                if updated!=paragraph.text:_change_text(paragraph,updated)


def copy_for_lesson(master,lesson,label=None):
    """Окрема Word-копія для дати/ДЗ; оригінал не змінюється."""
    source=Path(master)
    check_real_docx(source)
    output=word_path(lesson)
    output.parent.mkdir(parents=True,exist_ok=True)
    doc=Document(source)
    date_ddmm=lesson.day[8:10]+"."+lesson.day[5:7]
    heading=f"{label or lesson.stream}, Урок {date_ddmm} — {lesson.topic}"
    # Перший рядок: «9-Б ВІ, Урок 05.10 — Тема» (новий) або «Урок 05.10 — Тема» (старі Word).
    regular=re.compile(r"^(?:.{1,40}?,\s*)?Урок\s+\d{1,2}\.\d{2}(?:\.\d{4})?\s*[—–-]",re.I)
    found_title=False
    date_full=lesson.day[8:10]+"."+lesson.day[5:7]+"."+lesson.day[:4]
    info_strip=re.compile(r"^Урок\s*№\s*\d+\s*[•·|\-–—]",re.I)
    for p in doc.paragraphs[:7]:
        if regular.search(p.text.strip()):
            _change_text(p,heading)
            found_title=True
            break
        if info_strip.search(p.text.strip()):
            # Інформаційна смуга Word від ChatGPT: «Урок № 8 • 02.10.2026».
            _change_text(p,f"Урок № {lesson.lesson_number} • {date_full}")
            found_title=True
            break
    _fix_header_dates(doc,lesson)
    if not found_title:
        p=doc.paragraphs[0].insert_paragraph_before(heading) if doc.paragraphs else doc.add_paragraph(heading)
        p.style="Heading 1"
    from .documents import ensure_closing
    ensure_closing(doc,lesson)
    if output.resolve()!=source.resolve():
        doc.save(output)
    return output

def attach_document(item,lessons,state,replace_existing=False,primary=None):
    """Автоприв'язка, але НІКОЛИ не переписує матеріали, для яких є чернетка.

    primary — урок, для якого Word створено; для паралелей в інші дні запам'ятовуємо, звідки він
    («наперед з 8-Б ГО, 05.10»), щоб учитель це бачив у програмі.
    """
    from .engine import same_day_label
    outcome=[]
    for lesson in lessons:
        key=lesson.unique_key
        if key in state.get("drafts",{}):
            continue
        current=state.get("files",{}).get(key)
        if current and current.get("validated") and Path(current["path"]).is_file():
            # Retain confirmed material unless the teacher deliberately chose a replacement.
            if not replace_existing:continue
        new_path=copy_for_lesson(item["path"],lesson,label=same_day_label(lesson,lessons))
        entry={"path":str(new_path),"validated":True,"complete":True,"library_id":item["id"]}
        if primary is not None and lesson.day!=primary.day:
            entry["from"]=f"{primary.stream}, {primary.day[8:10]}.{primary.day[5:7]}"
            entry["ahead"]=lesson.day>primary.day
        state.setdefault("files",{})[key]=entry
        outcome.append(lesson)
    return outcome
